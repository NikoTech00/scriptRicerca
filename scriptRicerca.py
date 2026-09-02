from __future__ import annotations

import argparse
import hashlib
import ipaddress
import json
import logging
import os
import re
import shutil
import socket
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, asdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

# Le dipendenze vengono caricate nel bootstrap, così il log viene creato
# anche se manca un pacchetto.
requests = None
BeautifulSoup = None
DDGS = None
load_dotenv = None
load_workbook = None
Alignment = Font = PatternFill = None
PdfReader = None
genai = None

VERSION = "V2.0 MEDICI - GOOGLE GROUNDING + CV VERIFICATO"

DEFAULT_SHEET = "Foglio1"
DEFAULT_WORKERS = 6
DEFAULT_AI_CONCURRENCY = 3
DEFAULT_SEARCH_CONCURRENCY = 3
DEFAULT_SEARCH_RESULTS = 8
DEFAULT_SAVE_EVERY = 20
DEFAULT_MODEL = "gemini-3.7-flash"

HTTP_TIMEOUT = 10
MAX_PDF_BYTES = 25 * 1024 * 1024
MAX_HTML_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_PDF_TEXT = 100_000

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0 Safari/537.36"
)

AI_SEMAPHORE: threading.BoundedSemaphore | None = None
SEARCH_SEMAPHORE: threading.BoundedSemaphore | None = None
AI_CIRCUIT_LOCK = threading.Lock()
AI_DISABLED_FOR_RUN = False
AI_DISABLE_REASON = ""

OUTPUT_COLUMNS = (
    "Ricerca_Stato",
    "Specialita",
    "Specialita_Confidenza",
    "Specialita_Evidenza",
    "CV_Salvato",
    "CV_URL",
    "CV_Confidenza",
    "Fonti_Ricerca",
    "Ricerca_Metodo",
    "Ricerca_Note",
    "Ultimo_Aggiornamento_UTC",
    "Ricerca_Errore",
)


# ============================================================
# MODELLI
# ============================================================

@dataclass(frozen=True)
class Person:
    row: int
    pers_id: str
    medico_id: str
    surname: str
    name: str
    birth_date: str
    fiscal_code: str
    city: str
    email: str

    @property
    def full_name(self) -> str:
        return f"{self.name} {self.surname}".strip()

    @property
    def output_code(self) -> str:
        mid = numericish(clean(self.medico_id))
        if mid and mid not in {"0", "0.0"}:
            return mid
        return numericish(clean(self.pers_id))


@dataclass(frozen=True)
class WebHit:
    title: str
    url: str
    snippet: str
    provider: str = ""


# ============================================================
# BOOTSTRAP
# ============================================================

def clean(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.casefold() in {"none", "null", "nan", "n/a", "n.d.", "nd"}:
        return ""
    return s


def numericish(value: str) -> str:
    s = clean(value)
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    return s


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def normalize(text: str) -> str:
    s = clean(text).casefold()
    s = s.replace("’", "'").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", s).strip()


def safe_part(value: str) -> str:
    s = re.sub(r'[<>:"/\\|?*]', "_", clean(value))
    s = re.sub(r"\s+", "_", s).strip(" ._")
    return s or "ND"


def canonical_url(url: str) -> str:
    return clean(url).split("#", 1)[0].rstrip("/")


def domain(url: str) -> str:
    try:
        return (urlparse(url).hostname or "").casefold()
    except Exception:
        return ""


def bootstrap_log_path(log_dir: Path) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    return log_dir / f"startup_medici_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"


def write_bootstrap_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | {message}\n")


def import_dependencies(log_path: Path) -> tuple[bool, list[str]]:
    global requests, BeautifulSoup, DDGS, load_dotenv
    global load_workbook, Alignment, Font, PatternFill, PdfReader, genai

    missing: list[str] = []

    try:
        import requests as _requests
        requests = _requests
    except Exception as exc:
        missing.append(f"requests ({exc})")

    try:
        from bs4 import BeautifulSoup as _BeautifulSoup
        BeautifulSoup = _BeautifulSoup
    except Exception as exc:
        missing.append(f"beautifulsoup4 ({exc})")

    try:
        from dotenv import load_dotenv as _load_dotenv
        load_dotenv = _load_dotenv
    except Exception as exc:
        missing.append(f"python-dotenv ({exc})")

    try:
        from openpyxl import load_workbook as _load_workbook
        from openpyxl.styles import Alignment as _Alignment, Font as _Font, PatternFill as _PatternFill
        load_workbook = _load_workbook
        Alignment, Font, PatternFill = _Alignment, _Font, _PatternFill
    except Exception as exc:
        missing.append(f"openpyxl ({exc})")

    try:
        from pypdf import PdfReader as _PdfReader
        PdfReader = _PdfReader
    except Exception as exc:
        missing.append(f"pypdf ({exc})")

    # Gemini è il motore principale, ma teniamo provider esterni come fallback.
    try:
        from google import genai as _genai
        genai = _genai
    except Exception as exc:
        genai = None
        write_bootstrap_log(
            log_path,
            f"WARNING | google-genai non disponibile: {type(exc).__name__}: {exc}"
        )

    try:
        from ddgs import DDGS as _DDGS
        DDGS = _DDGS
    except Exception:
        DDGS = None  # fallback opzionale, non blocca

    if missing:
        write_bootstrap_log(log_path, "ERRORE | Dipendenze mancanti: " + "; ".join(missing))
        return False, missing

    write_bootstrap_log(log_path, "OK | Dipendenze obbligatorie caricate.")
    return True, []


def parse_early_paths(argv: list[str]) -> tuple[Path | None, Path | None, Path]:
    input_path = None
    output_path = None
    log_dir = Path("logs")

    for arg in argv[1:]:
        if not arg.startswith("-"):
            input_path = Path(arg)
            break

    if "--output" in argv:
        try:
            output_path = Path(argv[argv.index("--output") + 1])
        except Exception:
            pass

    if "--log-dir" in argv:
        try:
            log_dir = Path(argv[argv.index("--log-dir") + 1])
        except Exception:
            pass

    return input_path, output_path, log_dir


def default_output_path(input_path: Path) -> Path:
    return Path.cwd() / "output" / f"{input_path.stem}_specialita_cv.xlsx"


def create_initial_output_copy(input_path: Path | None, output_path: Path | None) -> Path | None:
    if input_path is None:
        return None

    inp = input_path.expanduser().resolve()
    out = output_path.expanduser().resolve() if output_path else default_output_path(inp).resolve()

    if inp.exists() and inp.suffix.casefold() == ".xlsx":
        out.parent.mkdir(parents=True, exist_ok=True)
        if not out.exists() and inp != out:
            shutil.copy2(inp, out)
    return out


# ============================================================
# CLI / LOG
# ============================================================

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description=f"{VERSION}: ricerca web grounded della specialità e CV dei medici."
    )
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--sheet", default=DEFAULT_SHEET)
    p.add_argument("--limit", "--max-rows", dest="limit", type=int)
    p.add_argument("--start-row", type=int, default=2)

    p.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    p.add_argument("--ai-concurrency", type=int, default=DEFAULT_AI_CONCURRENCY)
    p.add_argument("--search-concurrency", type=int, default=DEFAULT_SEARCH_CONCURRENCY)
    p.add_argument("--search-results", type=int, default=DEFAULT_SEARCH_RESULTS)

    p.add_argument("--model", default=os.getenv("GEMINI_MODEL", DEFAULT_MODEL))
    p.add_argument(
        "--provider",
        choices=("auto", "gemini", "serper", "brave", "ddgs"),
        default="auto",
        help="auto: Gemini Google Search -> Serper/Brave -> DDGS fallback.",
    )
    p.add_argument(
        "--deep",
        action="store_true",
        help="Seconda ricerca mirata se il primo passaggio non trova CV o specialità.",
    )
    p.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    p.add_argument("--retry-all", action="store_true")
    p.add_argument("--retry-errors", action="store_true")
    p.add_argument("--ignore-cache", action="store_true")

    p.add_argument("--cv-dir", type=Path, default=Path("cv_medici"))
    p.add_argument("--cache-dir", type=Path, default=Path("cache_medici"))
    p.add_argument("--log-dir", type=Path, default=Path("logs"))

    args = p.parse_args()

    for field in ("workers", "ai_concurrency", "search_concurrency", "search_results", "save_every"):
        if getattr(args, field) <= 0:
            p.error(f"--{field.replace('_', '-')} deve essere > 0")
    if args.limit is not None and args.limit <= 0:
        p.error("--limit deve essere > 0")
    if args.start_row < 2:
        p.error("--start-row deve essere >= 2")
    return args


def configure_logging(log_dir: Path) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    path = log_dir / f"run_medici_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

    root = logging.getLogger()
    root.handlers.clear()
    root.setLevel(logging.INFO)

    fmt = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(threadName)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    fh = logging.FileHandler(path, encoding="utf-8")
    sh = logging.StreamHandler()
    fh.setFormatter(fmt)
    sh.setFormatter(fmt)
    root.addHandler(fh)
    root.addHandler(sh)
    return path


# ============================================================
# EXCEL
# ============================================================

def headers_map(ws) -> dict[str, int]:
    out = {}
    for c in ws[1]:
        label = clean(c.value)
        if label:
            out[label.casefold()] = c.column
    return out


def require_input_columns(headers: dict[str, int]) -> None:
    req = ("Pers_Id", "Pers_Cognome", "Pers_Nome")
    missing = [x for x in req if x.casefold() not in headers]
    if missing:
        raise ValueError("Colonne obbligatorie mancanti: " + ", ".join(missing))


def ensure_output_columns(ws) -> dict[str, int]:
    headers = headers_map(ws)
    for label in OUTPUT_COLUMNS:
        if label.casefold() in headers:
            continue
        col = ws.max_column + 1
        cell = ws.cell(1, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers[label.casefold()] = col
    ws.freeze_panes = "A2"
    return headers


def getv(ws, row: int, headers: dict[str, int], label: str) -> Any:
    col = headers.get(label.casefold())
    return ws.cell(row, col).value if col else ""


def person_from_row(ws, row: int, headers: dict[str, int]) -> Person:
    dob = getv(ws, row, headers, "Pers_DataNascita")
    if isinstance(dob, (datetime, date)):
        dob = dob.strftime("%d/%m/%Y")
    return Person(
        row=row,
        pers_id=numericish(clean(getv(ws, row, headers, "Pers_Id"))),
        medico_id=numericish(clean(getv(ws, row, headers, "Medico_Id"))),
        surname=clean(getv(ws, row, headers, "Pers_Cognome")),
        name=clean(getv(ws, row, headers, "Pers_Nome")),
        birth_date=clean(dob),
        fiscal_code=clean(getv(ws, row, headers, "Pers_CodFis")).upper(),
        city=clean(getv(ws, row, headers, "Indirizzi_Citta")),
        email=clean(getv(ws, row, headers, "emailPredefinita"))
              or clean(getv(ws, row, headers, "Email")),
    )


def prepare_output(input_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not output_path.exists():
        shutil.copy2(input_path, output_path)


def atomic_save(wb, path: Path) -> None:
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def set_result(ws, headers: dict[str, int], row: int, result: dict[str, Any]) -> None:
    mapping = {
        "Ricerca_Stato": result.get("status", ""),
        "Specialita": result.get("specialty", ""),
        "Specialita_Confidenza": result.get("specialty_confidence", ""),
        "Specialita_Evidenza": result.get("specialty_evidence", ""),
        "CV_Salvato": result.get("cv_path", ""),
        "CV_URL": result.get("cv_url", ""),
        "CV_Confidenza": result.get("cv_confidence", ""),
        "Fonti_Ricerca": result.get("sources", ""),
        "Ricerca_Metodo": result.get("method", ""),
        "Ricerca_Note": result.get("notes", ""),
        "Ultimo_Aggiornamento_UTC": result.get("updated", utc_now()),
        "Ricerca_Errore": result.get("error", ""),
    }
    for label, value in mapping.items():
        ws.cell(row, headers[label.casefold()], value)


# ============================================================
# CACHE
# ============================================================

def cache_path(person: Person, cache_dir: Path) -> Path:
    identity = f"{VERSION}|{person.pers_id}|{person.full_name}|{person.birth_date}|{person.fiscal_code}"
    digest = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:14]
    return cache_dir / f"{safe_part(person.pers_id)}_{digest}.json"


def read_cache(person: Person, cache_dir: Path) -> dict[str, Any] | None:
    p = cache_path(person, cache_dir)
    try:
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            if data.get("version") == VERSION:
                return data.get("result")
    except Exception:
        pass
    return None


def write_cache(person: Person, cache_dir: Path, result: dict[str, Any]) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    p = cache_path(person, cache_dir)
    payload = {"version": VERSION, "person": asdict(person), "result": result}
    fd, tmp_name = tempfile.mkstemp(prefix=".cache_", suffix=".json", dir=cache_dir)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, p)
    finally:
        tmp.unlink(missing_ok=True)


# ============================================================
# SAFE HTTP / PDF
# ============================================================

def is_public_http_url(url: str) -> bool:
    try:
        p = urlparse(url)
        if p.scheme not in {"http", "https"} or not p.hostname:
            return False
        try:
            infos = socket.getaddrinfo(
                p.hostname,
                p.port or (443 if p.scheme == "https" else 80),
                type=socket.SOCK_STREAM,
            )
            for info in infos:
                ip = ipaddress.ip_address(info[4][0])
                if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                    return False
        except OSError:
            pass
        return True
    except Exception:
        return False


def get_bytes(url: str, max_bytes: int) -> tuple[Any, bytes] | None:
    if not is_public_http_url(url):
        return None
    try:
        r = requests.get(
            url,
            headers={"User-Agent": USER_AGENT, "Accept-Language": "it-IT,it;q=0.9,en;q=0.5"},
            timeout=HTTP_TIMEOUT,
            allow_redirects=True,
            stream=True,
        )
        r.raise_for_status()
        buf = bytearray()
        for chunk in r.iter_content(65536):
            if chunk:
                buf.extend(chunk)
                if len(buf) > max_bytes:
                    r.close()
                    return None
        return r, bytes(buf)
    except Exception:
        return None


def pdf_text(raw: bytes) -> str:
    import io
    try:
        reader = PdfReader(io.BytesIO(raw))
        chunks = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            t = page.extract_text() or ""
            if t:
                chunks.append(t)
        return "\n".join(chunks)[:MAX_PDF_TEXT]
    except Exception:
        return ""


CV_POSITIVE = (
    "curriculum vitae", "curriculum professionale", "curriculum formativo",
    "europass", "esperienza professionale", "esperienze professionali",
    "istruzione e formazione", "titoli di studio", "attività professionale",
)

CV_NEGATIVE = (
    "graduatoria", "elenco candidati", "candidati ammessi", "verbale",
    "bando", "avviso pubblico", "prova orale", "prova scritta",
    "commissione esaminatrice", "delibera", "deliberazione",
)

MEDICAL_TERMS = (
    "medico", "medicina", "chirurgo", "specialista", "specializzazione",
    "ospedale", "azienda sanitaria", "asl", "ausl", "asst", "aou",
    "irccs", "policlinico", "ordine dei medici",
)


def identity_score(text: str, person: Person) -> int:
    n = normalize(text)
    full = normalize(person.full_name)
    reverse = normalize(f"{person.surname} {person.name}")
    score = 0
    if full and full in n:
        score += 180
    if reverse and reverse in n:
        score += 180
    if normalize(person.surname) in n:
        score += 50
    if normalize(person.name) in n:
        score += 35
    if person.fiscal_code and normalize(person.fiscal_code) in n:
        score += 600
    return score


def date_variants(dob: str) -> set[str]:
    if not dob:
        return set()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            d = datetime.strptime(dob[:10], fmt)
            return {
                d.strftime("%d/%m/%Y"), d.strftime("%d-%m-%Y"),
                d.strftime("%d.%m.%Y"), d.strftime("%Y-%m-%d")
            }
        except Exception:
            continue
    return {dob}


def verify_cv(text: str, person: Person) -> tuple[bool, str, str]:
    n = normalize(text)

    if identity_score(text, person) < 180:
        return False, "bassa", "Nome/cognome non verificati nel documento."

    # Codice fiscale: se nel PDF ce n'è uno diverso, rifiuta.
    if person.fiscal_code:
        matches = re.findall(r"\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b", text.upper())
        if matches and person.fiscal_code not in matches:
            return False, "bassa", "Codice fiscale incompatibile."

    # Se il CV dichiara esplicitamente una DOB e non coincide, rifiuta.
    if person.birth_date:
        m = re.search(
            r"(?:data\s+di\s+nascita|nato\s+il|nata\s+il).{0,50}"
            r"(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
            text, re.I | re.S,
        )
        if m:
            found = m.group(1).replace(".", "/").replace("-", "/")
            wanted = {x.replace(".", "/").replace("-", "/") for x in date_variants(person.birth_date)}
            if found not in wanted:
                return False, "bassa", "Data di nascita incompatibile."

    pos = sum(1 for x in CV_POSITIVE if x in n)
    neg = sum(1 for x in CV_NEGATIVE if x in n)
    med = sum(1 for x in MEDICAL_TERMS if x in n)

    if pos < 2 and "curriculum vitae" not in n and "europass" not in n:
        return False, "bassa", "Il PDF non appare un CV."
    if med == 0:
        return False, "bassa", "Manca il contesto medico."
    if neg >= 3 and "curriculum vitae" not in n:
        return False, "bassa", "Documento concorsuale/amministrativo."

    conf = "alta" if (person.fiscal_code and normalize(person.fiscal_code) in n) or pos >= 4 else "media"
    return True, conf, "CV verificato sul contenuto."


def cv_destination(person: Person, cv_dir: Path) -> Path:
    return cv_dir / f"{safe_part(person.output_code)}-{safe_part(person.surname)}-{safe_part(person.name)}.pdf"


def try_pdf_url(url: str, person: Person, cv_dir: Path) -> tuple[str, str, str] | None:
    got = get_bytes(url, MAX_PDF_BYTES)
    if not got:
        return None
    r, raw = got
    try:
        ctype = clean(r.headers.get("Content-Type")).casefold()
        if not (raw.startswith(b"%PDF") or "application/pdf" in ctype):
            return None
        text = pdf_text(raw)
        if not text:
            return None
        ok, conf, reason = verify_cv(text, person)
        if not ok:
            logging.info("CV SCARTATO | %s | %s", url, reason)
            return None
        cv_dir.mkdir(parents=True, exist_ok=True)
        dest = cv_destination(person, cv_dir)
        dest.write_bytes(raw)
        logging.info("CV VERIFICATO | %s | %s", person.pers_id, dest)
        return str(dest), text, conf
    finally:
        r.close()


def try_cv_candidate(url: str, person: Person, cv_dir: Path) -> tuple[str, str, str, str] | None:
    direct = try_pdf_url(url, person, cv_dir)
    if direct:
        path, text, conf = direct
        return path, url, text, conf

    got = get_bytes(url, MAX_HTML_BYTES)
    if not got:
        return None
    r, raw = got
    try:
        soup = BeautifulSoup(raw, "html.parser")
        links: list[tuple[int, str]] = []
        for a in soup.select("a[href]"):
            href = clean(a.get("href"))
            if not href:
                continue
            candidate = urljoin(url, href)
            label = normalize(f"{a.get_text(' ', strip=True)} {candidate}")
            score = 0
            if ".pdf" in candidate.casefold():
                score += 50
            if "curriculum" in label or "europass" in label:
                score += 100
            if normalize(person.surname) in label:
                score += 40
            if normalize(person.name) in label:
                score += 25
            if any(x in label for x in CV_NEGATIVE):
                score -= 100
            if score >= 80:
                links.append((score, candidate))

        seen = set()
        for _, candidate in sorted(links, reverse=True)[:12]:
            key = canonical_url(candidate)
            if key in seen:
                continue
            seen.add(key)
            found = try_pdf_url(candidate, person, cv_dir)
            if found:
                path, text, conf = found
                return path, candidate, text, conf
    finally:
        r.close()
    return None


# ============================================================
# GEMINI GROUNDED SEARCH
# ============================================================

GROUND_SCHEMA = {
    "type": "object",
    "properties": {
        "identity_confidence": {
            "type": "string",
            "enum": ["alta", "media", "bassa", "nessuna"],
        },
        "specialty_found": {"type": "boolean"},
        "specialty": {"type": "string"},
        "specialty_confidence": {
            "type": "string",
            "enum": ["alta", "media", "bassa", "nessuna"],
        },
        "specialty_evidence": {"type": "string"},
        "specialty_source_urls": {
            "type": "array",
            "items": {"type": "string"},
        },
        "cv_candidates": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "url": {"type": "string"},
                    "title": {"type": "string"},
                    "confidence": {
                        "type": "string",
                        "enum": ["alta", "media", "bassa"],
                    },
                },
                "required": ["url", "title", "confidence"],
            },
        },
        "other_source_urls": {
            "type": "array",
            "items": {"type": "string"},
        },
        "notes": {"type": "string"},
    },
    "required": [
        "identity_confidence",
        "specialty_found",
        "specialty",
        "specialty_confidence",
        "specialty_evidence",
        "specialty_source_urls",
        "cv_candidates",
        "other_source_urls",
        "notes",
    ],
}


def gemini_prompt(person: Person, deep: bool = False) -> str:
    extra = (
        "\nEsegui una ricerca più approfondita e mirata: prova anche siti di ASL/AUSL/ASST/AOU/"
        "IRCCS, università, amministrazione trasparente, ordini dei medici e PDF curriculum."
        if deep else ""
    )

    return f"""
Stai facendo una ricerca professionale sul web su UNO specifico medico italiano.

PERSONA DA IDENTIFICARE
- Nome: {person.name}
- Cognome: {person.surname}
- Data di nascita: {person.birth_date or "non disponibile"}
- Codice fiscale: {person.fiscal_code or "non disponibile"}
- Città nel database: {person.city or "non disponibile"}
- Email nel database: {person.email or "non disponibile"}

OBIETTIVI
1. Identificare la SPECIALITÀ MEDICA effettiva della persona.
2. Trovare il suo CURRICULUM VITAE, preferibilmente PDF ufficiale o pubblicato
   da ospedale, ASL/AUSL/ASST/AOU, IRCCS, università, ordine professionale,
   ente pubblico o struttura sanitaria.

REGOLE CRITICHE
- Usa Google Search.
- Non inventare nulla.
- Non dedurre automaticamente la specialità dal semplice reparto in cui lavora.
- "Medico chirurgo", "laureato in medicina e chirurgia" e "dirigente medico"
  NON sono specialità.
- Considera valida una specialità solo se una fonte riferita alla stessa persona
  dichiara chiaramente "specialista in", "specializzazione in", un titolo equivalente,
  oppure un CV verificabile la documenta.
- Usa data di nascita, codice fiscale, città ed email solo per disambiguare gli omonimi.
- Se ci sono più omonimi e non riesci a distinguere la persona, abbassa la confidenza
  o restituisci specialty_found=false.
- Per il CV restituisci URL REALI trovati nella ricerca. Non costruire URL.
- Non indicare come CV graduatorie, elenchi candidati, verbali, bandi, delibere,
  determinazioni o pagine che citano semplicemente il medico.
- Dai priorità a fonti istituzionali e sanitarie.
- specialty_evidence deve essere una frase breve che descrive cosa dimostra la fonte,
  senza inventare citazioni testuali.
{extra}

Restituisci esclusivamente il JSON richiesto dallo schema.
""".strip()


def recursively_extract_urls(obj: Any) -> list[str]:
    urls: list[str] = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in {"url", "uri"} and isinstance(v, str) and v.startswith(("http://", "https://")):
                urls.append(v)
            urls.extend(recursively_extract_urls(v))
    elif isinstance(obj, list):
        for v in obj:
            urls.extend(recursively_extract_urls(v))
    return urls


def is_quota_error(exc: Exception) -> bool:
    s = f"{type(exc).__name__}: {exc}".casefold()
    return any(x in s for x in ("429", "quota", "rate limit", "resource_exhausted", "too many requests"))


def disable_ai_for_run(reason: str) -> None:
    global AI_DISABLED_FOR_RUN, AI_DISABLE_REASON
    with AI_CIRCUIT_LOCK:
        AI_DISABLED_FOR_RUN = True
        AI_DISABLE_REASON = reason


def gemini_grounded(person: Person, model: str, deep: bool = False) -> dict[str, Any] | None:
    global AI_DISABLED_FOR_RUN

    if genai is None or not os.getenv("GEMINI_API_KEY"):
        return None
    if AI_DISABLED_FOR_RUN:
        return None
    if AI_SEMAPHORE is None:
        raise RuntimeError("AI semaphore non inizializzato")

    with AI_SEMAPHORE:
        try:
            client = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
            interaction = client.interactions.create(
                model=model,
                input=gemini_prompt(person, deep=deep),
                tools=[{"type": "google_search"}],
                response_format={
                    "type": "text",
                    "mime_type": "application/json",
                    "schema": GROUND_SCHEMA,
                },
            )
            raw = clean(interaction.output_text)
            data = json.loads(raw)

            # Le annotazioni di grounding forniscono fonti aggiuntive reali.
            try:
                dumped = interaction.model_dump()
                annotation_urls = recursively_extract_urls(dumped)
            except Exception:
                annotation_urls = []

            data["_annotation_urls"] = list(dict.fromkeys(annotation_urls))
            logging.info(
                "GEMINI SEARCH OK | Pers_Id=%s | specialita=%s | cv_candidates=%s | deep=%s",
                person.pers_id,
                clean(data.get("specialty")) or "N/D",
                len(data.get("cv_candidates") or []),
                deep,
            )
            return data

        except Exception as exc:
            if is_quota_error(exc):
                reason = f"{type(exc).__name__}: {exc}"
                disable_ai_for_run(reason)
                logging.error("GEMINI CIRCUIT BREAKER | %s", reason)
            else:
                logging.warning(
                    "GEMINI SEARCH FAIL | Pers_Id=%s | %s: %s",
                    person.pers_id, type(exc).__name__, exc
                )
            return None


# ============================================================
# SEARCH API FALLBACK
# ============================================================

def search_serper(query: str, max_results: int) -> list[WebHit]:
    key = os.getenv("SERPER_API_KEY")
    if not key:
        return []
    r = requests.post(
        "https://google.serper.dev/search",
        headers={"X-API-KEY": key, "Content-Type": "application/json"},
        json={"q": query, "gl": "it", "hl": "it", "num": max_results},
        timeout=HTTP_TIMEOUT,
    )
    r.raise_for_status()
    data = r.json()
    return [
        WebHit(clean(x.get("title")), clean(x.get("link")), clean(x.get("snippet")), "serper")
        for x in data.get("organic", [])[:max_results]
        if clean(x.get("link"))
    ]


def search_brave(query: str, max_results: int) -> list[WebHit]:
    key = os.getenv("BRAVE_SEARCH_API_KEY")
    if not key:
        return []
    r = requests.get(
        "https://api.search.brave.com/res/v1/web/search",
        headers={"X-Subscription-Token": key, "Accept": "application/json"},
        params={"q": query, "count": max_results, "country": "IT", "search_lang": "it"},
        timeout=HTTP_TIMEOUT,
    )
    r.raise_for_status()
    data = r.json()
    return [
        WebHit(clean(x.get("title")), clean(x.get("url")), clean(x.get("description")), "brave")
        for x in data.get("web", {}).get("results", [])[:max_results]
        if clean(x.get("url"))
    ]


def search_ddgs(query: str, max_results: int) -> list[WebHit]:
    if DDGS is None:
        return []
    raw = DDGS(timeout=HTTP_TIMEOUT).text(
        query, region="it-it", safesearch="moderate",
        max_results=max_results, backend="duckduckgo"
    ) or []
    return [
        WebHit(
            clean(x.get("title")),
            clean(x.get("href") or x.get("url")),
            clean(x.get("body") or x.get("snippet")),
            "ddgs",
        )
        for x in raw
        if clean(x.get("href") or x.get("url"))
    ]


def external_search(query: str, max_results: int, provider: str = "auto") -> list[WebHit]:
    if SEARCH_SEMAPHORE is None:
        raise RuntimeError("Search semaphore non inizializzato")

    with SEARCH_SEMAPHORE:
        chain = []
        if provider in {"auto", "serper"} and os.getenv("SERPER_API_KEY"):
            chain.append(("serper", lambda: search_serper(query, max_results)))
        if provider in {"auto", "brave"} and os.getenv("BRAVE_SEARCH_API_KEY"):
            chain.append(("brave", lambda: search_brave(query, max_results)))
        if provider in {"auto", "ddgs"} and DDGS is not None:
            chain.append(("ddgs", lambda: search_ddgs(query, max_results)))

        for name, fn in chain:
            try:
                hits = fn()
                if hits:
                    logging.info("SEARCH API OK | %s | %s | %s risultati", name, query, len(hits))
                    return dedupe_hits(hits)
            except Exception as exc:
                logging.warning("SEARCH API FAIL | %s | %s: %s", name, type(exc).__name__, exc)
        return []


def dedupe_hits(hits: Iterable[WebHit]) -> list[WebHit]:
    out, seen = [], set()
    for hit in hits:
        key = canonical_url(hit.url)
        if key and key not in seen:
            seen.add(key)
            out.append(hit)
    return out


def fallback_queries(person: Person) -> list[str]:
    full = f'"{person.full_name}"'
    qs = [
        f'{full} medico "specialista in"',
        f'{full} medico "specializzazione in"',
        f'{full} "curriculum vitae" filetype:pdf',
    ]
    if person.city:
        qs.append(f'{full} medico "{person.city}" curriculum')
    return qs


SPECIALTY_HINT = re.compile(
    r"(?:specialista|specializzato|specializzata|specializzazione)\s+in\s+"
    r"([A-ZÀ-ÖØ-Ýa-zà-öø-ÿ][A-ZÀ-ÖØ-Ýa-zà-öø-ÿ0-9 '&/().,+\-]{2,90})",
    re.I,
)


def extract_specialty_from_hits(hits: list[WebHit], person: Person) -> tuple[str, str, str, str]:
    candidates = []
    for hit in hits:
        text = f"{hit.title} {hit.snippet}"
        if identity_score(text, person) < 180:
            continue
        for m in SPECIALTY_HINT.finditer(text):
            spec = clean(m.group(1))
            spec = re.split(r"[|;•]", spec)[0].strip(" .,-")
            if not spec:
                continue
            score = identity_score(text, person)
            d = domain(hit.url)
            if any(x in d for x in ("asl", "ausl", "asst", "aou", "irccs", "osped", "policlin", "univ", "ordinemedici")):
                score += 150
            candidates.append((score, spec, hit.url, clean(hit.snippet)))

    if not candidates:
        return "", "", "", ""

    candidates.sort(reverse=True)
    best = candidates[0]
    if len(candidates) > 1 and normalize(candidates[1][1]) != normalize(best[1]) and candidates[1][0] >= best[0] - 30:
        return "", "", "", ""
    return best[1], ("alta" if best[0] >= 400 else "media"), best[2], best[3]


# ============================================================
# PIPELINE PERSONA
# ============================================================

def collect_urls_from_ground(data: dict[str, Any]) -> list[str]:
    urls = []
    for key in ("specialty_source_urls", "other_source_urls", "_annotation_urls"):
        for url in data.get(key) or []:
            if isinstance(url, str) and url.startswith(("http://", "https://")):
                urls.append(url)
    for item in data.get("cv_candidates") or []:
        url = clean(item.get("url"))
        if url.startswith(("http://", "https://")):
            urls.append(url)
    return list(dict.fromkeys(urls))


def validate_ground_specialty(data: dict[str, Any]) -> tuple[str, str, str]:
    if not data.get("specialty_found"):
        return "", "nessuna", ""

    identity_conf = clean(data.get("identity_confidence"))
    spec_conf = clean(data.get("specialty_confidence"))
    specialty = clean(data.get("specialty"))
    evidence = clean(data.get("specialty_evidence"))

    # Non accettiamo risultati deboli come dati definitivi.
    if identity_conf not in {"alta", "media"}:
        return "", "nessuna", ""
    if spec_conf not in {"alta", "media"}:
        return "", "nessuna", ""
    if not specialty:
        return "", "nessuna", ""

    n = normalize(specialty)
    invalid = (
        "medico chirurgo", "dirigente medico", "medicina e chirurgia",
        "medico", "chirurgo",
    )
    if n in invalid:
        return "", "nessuna", ""

    return specialty, spec_conf, evidence


def research_person(
    person: Person,
    args: argparse.Namespace,
    cv_dir: Path,
    cache_dir: Path,
) -> dict[str, Any]:
    if not args.ignore_cache:
        cached = read_cache(person, cache_dir)
        if cached:
            cached["row"] = person.row
            cached["method"] = clean(cached.get("method")) + " + CACHE"
            return cached

    started = time.perf_counter()
    sources: list[str] = []
    notes: list[str] = []
    method_parts: list[str] = []

    specialty = ""
    specialty_conf = "nessuna"
    specialty_evidence = ""
    cv_path = ""
    cv_url = ""
    cv_conf = "nessuna"

    grounded: dict[str, Any] | None = None

    # 1) GROUNDING GOOGLE come motore principale
    if args.provider in {"auto", "gemini"}:
        grounded = gemini_grounded(person, args.model, deep=False)
        if grounded:
            method_parts.append("Gemini Google Search")
            sources.extend(collect_urls_from_ground(grounded))

            specialty, specialty_conf, specialty_evidence = validate_ground_specialty(grounded)

            # Verifica davvero i CV proposti dal modello.
            cv_candidates = grounded.get("cv_candidates") or []
            cv_candidates = sorted(
                cv_candidates,
                key=lambda x: {"alta": 3, "media": 2, "bassa": 1}.get(clean(x.get("confidence")), 0),
                reverse=True,
            )
            for item in cv_candidates[:5]:
                candidate = clean(item.get("url"))
                if not candidate.startswith(("http://", "https://")):
                    continue
                found = try_cv_candidate(candidate, person, cv_dir)
                if found:
                    cv_path, cv_url, cv_text, cv_conf = found
                    sources.insert(0, cv_url)
                    notes.append("CV trovato tramite Google Search e verificato sul PDF.")
                    break

    # 2) Secondo passaggio grounded, SOLO se richiesto e manca qualcosa.
    if args.deep and args.provider in {"auto", "gemini"} and (not specialty or not cv_path):
        deep_data = gemini_grounded(person, args.model, deep=True)
        if deep_data:
            method_parts.append("Gemini Deep Search")
            sources.extend(collect_urls_from_ground(deep_data))

            if not specialty:
                specialty, specialty_conf, specialty_evidence = validate_ground_specialty(deep_data)

            if not cv_path:
                for item in (deep_data.get("cv_candidates") or [])[:6]:
                    candidate = clean(item.get("url"))
                    if not candidate.startswith(("http://", "https://")):
                        continue
                    found = try_cv_candidate(candidate, person, cv_dir)
                    if found:
                        cv_path, cv_url, cv_text, cv_conf = found
                        sources.insert(0, cv_url)
                        notes.append("CV trovato nel secondo passaggio e verificato.")
                        break

    # 3) Search API fallback se Gemini non ha prodotto abbastanza.
    need_fallback = not specialty or not cv_path
    if need_fallback and args.provider != "gemini":
        provider_for_api = args.provider if args.provider in {"serper", "brave", "ddgs"} else "auto"
        all_hits: list[WebHit] = []

        for q in fallback_queries(person):
            hits = external_search(q, args.search_results, provider_for_api)
            all_hits.extend(hits)

            # Prova CV subito dai risultati più pertinenti.
            if not cv_path:
                ranked = sorted(
                    hits,
                    key=lambda h: (
                        100 if "curriculum" in normalize(f"{h.title} {h.snippet} {h.url}") else 0
                    ) + identity_score(f"{h.title} {h.snippet}", person),
                    reverse=True,
                )
                for hit in ranked[:4]:
                    if identity_score(f"{hit.title} {hit.snippet}", person) < 180:
                        continue
                    found = try_cv_candidate(hit.url, person, cv_dir)
                    if found:
                        cv_path, cv_url, cv_text, cv_conf = found
                        sources.insert(0, cv_url)
                        notes.append(f"CV verificato tramite fallback {hit.provider}.")
                        break

            if not specialty:
                s, c, u, ev = extract_specialty_from_hits(dedupe_hits(all_hits), person)
                if s:
                    specialty, specialty_conf, specialty_evidence = s, c, ev
                    sources.append(u)

            if specialty and cv_path:
                break

        if all_hits:
            method_parts.append("Search API fallback")
            sources.extend([h.url for h in dedupe_hits(all_hits)[:6]])

    # 4) Se Gemini è fuori quota, distinguiamolo da "nessun risultato".
    if not grounded and AI_DISABLED_FOR_RUN and args.provider in {"auto", "gemini"}:
        notes.append("Gemini Google Search disabilitato per quota/rate limit durante questa esecuzione.")

    sources = list(dict.fromkeys(u for u in sources if clean(u).startswith(("http://", "https://"))))

    if specialty and cv_path:
        status = "COMPLETATO"
    elif specialty:
        status = "SPECIALITA_TROVATA"
    elif cv_path:
        status = "CV_TROVATO_SPECIALITA_DA_VERIFICARE"
    elif AI_DISABLED_FOR_RUN and not sources:
        status = "BLOCCATO_QUOTA_RICERCA"
    elif sources:
        status = "DA_VERIFICARE"
    else:
        status = "NESSUN_RISULTATO"

    elapsed = time.perf_counter() - started

    result = {
        "row": person.row,
        "status": status,
        "specialty": specialty,
        "specialty_confidence": specialty_conf,
        "specialty_evidence": specialty_evidence,
        "cv_path": cv_path,
        "cv_url": cv_url,
        "cv_confidence": cv_conf,
        "sources": "\n".join(sources[:10]),
        "method": " | ".join(method_parts) or "N/D",
        "notes": " ".join(notes),
        "updated": utc_now(),
        "error": "",
        "elapsed": elapsed,
    }

    write_cache(person, cache_dir, result)

    logging.info(
        "DONE | Pers_Id=%s | %.1fs | %s | specialita=%s | cv=%s | metodo=%s",
        person.pers_id, elapsed, status, specialty or "N/D",
        "SI" if cv_path else "NO", result["method"]
    )
    return result


# ============================================================
# MAIN
# ============================================================

def main() -> int:
    global AI_SEMAPHORE, SEARCH_SEMAPHORE

    args = parse_args()
    AI_SEMAPHORE = threading.BoundedSemaphore(args.ai_concurrency)
    SEARCH_SEMAPHORE = threading.BoundedSemaphore(args.search_concurrency)

    log_file = configure_logging(args.log_dir)

    input_path = args.input.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else default_output_path(input_path).resolve()
    )
    cv_dir = args.cv_dir.expanduser().resolve()
    cache_dir = args.cache_dir.expanduser().resolve()

    if not input_path.exists():
        raise FileNotFoundError(input_path)

    prepare_output(input_path, output_path)

    wb = load_workbook(output_path)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"Foglio '{args.sheet}' non trovato. Disponibili: {wb.sheetnames}")

    ws = wb[args.sheet]
    headers = headers_map(ws)
    require_input_columns(headers)
    headers = ensure_output_columns(ws)
    atomic_save(wb, output_path)

    logging.info("=" * 76)
    logging.info("%s", VERSION)
    logging.info("Input: %s", input_path)
    logging.info("Output: %s", output_path)
    logging.info("Workers: %s", args.workers)
    logging.info("AI concurrency: %s", args.ai_concurrency)
    logging.info("Provider: %s", args.provider)
    logging.info("Gemini model: %s", args.model)
    logging.info("Gemini key: %s", "SI" if os.getenv("GEMINI_API_KEY") else "NO")
    logging.info("Serper key: %s", "SI" if os.getenv("SERPER_API_KEY") else "NO")
    logging.info("Brave key: %s", "SI" if os.getenv("BRAVE_SEARCH_API_KEY") else "NO")
    logging.info("DDGS fallback: %s", "SI" if DDGS is not None else "NO")
    logging.info("Deep mode: %s", "SI" if args.deep else "NO")
    logging.info("CV dir: %s", cv_dir)
    logging.info("Cache dir: %s", cache_dir)
    logging.info("Log: %s", log_file)
    logging.info("=" * 76)

    print(f"Excel output: {output_path}")
    print(f"Log: {log_file.resolve()}")

    people: list[Person] = []
    for row in range(args.start_row, ws.max_row + 1):
        if args.limit is not None and len(people) >= args.limit:
            break

        person = person_from_row(ws, row, headers)
        if not person.pers_id or not person.surname or not person.name:
            continue

        status = clean(getv(ws, row, headers, "Ricerca_Stato")).upper()
        if not args.retry_all and status in {"COMPLETATO", "SPECIALITA_TROVATA"}:
            continue
        if status == "ERRORE" and not args.retry_errors:
            continue

        people.append(person)

    logging.info("Medici da elaborare: %s", len(people))

    if not people:
        atomic_save(wb, output_path)
        return 0

    start_all = time.perf_counter()
    completed = 0
    unsaved = 0
    counts: dict[str, int] = {}

    with ThreadPoolExecutor(max_workers=args.workers, thread_name_prefix="medico") as pool:
        futures = {
            pool.submit(research_person, p, args, cv_dir, cache_dir): p
            for p in people
        }

        for future in as_completed(futures):
            person = futures[future]
            try:
                result = future.result()
            except Exception as exc:
                logging.exception("ERRORE PERSONA | %s | %s", person.pers_id, person.full_name)
                result = {
                    "row": person.row,
                    "status": "ERRORE",
                    "specialty": "",
                    "specialty_confidence": "nessuna",
                    "specialty_evidence": "",
                    "cv_path": "",
                    "cv_url": "",
                    "cv_confidence": "nessuna",
                    "sources": "",
                    "method": "",
                    "notes": "",
                    "updated": utc_now(),
                    "error": f"{type(exc).__name__}: {exc}",
                }

            set_result(ws, headers, person.row, result)

            status = clean(result.get("status")) or "?"
            counts[status] = counts.get(status, 0) + 1
            completed += 1
            unsaved += 1

            if unsaved >= args.save_every:
                atomic_save(wb, output_path)
                unsaved = 0
                logging.info("CHECKPOINT | %s/%s | %s", completed, len(people), output_path)

            if completed % 10 == 0 or completed == len(people):
                elapsed = time.perf_counter() - start_all
                rate = completed / elapsed * 60 if elapsed else 0.0
                eta = (len(people) - completed) / rate if rate else 0.0
                logging.info(
                    "PROGRESS | %s/%s | %.2f medici/min | ETA %.1f min | %s",
                    completed, len(people), rate, eta, counts
                )

    atomic_save(wb, output_path)

    elapsed = time.perf_counter() - start_all
    rate = completed / elapsed * 60 if elapsed else 0.0

    logging.info("=" * 76)
    logging.info("FINE | elaborati=%s | %.1f min | %.2f medici/min", completed, elapsed / 60, rate)
    logging.info("STATI | %s", counts)
    if AI_DISABLED_FOR_RUN:
        logging.info("AI CIRCUIT BREAKER | %s", AI_DISABLE_REASON)
    logging.info("OUTPUT | %s", output_path)
    logging.info("=" * 76)

    print("")
    print(f"Excel salvato in: {output_path}")
    print(f"Log esecuzione: {log_file.resolve()}")
    print(f"CV salvati in: {cv_dir}")
    return 0


if __name__ == "__main__":
    import sys

    early_input, early_output, early_log_dir = parse_early_paths(sys.argv)
    startup_log = bootstrap_log_path(early_log_dir)

    try:
        initial_output = create_initial_output_copy(early_input, early_output)
        if initial_output:
            write_bootstrap_log(startup_log, f"INFO | Excel iniziale: {initial_output}")
            print(f"Excel di output: {initial_output}")
            print(f"Log startup: {startup_log.resolve()}")
    except Exception as exc:
        write_bootstrap_log(
            startup_log,
            f"ERRORE | Creazione output iniziale: {type(exc).__name__}: {exc}"
        )

    ok, missing = import_dependencies(startup_log)
    if not ok:
        print("\nERRORE: mancano dipendenze obbligatorie.")
        print("Installa con: python -m pip install -r requirements_medici_v2.txt")
        print(f"Log diagnostico: {startup_log.resolve()}")
        for item in missing:
            print(f" - {item}")
        raise SystemExit(2)

    try:
        load_dotenv()
        raise SystemExit(main())
    except KeyboardInterrupt:
        write_bootstrap_log(startup_log, "INTERRUZIONE | KeyboardInterrupt")
        print("\nInterrotto.")
        raise SystemExit(130)
    except Exception as exc:
        write_bootstrap_log(startup_log, f"FATAL | {type(exc).__name__}: {exc}")
        print(f"\nERRORE FATALE: {type(exc).__name__}: {exc}")
        print(f"Log diagnostico: {startup_log.resolve()}")
        raise
