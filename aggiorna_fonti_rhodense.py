"""Incrocia un indice pubblico ospedaliero con l'Excel, senza Search API."""
import argparse
import csv
import logging
import os
from pathlib import Path
import tempfile
from urllib.parse import urljoin, urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook
import requests

import scriptMedici as medici


INDEX_URL = 'https://www.asst-rhodense.it/AmministrazioneTrasparente/Personale/Tassi-assenza-2016/CVDirigenti.html'


def parse_index(raw: bytes, base_url: str = INDEX_URL) -> dict[tuple[str, str], list[str]]:
    soup = BeautifulSoup(raw, 'html.parser')
    result = {}
    valid_header = False
    for tr in soup.select('tr'):
        cells = tr.find_all(['th', 'td'], recursive=False)
        values = [medici.normalize(td.get_text(' ', strip=True)) for td in cells]
        if values[:3] == ['cognome', 'nome', 'profilo']:
            valid_header = True
            continue
        if not valid_header or len(values) < 5 or values[2] != 'medici':
            continue
        key = (values[0], values[1])
        for anchor in cells[-1].find_all('a', href=True):
            url = urljoin(base_url, anchor['href'])
            parsed = urlparse(url)
            if parsed.hostname == urlparse(base_url).hostname and parsed.path.lower().endswith('.pdf'):
                if url not in result.setdefault(key, []):
                    result[key].append(url)
    if not valid_header or not result:
        raise ValueError('Indice non riconosciuto o privo di medici: nessuna modifica al CSV.')
    return result


def match_people(input_path: Path, sheet: str, index: dict) -> list[dict[str, str]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    matches = []
    try:
        rows = workbook[sheet].iter_rows(values_only=True)
        positions = {medici.normalize(label): i for i, label in enumerate(next(rows)) if label}
        medici.require_input_columns(positions)
        for row in rows:
            pid = medici.numericish(medici.clean(row[positions['pers_id']]))
            surname = medici.normalize(row[positions['pers_cognome']])
            name = medici.normalize(row[positions['pers_nome']])
            if not pid or not surname or not name:
                continue
            for url in index.get((surname, name), []):
                matches.append({'Pers_Id': pid, 'URL': url,
                    'Nota': f'Candidato da indice ufficiale ASST Rhodense; nome e cognome esatti ({surname} {name}); identita e titoli da verificare nel CV. Indice: {INDEX_URL}'})
    finally:
        workbook.close()
    return matches


def merge_manifest(path: Path, candidates: list[dict], max_new: int) -> int:
    fields = ['Pers_Id', 'URL', 'Nota']
    existing = []
    if path.exists():
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream)
            if not {'Pers_Id', 'URL'}.issubset(reader.fieldnames or []):
                raise ValueError('CSV esistente non valido: nessuna modifica.')
            fields = list(reader.fieldnames)
            if 'Nota' not in fields:
                fields.append('Nota')
            existing = list(reader)
    seen = {(medici.numericish(medici.clean(row['Pers_Id'])), medici.clean(row['URL'])) for row in existing}
    additions = []
    for row in candidates:
        key = (row['Pers_Id'], row['URL'])
        if key in seen:
            continue
        if len(additions) >= max_new:
            break
        seen.add(key)
        additions.append(row)
    if not additions:
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, filename = tempfile.mkstemp(prefix='.fonti_', suffix='.csv', dir=path.parent)
    os.close(fd)
    tmp = Path(filename)
    try:
        with tmp.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
            writer.writeheader()
            writer.writerows(existing + additions)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
    return len(additions)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--sheet', default='Foglio1')
    parser.add_argument('--output', type=Path, default=Path('fonti_campione.csv'))
    parser.add_argument('--max-new', type=int, default=40)
    args = parser.parse_args()
    if args.max_new < 1:
        parser.error('--max-new deve essere positivo')
    medici.requests = requests
    got = medici.get_bytes(INDEX_URL, medici.MAX_HTML_BYTES)
    if got is None:
        raise RuntimeError('Indice non scaricabile: il CSV esistente non e stato modificato.')
    response, raw = got
    try:
        index = parse_index(raw, response.url)
    finally:
        response.close()
    matches = match_people(args.input, args.sheet, index)
    added = merge_manifest(args.output, matches, args.max_new)
    print(f'Nominativi medici nell’indice: {len(index)}')
    print(f'Abbinamenti candidati nell’Excel: {len(matches)}')
    print(f'Nuove fonti aggiunte: {added} | CSV: {args.output}')
    print('Nessuna Search API. Le corrispondenze di nome non costituiscono verifica di identita.')
    return 0


if __name__ == '__main__':
    logging.basicConfig(level=logging.INFO)
    raise SystemExit(main())
