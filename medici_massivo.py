"""Ricerca per fonti su tutto l'archivio, con coda SQLite e ripresa persistente.

Nessuna Search API: gli indici degli enti vengono incrociati localmente.
L'identità nominale rimane distinta dalla verifica con data di nascita/CF.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import re
import shutil
import sqlite3
import tempfile
import threading
import time
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict, deque
from concurrent.futures import ThreadPoolExecutor, as_completed, wait, FIRST_COMPLETED
from dataclasses import asdict
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse, urldefrag

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from pypdf import PdfReader

import scriptMedici as core

SCHEMA = 1
ANALYZER = 5
DISCOVERY_VERSION = 3
UA = 'MediciResearch/1.0'
MAX_BYTES = 25 * 1024 * 1024
EXTRA_COLUMNS = [
    'Massivo_Stato', 'Identita_Verifica', 'Specialita_Documentata',
    'Specialita_Proposta', 'Specialita_Evidenza_Massivo', 'Tipo_Fonte',
    'CV_Stato_Massivo', 'CV_File_Massivo', 'Fonti_Massivo', 'Motivo_Massivo',
    'Disciplina_Dichiarata_Profilo', 'Disciplina_Evidenza', 'Fonti_Ancora_In_Coda', 'Fonti_Con_Errore',
]


def key(value):
    text = unicodedata.normalize('NFKD', core.clean(value).casefold())
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return ' '.join(re.findall(r'[a-z0-9]+', text))


def sha_file(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def atomic_text(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, prefix='.' + path.name)
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            stream.write(value)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def readable_cv_path(source, person, cv_dir, url=''):
    """Copia un CV verificato fuori dalla cache usando Pers_Id, cognome e nome."""
    source = Path(source)
    cv_dir = Path(cv_dir).resolve()
    cv_dir.mkdir(parents=True, exist_ok=True)
    ext = source.suffix.lower() if source.suffix.lower() in {'.pdf', '.doc', '.docx', '.html'} else '.bin'
    base = '_'.join((core.safe_part(person.pers_id), core.safe_part(person.surname), core.safe_part(person.name)))
    target = cv_dir / f'{base}{ext}'
    if target.exists():
        if sha_file(target) == sha_file(source):
            return target
        digest = hashlib.sha1((url or str(source)).encode('utf-8')).hexdigest()[:10]
        target = cv_dir / f'{base}_{digest}{ext}'
    if not target.exists() or sha_file(target) != sha_file(source):
        shutil.copy2(source, target)
    return target


def materialize_verified_cvs(store, cv_dir):
    """Rende leggibili anche i CV acquisiti nei cicli massivi precedenti."""
    people = {row['pid']: core.Person(**json.loads(row['data'])) for row in store.db.execute('SELECT pid,data FROM people')}
    changed = 0
    for row in store.db.execute('SELECT pid,url,data FROM evidence').fetchall():
        data = json.loads(row['data'])
        if not data.get('cv') or data.get('identity') not in ('anagrafica_concordante', 'solo_nome_completo'):
            continue
        source = Path(data.get('path', ''))
        person = people.get(row['pid'])
        if not person or not source.is_file():
            continue
        target = readable_cv_path(source, person, cv_dir, data.get('url', row['url']))
        if str(target) != data.get('path'):
            data['path'] = str(target.resolve())
            store.db.execute('UPDATE evidence SET data=? WHERE pid=? AND url=?',
                             (json.dumps(data, ensure_ascii=False), row['pid'], row['url']))
            changed += 1
    if changed:
        store.db.commit()
        logging.info('CV rinominati e copiati in %s: %s', Path(cv_dir).resolve(), changed)
    return changed


class RunLock:
    """Lock di processo rilasciato dal sistema anche dopo un crash, Windows/macOS/Linux."""
    def __init__(self, folder):
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        self.stream = (folder / '.run.lock').open('a+b')
        self.stream.seek(0)
        if not self.stream.read(1):
            self.stream.write(b'0')
            self.stream.flush()
        self.stream.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(self.stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.stream.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            self.stream.close()
            raise ValueError('Un’altra esecuzione usa questa --state-dir. Attendere che termini.') from exc

    def close(self):
        if os.name == 'nt':
            import msvcrt
            self.stream.seek(0)
            msvcrt.locking(self.stream.fileno(), msvcrt.LK_UNLCK, 1)
        self.stream.close()


class Store:
    def __init__(self, folder):
        self.folder = Path(folder).resolve()
        self.folder.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(self.folder / 'ricerca.sqlite', timeout=30)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.executescript('''
          CREATE TABLE IF NOT EXISTS meta (name TEXT PRIMARY KEY, value TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS people (pid TEXT PRIMARY KEY, input_row INTEGER, data TEXT NOT NULL, name_key TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS urls (url TEXT PRIMARY KEY, kind TEXT NOT NULL, state TEXT NOT NULL DEFAULT 'pending', error TEXT NOT NULL DEFAULT '', path TEXT, final_url TEXT, sha TEXT, text TEXT, analyzed INTEGER DEFAULT 0);
          CREATE TABLE IF NOT EXISTS candidates (pid TEXT, url TEXT, source TEXT, PRIMARY KEY(pid,url));
          CREATE TABLE IF NOT EXISTS probes (url TEXT PRIMARY KEY, source TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS evidence (pid TEXT, url TEXT, data TEXT, PRIMARY KEY(pid,url));
          CREATE TABLE IF NOT EXISTS sources (id TEXT PRIMARY KEY, state TEXT, count INTEGER, error TEXT, updated TEXT);
          CREATE INDEX IF NOT EXISTS candidate_url ON candidates(url);
          CREATE INDEX IF NOT EXISTS url_state ON urls(state);
        ''')
        old = self.get('schema')
        if old and old != str(SCHEMA):
            raise ValueError('Schema archivio non compatibile: usare una nuova --state-dir.')
        self.put('schema', str(SCHEMA))

    def get(self, name):
        row = self.db.execute('SELECT value FROM meta WHERE name=?', (name,)).fetchone()
        return row[0] if row else None

    def put(self, name, value):
        self.db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)', (name, value))
        self.db.commit()

    def import_people(self, path, sheet):
        fingerprint = sha_file(path) + ':' + sheet
        previous = self.get('input')
        if previous and previous != fingerprint:
            raise ValueError('Input diverso da quello nella coda. Usare una nuova --state-dir per non mescolare gli archivi.')
        if self.get('import_complete') == fingerprint:
            return
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook[sheet].iter_rows(values_only=True)
            labels = next(rows)
            positions = {core.clean(v).casefold(): i for i, v in enumerate(labels) if v}
            core.require_input_columns(positions)
            self.db.execute('DELETE FROM people')
            for n, row in enumerate(rows, 2):
                def val(name):
                    index = positions.get(name.casefold())
                    return row[index] if index is not None and index < len(row) else ''
                dob = val('Pers_DataNascita')
                person = core.Person(n, core.numericish(core.clean(val('Pers_Id'))),
                    core.numericish(core.clean(val('Medico_Id'))), core.clean(val('Pers_Cognome')),
                    core.clean(val('Pers_Nome')), dob.strftime('%d/%m/%Y') if isinstance(dob, (datetime, date)) else core.clean(dob),
                    core.clean(val('Pers_CodFis')).upper(), core.clean(val('Indirizzi_Citta')),
                    core.clean(val('emailPredefinita')) or core.clean(val('Email')))
                if not person.pers_id:
                    raise ValueError(f'Pers_Id mancante alla riga {n}: importazione annullata.')
                self.db.execute('INSERT INTO people VALUES (?,?,?,?)', (person.pers_id, n, json.dumps(asdict(person), ensure_ascii=False), key(person.full_name)))
                if n % 20000 == 0:
                    logging.info('Importazione: %s righe', n - 1)
            self.db.commit()
            self.put('input', fingerprint)
            self.put('import_complete', fingerprint)
        except BaseException:
            self.db.rollback()
            raise
        finally:
            workbook.close()

    def add(self, pid, url, source, kind='profile'):
        url = urldefrag(url)[0]
        if not urlparse(url).scheme in ('http', 'https'):
            return
        self.db.execute('INSERT OR IGNORE INTO urls(url,kind) VALUES (?,?)', (url, kind))
        self.db.execute('INSERT OR IGNORE INTO candidates VALUES (?,?,?)', (pid, url, source))

    def add_probe(self, url, source):
        """Accoda una scheda con URL numerico; il nome verrà letto dall'intestazione."""
        url = urldefrag(url)[0]
        if urlparse(url).scheme not in ('http', 'https'):
            return
        self.db.execute("INSERT OR IGNORE INTO urls(url,kind) VALUES (?,'profile')", (url,))
        self.db.execute('INSERT OR REPLACE INTO probes VALUES (?,?)', (url, source))

    def close(self):
        self.db.close()


class Names:
    """Indice di sequenze esatte: nessuna scansione 81.000 x URL e niente fuzzy automatico."""
    def __init__(self, people):
        self.people = people
        self.names = defaultdict(set)
        self.full = defaultdict(set)
        self.compact = defaultdict(set)
        for pid, person in people.items():
            for label in (person.full_name, f'{person.surname} {person.name}'):
                self.names[key(label)].add(pid)
                self.compact[key(label).replace(' ', '')].add(pid)
            self.full[key(person.full_name)].add(pid)
        self.lengths = sorted({len(k.split()) for k in self.names}, reverse=True)

    def match(self, text):
        words = key(unquote(text)).split()
        hits = []
        occupied = set()
        for length in self.lengths:
            for start in range(len(words) - length + 1):
                if any(i in occupied for i in range(start, start + length)):
                    continue
                ids = self.names.get(' '.join(words[start:start + length]))
                if ids:
                    hits.extend(ids)
                    occupied.update(range(start, start + length))
        if not hits:
            for word in words:
                hits.extend(self.compact.get(word, set()))
        return set(hits)

    def ambiguous(self, person):
        return len(self.full[key(person.full_name)]) > 1


class FetchError(Exception):
    pass


class RobotPolicy:
    """Regole robots con gruppi ripetuti, wildcard e precedenza al percorso più lungo."""
    def parse(self, lines):
        self.groups = []
        agents, rules = [], []
        delay = None
        def finish():
            if agents:
                self.groups.append((list(agents), list(rules), delay))
        for line in list(lines) + ['User-agent: __end__']:
            line = line.split('#', 1)[0].strip()
            if ':' not in line:
                continue
            field, value = [part.strip() for part in line.split(':', 1)]
            field = field.casefold()
            if field == 'user-agent':
                if rules or delay is not None:
                    finish()
                    agents, rules, delay = [], [], None
                agents.append(value.casefold())
            elif field in ('allow', 'disallow') and agents and value:
                rules.append((field == 'allow', value))
            elif field == 'crawl-delay' and agents:
                try:
                    delay = max(0, float(value))
                except ValueError:
                    pass
        finish()

    def selected(self, agent):
        matches = []
        for agents, rules, delay in self.groups:
            scores = [len(a) if a != '*' else 0 for a in agents if a == '*' or a in agent.casefold()]
            if scores:
                matches.append((max(scores), rules, delay))
        best = max((m[0] for m in matches), default=-1)
        return [(rules, delay) for score, rules, delay in matches if score == best]

    def can_fetch(self, agent, url):
        parsed = urlparse(url)
        path = unquote(parsed.path + ('?' + parsed.query if parsed.query else ''))
        matches = []
        for rules, _ in self.selected(agent):
            for allow, pattern in rules:
                end = pattern.endswith('$')
                literal = unquote(pattern[:-1] if end else pattern)
                regex = '^' + '.*'.join(re.escape(part) for part in literal.split('*')) + ('$' if end else '')
                if re.search(regex, path):
                    matches.append((len(literal.replace('*', '')), allow))
        return max(matches)[1] if matches else True

    def crawl_delay(self, agent):
        return max((delay or 0 for _, delay in self.selected(agent)), default=0)


class Fetcher:
    """Cache URL con hash, robots, un trasferimento per host e limite globale."""
    def __init__(self, folder, limit=6000, delay=0.5):
        self.folder = Path(folder) / 'http'
        self.folder.mkdir(parents=True, exist_ok=True)
        self.limit, self.delay, self.requests = limit, delay, 0
        self.lock = threading.Lock()
        self.host_locks = defaultdict(threading.Lock)
        self.last = {}
        self.robots = {}
        self.blocked = {}
        self.local = threading.local()

    def session(self):
        if not hasattr(self.local, 'session'):
            self.local.session = requests.Session()
            self.local.session.headers.update({'User-Agent': UA, 'Accept-Language': 'it-IT,it;q=0.9'})
        return self.local.session

    def _request(self, url):
        if not core.is_public_http_url(url) or urlparse(url).username:
            raise FetchError('URL_NON_PUBBLICO')
        host = urlparse(url).netloc
        with self.host_locks[host]:
            if host in self.blocked:
                raise FetchError(self.blocked[host])
            with self.lock:
                if self.requests >= self.limit:
                    raise FetchError('LIMITE_RICHIESTE')
                self.requests += 1
            pause = self.delay - (time.monotonic() - self.last.get(host, 0))
            if pause > 0:
                time.sleep(pause)
            self.last[host] = time.monotonic()
            try:
                with self.session().get(url, timeout=(8, 18), stream=True, allow_redirects=False) as response:
                    if response.status_code in (401, 403, 429):
                        self.blocked[host] = f'HTTP_{response.status_code}: host sospeso per questa esecuzione'
                    data = bytearray()
                    for chunk in response.iter_content(65536):
                        data.extend(chunk)
                        if len(data) > MAX_BYTES:
                            raise FetchError('DOCUMENTO_TROPPO_GRANDE')
                    return response.status_code, dict(response.headers), bytes(data)
            except requests.RequestException as exc:
                raise FetchError(type(exc).__name__) from exc

    def allowed(self, url):
        parsed = urlparse(url)
        origin = f'{parsed.scheme}://{parsed.netloc}'
        # Un solo caricamento robots per origine, anche con worker simultanei.
        with self.lock:
            robot_lock = self.host_locks['robots:' + origin]
        with robot_lock:
            if origin not in self.robots:
                robots_url = origin + '/robots.txt'
                status, headers, raw = self._request(robots_url)
                if status in (301, 302, 303, 307, 308):
                    target = urljoin(robots_url, headers.get('Location', ''))
                    if urlparse(target).netloc != parsed.netloc:
                        raise FetchError('REDIRECT_ROBOTS_DA_VERIFICARE')
                    status, headers, raw = self._request(target)
                parser = RobotPolicy()
                if status == 200:
                    parser.parse(raw.decode('utf-8', 'replace').splitlines())
                elif status in (404, 410):
                    parser.parse(['User-agent: *', 'Allow: /'])
                else:
                    raise FetchError(f'ROBOTS_NON_DISPONIBILE_HTTP_{status}')
                self.robots[origin] = parser
            parser = self.robots[origin]
            crawl_delay = parser.crawl_delay(UA) or parser.crawl_delay('*')
            if crawl_delay:
                self.delay = max(self.delay, crawl_delay)
            return parser.can_fetch(UA, url)

    def get(self, url, refresh=False):
        digest = hashlib.sha256(url.encode()).hexdigest()
        meta = self.folder / (digest + '.json')
        if not refresh:
            try:
                saved = json.loads(meta.read_text())
                path = self.folder / saved['file']
                if path.parent == self.folder and sha_file(path) == saved['sha']:
                    return saved, path.read_bytes()
            except (OSError, ValueError, KeyError):
                pass
        current = url
        for hop in range(6):
            if not self.allowed(current):
                raise FetchError('ROBOTS_NON_CONSENTE')
            status, headers, raw = self._request(current)
            if status in (301, 302, 303, 307, 308):
                if not headers.get('Location') or hop == 5:
                    raise FetchError('TROPPI_REDIRECT')
                current = urljoin(current, headers['Location'])
                continue
            if status != 200:
                raise FetchError(f'HTTP_{status}')
            ctype = next((v for k, v in headers.items() if k.lower() == 'content-type'), '')
            ext = '.pdf' if raw.startswith(b'%PDF') else ('.docx' if raw.startswith(b'PK') else '.doc' if raw.startswith(bytes.fromhex('d0cf11e0')) else '.html')
            if current.lower().endswith('.xml') or 'xml' in ctype:
                ext = '.xml'
            path = self.folder / (digest + ext)
            path.write_bytes(raw)
            saved = {'url': url, 'final_url': current, 'sha': hashlib.sha256(raw).hexdigest(),
                     'file': path.name, 'type': ctype, 'retrieved': core.utc_now()}
            atomic_text(meta, json.dumps(saved))
            return saved, raw
        raise FetchError('REDIRECT_NON_RISOLTO')


def sitemap_entries(raw):
    root = ET.fromstring(raw)
    # Solo <loc> del protocollo sitemap, non image:loc o video:loc.
    ns = '{http://www.sitemaps.org/schemas/sitemap/0.9}'
    if root.tag not in (ns + 'urlset', ns + 'sitemapindex', 'urlset', 'sitemapindex'):
        raise ValueError('Formato sitemap non riconosciuto')
    tag = ns + 'loc' if root.tag.startswith(ns) else 'loc'
    urls = [node.text.strip() for entry in root for node in entry if node.tag == tag and node.text]
    return root.tag.endswith('sitemapindex'), urls


def discover_one(spec, fetcher, names, refresh=False):
    found = set()
    queue = list(spec['urls'])
    seen = set()
    errors = []
    while queue and len(seen) < spec.get('max_index_pages', 40):
        url = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        try:
            info, raw = fetcher.get(url, refresh=refresh)
            if spec['type'] == 'sitemap':
                is_index, links = sitemap_entries(raw)
                if is_index:
                    for link in links:
                        if link.startswith('http://') and url.startswith('https://'):
                            link = 'https://' + link[7:]
                        if urlparse(link).hostname in spec['hosts'] and re.search(spec.get('sitemap_pattern', '.*'), link):
                            queue.append(link)
                    continue
                for link in links:
                    if link.startswith('http://') and url.startswith('https://'):
                        link = 'https://' + link[7:]
                    if urlparse(link).hostname not in spec['hosts'] or not re.search(spec['profile_pattern'], link):
                        continue
                    if spec.get('content_match'):
                        found.add(('', link, 'probe'))
                        continue
                    path = unquote(urlparse(link).path).rstrip('/')
                    if spec.get('name_pattern'):
                        match = re.search(spec['name_pattern'], path)
                        if not match:
                            continue
                        label = match[1]
                    else:
                        label = path.rsplit('/', 1)[-1]
                    for pid in names.match(label):
                        found.add((pid, link, 'indexed_activity' if spec.get('indexed_activity') else 'profile'))
            else:
                soup = BeautifulSoup(raw, 'html.parser')
                for a in soup.select('a[href]'):
                    link = urljoin(info['final_url'], a['href'])
                    if urlparse(link).hostname not in spec['hosts']:
                        continue
                    if not re.search(spec.get('link_pattern', '.*'), link):
                        continue
                    label = a.get_text(' ', strip=True)
                    row = a.find_parent('tr')
                    if row:
                        label = row.get_text(' ', strip=True)
                    if len(label) > 1200:
                        continue
                    ids = names.match(label) or names.match(unquote(urlparse(link).path))
                    # Una riga con più persone differenti non assegna il CV a tutte.
                    if len({key(names.people[pid].full_name) for pid in ids}) > 1:
                        continue
                    for pid in ids:
                        found.add((pid, link, 'document' if spec.get('documents') else 'profile'))
        except Exception as exc:
            errors.append(f'{url}: {type(exc).__name__}: {exc}')
    if queue:
        errors.append('Limite pagine indice raggiunto; aumentare max_index_pages nel catalogo.')
    return found, errors


def indexed_activity_result(spec, person, url, ambiguous):
    """Crea evidenza nominale dalla categoria professionale esplicita nell'URL indicizzato."""
    path = unquote(urlparse(url).path).rstrip('/')
    match = re.search(spec.get('activity_pattern', r'^/([^/]+)/'), path)
    slug = match[1].casefold() if match else ''
    label = spec.get('activity_map', {}).get(slug, '')
    if not label:
        raw = slug.replace('-', ' ')
        aliases = {key(k): v for k, v in core.SPECIALTY_ALIASES.items()}
        label = aliases.get(key(raw), '')
    identity = 'omonimia' if ambiguous else 'solo_nome_completo'
    reason = ('Più persone nell’input con lo stesso nome: manca un discriminante anagrafico'
              if ambiguous else 'Nome completo concordante con il profilo indicizzato; identità anagrafica non confermata dalla fonte')
    return {
        'identity': identity,
        'specialties': [],
        'evidence': [],
        'activities': [label] if label and not ambiguous else [],
        'activity_evidence': [f"Categoria pubblica {spec['id']}: {label}"] if label and not ambiguous else [],
        'cv': False,
        'reason': reason if label else reason + '; categoria non inclusa nella tassonomia medica',
        'path': '',
        'kind': 'Profilo pubblico indicizzato',
        'url': url,
    }


def discover(store, fetcher, names, catalog, refresh=False, workers=6):
    specs = json.loads(Path(catalog).read_text(encoding='utf-8'))['sources']
    def signature(spec):
        return hashlib.sha256((str(DISCOVERY_VERSION) + json.dumps(spec, sort_keys=True)).encode()).hexdigest()
    todo = [s for s in specs if s.get('enabled', True) and (refresh or store.get('source:' + s['id']) != signature(s) or not store.db.execute("SELECT 1 FROM sources WHERE id=? AND state='done'", (s['id'],)).fetchone())]
    with ThreadPoolExecutor(max_workers=workers) as pool:
        jobs = {pool.submit(discover_one, spec, fetcher, names, refresh): spec for spec in todo}
        for job in as_completed(jobs):
            spec = jobs[job]
            found, errors = job.result()
            if not errors:
                fresh = {(pid, url) for pid, url, _ in found}
                stale = [(row['pid'], row['url']) for row in store.db.execute('SELECT pid,url FROM candidates WHERE source=?', (spec['id'],)) if (row['pid'], row['url']) not in fresh]
                store.db.executemany('DELETE FROM candidates WHERE pid=? AND url=?', stale)
                store.db.execute('DELETE FROM evidence WHERE NOT EXISTS (SELECT 1 FROM candidates c WHERE c.pid=evidence.pid AND c.url=evidence.url)')
                fresh_probes = {url for pid, url, kind in found if not pid and kind == 'probe'}
                if spec.get('content_match'):
                    stale_probes = [(row['url'],) for row in store.db.execute('SELECT url FROM probes WHERE source=?', (spec['id'],)) if row['url'] not in fresh_probes]
                    store.db.executemany('DELETE FROM probes WHERE url=?', stale_probes)
                store.db.execute('DELETE FROM urls WHERE NOT EXISTS (SELECT 1 FROM candidates c WHERE c.url=urls.url) AND NOT EXISTS (SELECT 1 FROM probes p WHERE p.url=urls.url)')
            for pid, url, kind in sorted(found, key=lambda item: (item[1], item[0])):
                if kind == 'probe':
                    store.add_probe(url, spec['id'])
                else:
                    store.add(pid, url, spec['id'], kind)
                    if kind == 'indexed_activity':
                        person = names.people[pid]
                        result = indexed_activity_result(spec, person, url, names.ambiguous(person))
                        store.db.execute(
                            "UPDATE urls SET state='done',error='',path='',final_url=?,sha=?,text='',analyzed=? WHERE url=?",
                            (url, hashlib.sha256(url.encode()).hexdigest(), ANALYZER, url),
                        )
                        store.db.execute(
                            'INSERT OR REPLACE INTO evidence VALUES (?,?,?)',
                            (pid, url, json.dumps(result, ensure_ascii=False)),
                        )
            store.db.execute('INSERT OR REPLACE INTO sources VALUES (?,?,?,?,?)',
                (spec['id'], 'partial' if errors else 'done', len(found), '\n'.join(errors), core.utc_now()))
            store.db.commit()
            store.put('source:' + spec['id'], signature(spec))
            logging.info('FONTE %s: %s associazioni, %s problemi', spec['id'], len(found), len(errors))


def import_manifest(store, path, people):
    if not path or not Path(path).exists():
        return
    with Path(path).open(encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            pid = core.numericish(core.clean(row.get('Pers_Id')))
            if pid in people and row.get('URL'):
                store.add(pid, row['URL'], 'manifest', 'document')
    store.db.commit()


def import_local_documents(store, names, folders):
    index = core.local_document_index(*folders)
    for pid, person in names.people.items():
        for path in core.local_documents(person, index):
            if path.stat().st_size > MAX_BYTES:
                continue
            digest = sha_file(path)
            url = 'local:' + digest
            store.db.execute("INSERT OR IGNORE INTO urls(url,kind,path,sha) VALUES (?,'local',?,?)", (url, str(path.resolve()), digest))
            store.db.execute('INSERT OR IGNORE INTO candidates VALUES (?,?,?)', (pid, url, 'cv_locale'))
    store.db.commit()


def visible_profile(raw):
    soup = BeautifulSoup(raw, 'html.parser')
    seo_title = soup.title.get_text(' ', strip=True) if soup.title else ''
    title = ' '.join(x.get_text(' ', strip=True) for x in soup.select('h1'))
    if not title:
        title = seo_title
    roles = [a.get_text(' ', strip=True) for a in soup.select('[data-test-id=doctor-specializations] a[title]')]
    # Alcune directory espongono la disciplina soltanto nel titolo SEO.
    match = re.search(r'(?i)specialista\s+in\s+(.+?)(?=\s+(?:a|in)\s+[^|]+(?:\||$))', seo_title)
    if match:
        roles.append(match.group(1).strip())
    for item in soup.select('script,style,noscript,nav,footer,header,aside,form,#profile-reviews,[itemprop=review]'):
        item.decompose()
    main = soup.select_one('main') or soup.select_one('article') or soup.body or soup
    text = main.get_text('\n', strip=True)
    if roles:
        text = '\n'.join('Disciplina dichiarata: ' + role for role in roles) + '\n' + text
    return title, text, main


def analyze_content(person, text, is_cv, title, ambiguous):
    """Qualifiche esplicite e motivazione dell'identità, separati dalla presenza del CV."""
    result = {'identity': 'non_verificata', 'specialties': [], 'evidence': [], 'activities': [], 'activity_evidence': [], 'cv': is_cv, 'reason': ''}
    region = (title + '\n' + core.cv_header_region(text)) if is_cv else title
    full, reverse = key(person.full_name), key(f'{person.surname} {person.name}')
    header = ' ' + key(region) + ' '
    if not any(' ' + value + ' ' in header for value in (full, reverse)):
        result['reason'] = 'Nome completo assente dall’intestazione'
        return result
    found, wanted = core.cv_birth_date(text), core.parse_birth_date(person.birth_date)
    codes = re.findall(r'\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b', text.upper())
    if (found is not None and wanted is not None and found != wanted) or (codes and person.fiscal_code and person.fiscal_code not in codes):
        result['identity'] = 'incompatibile'
        result['reason'] = 'Data di nascita o codice fiscale incompatibile'
        return result
    strong = bool((found is not None and found == wanted) or (person.fiscal_code and person.fiscal_code in codes))
    if strong:
        result['identity'] = 'anagrafica_concordante'
    elif ambiguous:
        result['identity'] = 'omonimia'
        result['reason'] = 'Più persone nell’input con lo stesso nome: manca un discriminante anagrafico'
        return result
    else:
        result['identity'] = 'solo_nome_completo'
        result['reason'] = 'Nome completo concordante; identità anagrafica non confermata dalla fonte'
    if core.target_role_conflict(region, person):
        result['identity'] = 'incompatibile'
        result['reason'] = 'Ruolo incompatibile nell’intestazione'
        return result
    # Il testo HTML resta diviso in righe: evita di attribuire qualifiche di medici correlati.
    cleaned = re.split(r'(?im)^\s*(?:altri medici|medici correlati|potrebbero interessarti|i nostri medici|stessa specialit[aà]|pubblicazioni|bibliografia)\s*$', text)[0]
    cleaned = re.sub(r"([^\n])\n(?=(?:e |dell[’']|della |del ))", r'\1 ', cleaned, flags=re.I)
    candidates = core.verified_cv_specialty_candidates(cleaned, person, '')
    allowed = {key(value) for value in core.SPECIALTY_ALIASES.values()}
    allowed.update({key("Scienza dell'Alimentazione"), key('Malattie dell’Apparato Cardiovascolare')})
    unknown = [item[1] for item in candidates if key(item[1]) not in allowed]
    candidates = [item for item in candidates if key(item[1]) in allowed]
    if unknown:
        result['reason'] += ' Titoli fuori dalla tassonomia, da normalizzare: ' + '; '.join(unknown)
    result['specialties'] = sorted({item[1] for item in candidates})
    result['evidence'] = list(dict.fromkeys(item[3] for item in candidates))
    aliases = {key(label): value for label, value in core.SPECIALTY_ALIASES.items()}
    aliases.update({key(value): value for value in core.SPECIALTY_ALIASES.values()})
    aliases['medico di medicina generale'] = 'Medicina generale (attività dichiarata)'
    if not is_cv:
        for match in re.finditer(r'(?im)^(?:Disciplina dichiarata: *|(?:Specializzazion[ei]|Area Medica) *:?\s*\n)([^\n]{3,90})', cleaned):
            label = match[1].strip()
            if key(label) in aliases:
                result['activities'].append(aliases[key(label)])
                result['activity_evidence'].append(match[0])
    return result


def extract(raw, info):
    if raw.startswith(b'%PDF'):
        reader = PdfReader(io.BytesIO(raw))
        text = '\n'.join(p.extract_text() or '' for p in reader.pages[:100])[:120000]
        is_cv = core.looks_like_cv(text)[0]
        return text, is_cv, '', None
    if raw.startswith(bytes.fromhex('d0cf11e0')):
        text, method = core.legacy_doc_text_with_method(raw)
        if method == 'ole/binary-heuristic':
            return '', False, '', None
        return text, core.looks_like_cv(text)[0], '', None
    if raw.startswith(b'PK'):
        text = core.docx_text(raw)
        return text, core.looks_like_cv(text)[0], '', None
    title, text, main = visible_profile(raw)
    is_cv = any('curriculum' in key(h.get_text(' ', strip=True)) for h in main.select('h1,h2,h3'))
    return text[:120000], is_cv, title, main


def process_url(url, fetcher, local_path=None):
    try:
        if local_path:
            path = Path(local_path)
            if path.stat().st_size > MAX_BYTES:
                raise FetchError('DOCUMENTO_TROPPO_GRANDE')
            raw = path.read_bytes()
            if 'local:' + hashlib.sha256(raw).hexdigest() != url:
                raise FetchError('DOCUMENTO_LOCALE_MODIFICATO')
            info = {'file': str(path), 'sha': hashlib.sha256(raw).hexdigest(), 'final_url': str(path)}
        else:
            info, raw = fetcher.get(url)
        text, is_cv, title, main = extract(raw, info)
        links = []
        if main:
            for a in main.select('a[href]'):
                link = urljoin(info['final_url'], a['href'])
                if core.looks_like_cv_url(link, a.get_text(' ', strip=True)):
                    if urlparse(link).hostname == urlparse(info['final_url']).hostname:
                        links.append(link)
        return {'info': info, 'text': text, 'cv': is_cv, 'title': title, 'links': list(dict.fromkeys(links))[:5]}
    except Exception as exc:
        return {'error': f'{type(exc).__name__}: {exc}'}


def save_outcome(store, fetcher, names, url, outcome):
    added = []
    error = outcome.get('error')
    if error:
        state = 'pending' if 'LIMITE_RICHIESTE' in error else 'error'
        store.db.execute('UPDATE urls SET state=?,error=? WHERE url=?', (state, error, url))
    else:
        info = outcome['info']
        path = fetcher.folder / info['file']
        store.db.execute("UPDATE urls SET state='done',error='',path=?,final_url=?,sha=?,text=?,analyzed=? WHERE url=?", (str(path), info['final_url'], info['sha'], outcome['text'], ANALYZER, url))
        ids = [r[0] for r in store.db.execute('SELECT pid FROM candidates WHERE url=?', (url,))]
        if not ids:
            probe = store.db.execute('SELECT source FROM probes WHERE url=?', (url,)).fetchone()
            if probe:
                ids = sorted(names.match(outcome['title']))
                for pid in ids:
                    store.db.execute('INSERT OR IGNORE INTO candidates VALUES (?,?,?)', (pid, url, probe['source']))
        for pid in ids:
            person = names.people[pid]
            result = analyze_content(person, outcome['text'], outcome['cv'], outcome['title'], names.ambiguous(person))
            if not outcome['text'].strip():
                result['reason'] = 'Documento senza testo: richiede OCR o conversione'
            result.update({'path': str(path), 'kind': 'CV' if outcome['cv'] else 'Profilo pubblico', 'url': info['final_url']})
            store.db.execute('INSERT OR REPLACE INTO evidence VALUES (?,?,?)', (pid, url, json.dumps(result, ensure_ascii=False)))
            if result['identity'] in ('anagrafica_concordante', 'solo_nome_completo'):
                for link in outcome['links']:
                    store.add(pid, link, 'cv_da_profilo', 'document')
                    added.append(link)
    store.db.commit()
    return added


def process_queue(store, fetcher, names, workers=6, max_documents=5000, checkpoint=None,
                  checkpoint_every=1000):
    processed = 0
    store.db.execute('DELETE FROM evidence WHERE url IN (SELECT url FROM urls WHERE analyzed!=?)', (ANALYZER,))
    store.db.execute("UPDATE urls SET state='pending' WHERE state='done' AND (analyzed!=? OR EXISTS (SELECT 1 FROM candidates c LEFT JOIN evidence e ON e.pid=c.pid AND e.url=c.url WHERE c.url=urls.url AND e.pid IS NULL))", (ANALYZER,))
    store.db.commit()
    buckets = defaultdict(deque)
    known = set()
    def enqueue(row):
        if row['url'] not in known:
            known.add(row['url'])
            buckets[urlparse(row['url']).hostname or '_local'].append(row)
    # Prima le URL già associate per nome (resa alta), poi le sonde numeriche
    # che richiedono il download per scoprire a chi appartiene la scheda.
    for row in store.db.execute("SELECT url,kind,path FROM urls WHERE state='pending' ORDER BY EXISTS(SELECT 1 FROM candidates c WHERE c.url=urls.url) DESC, rowid"):
        enqueue(row)
    active, busy = {}, set()
    host_order = deque(sorted(buckets))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        while (active or any(buckets.values())) and processed < max_documents:
            # Massimo una richiesta in corso per host; nessuna barriera di fine lotto.
            for _ in range(len(host_order)):
                if len(active) >= workers or processed + len(active) >= max_documents or fetcher.requests >= fetcher.limit:
                    break
                host = host_order.popleft()
                host_order.append(host)
                if host in busy or not buckets[host]:
                    continue
                row = buckets[host].popleft()
                job = pool.submit(process_url, row['url'], fetcher, row['path'] if row['kind'] == 'local' else None)
                active[job] = (row['url'], host)
                busy.add(host)
            if not active:
                break
            finished, _ = wait(active, return_when=FIRST_COMPLETED)
            for job in finished:
                url, host = active.pop(job)
                busy.remove(host)
                outcome = job.result()
                added = save_outcome(store, fetcher, names, url, outcome)
                for link in added:
                    row = store.db.execute("SELECT url,kind,path FROM urls WHERE url=? AND state='pending'", (link,)).fetchone()
                    if row:
                        new_host = urlparse(link).hostname or '_local'
                        if new_host not in buckets:
                            host_order.append(new_host)
                        enqueue(row)
                if urlparse(url).netloc in fetcher.blocked:
                    reason = fetcher.blocked[urlparse(url).netloc]
                    while buckets[host]:
                        paused = buckets[host].popleft()
                        store.db.execute("UPDATE urls SET state='error',error=? WHERE url=?", (reason + '; fonte sospesa, pagina non consultata', paused['url']))
                    store.db.commit()
                processed += 1
                if processed % 50 == 0:
                    logging.info('Documenti elaborati in questa esecuzione: %s; richieste HTTP: %s', processed, fetcher.requests)
                if checkpoint and checkpoint_every and processed % checkpoint_every == 0:
                    checkpoint(processed)
    return processed


def summarize(evidence, pending=False, errors=False):
    good = [e for e in evidence if e['identity'] == 'anagrafica_concordante']
    nominal = [e for e in evidence if e['identity'] == 'solo_nome_completo']
    strong_specs = {s for e in good for s in e['specialties']}
    proposed = {s for e in nominal for s in e['specialties']} - strong_specs
    if strong_specs:
        state = 'SPECIALITA_DOCUMENTATA'
    elif good:
        state = 'IDENTITA_CONFERMATA_SENZA_TITOLO'
    elif proposed:
        state = 'SPECIALITA_CON_IDENTITA_NOMINALE'
    elif any(e.get('activities') for e in good + nominal):
        state = 'DISCIPLINA_DICHIARATA_NEL_PROFILO'
    elif nominal:
        state = 'PROFILO_CON_IDENTITA_NOMINALE'
    elif evidence:
        state = 'EVIDENZA_NON_CONFERMATA'
    elif pending:
        state = 'IN_CODA'
    elif errors:
        state = 'FONTE_NON_ACCESSIBILE'
    else:
        state = 'NON_COPERTO_DALLE_FONTI'
    selected = good + nominal
    cvs = [e for e in selected if e['cv']]
    return [state, '; '.join(sorted({e['identity'] for e in evidence})),
            '; '.join(sorted(strong_specs)), '; '.join(sorted(proposed)),
            '\n'.join(dict.fromkeys(v for e in selected for v in e['evidence'])),
            '; '.join(sorted({e['kind'] for e in selected})),
            ('CV con identità anagrafica confermata' if any(e['identity'] == 'anagrafica_concordante' for e in cvs) else 'CV con identità nominale' if cvs else 'CV non acquisito'),
            '\n'.join(e['path'] for e in cvs), '\n'.join(dict.fromkeys(e['url'] for e in evidence)),
            '\n'.join(dict.fromkeys(e['reason'] for e in evidence if e['reason'])),
            '; '.join(sorted({v for e in selected for v in e.get('activities', [])})),
            '\n'.join(dict.fromkeys(v for e in selected for v in e.get('activity_evidence', [])))]


def validate_output_target(store, input_path, output):
    output = Path(output).resolve()
    if output == Path(input_path).resolve():
        raise ValueError('Il report non può sostituire l’input.')
    marker = output.with_suffix('.xlsx.meta.json')
    if output.exists():
        try:
            owner = json.loads(marker.read_text())
        except (OSError, ValueError):
            raise ValueError(
                f'Output esistente non creato dalla modalità massiva: {output}. '
                'Scegliere un nome nuovo con --output; la run non è stata avviata.'
            )
        if owner.get('input') != store.get('input') or owner.get('state') != str(store.folder):
            raise ValueError(
                f'Output appartenente a un altro archivio o stato: {output}. '
                'Scegliere un nome nuovo con --output; la run non è stata avviata.'
            )
    return output, marker


def export_report(store, input_path, sheet, output, cv_dir=Path('cv_medici')):
    output, marker = validate_output_target(store, input_path, output)
    materialize_verified_cvs(store, cv_dir)
    evidence = defaultdict(list)
    for row in store.db.execute('SELECT pid,data FROM evidence'):
        evidence[row['pid']].append(json.loads(row['data']))
    pending, failed = Counter(), Counter()
    all_sources, failures = defaultdict(list), defaultdict(list)
    for row in store.db.execute('SELECT c.pid,u.state,u.url,u.error FROM candidates c JOIN urls u ON u.url=c.url'):
        all_sources[row['pid']].append(row['url'])
        if row['state'] == 'pending':
            pending[row['pid']] += 1
        elif row['state'] == 'error':
            failed[row['pid']] += 1
            failures[row['pid']].append(row['error'])
    workbook = Workbook(write_only=True)
    ws = workbook.create_sheet(sheet)
    ws.freeze_panes = 'C2'
    original = load_workbook(input_path, read_only=True, data_only=True)
    counts = Counter()
    try:
        rows = original[sheet].iter_rows(values_only=True)
        labels = list(next(rows))
        if any(col in labels for col in EXTRA_COLUMNS):
            raise ValueError('Usare come input l’archivio originale, non il report massivo.')
        pid_index = [core.clean(label).casefold() for label in labels].index('pers_id')
        ws.append(labels + EXTRA_COLUMNS)
        for values in rows:
            pid = core.numericish(core.clean(values[pid_index]))
            extra = summarize(evidence[pid], pid in pending, pid in failed)
            extra[8] = '\n'.join(dict.fromkeys([e['url'] for e in evidence[pid]] + [u for u in all_sources[pid] if not u.startswith('local:')]))
            extra[9] += '\n' + '\n'.join(dict.fromkeys(failures[pid]))
            extra.extend([pending[pid], failed[pid]])
            counts[extra[0]] += 1
            # I valori web devono essere stringhe Excel, mai formule eseguibili.
            cells = []
            for value in list(values) + extra:
                if isinstance(value, str):
                    cell = WriteOnlyCell(ws, value=value[:32767])
                    cell.data_type = 's'
                    cells.append(cell)
                else:
                    cells.append(value)
            ws.append(cells)
    finally:
        original.close()
    from openpyxl.utils import get_column_letter
    ws.auto_filter.ref = f'A1:{get_column_letter(len(labels) + len(EXTRA_COLUMNS))}{sum(counts.values()) + 1}'
    summary = workbook.create_sheet('Copertura', 0)
    summary.freeze_panes = 'A2'
    summary.column_dimensions['A'].width = 47
    summary.column_dimensions['B'].width = 25
    summary.column_dimensions['C'].width = 22
    summary.append(['Indicatore', 'Record', 'Quota sul totale'])
    total = sum(counts.values())
    summary.append(['Totale input', total, 1])
    for status, count in sorted(counts.items()):
        percent = WriteOnlyCell(summary, value=count / total if total else 0)
        percent.number_format = '0.00%'
        summary.append([status, count, percent])
    summary.append(['Aggiornamento UTC', core.utc_now()])
    summary.append(['Identità nominale', 'Il nome coincide, ma mancano CF o nascita concordanti; non equivale a identità verificata.'])
    sources = workbook.create_sheet('Fonti')
    sources.column_dimensions['A'].width = 26
    sources.column_dimensions['B'].width = 18
    sources.column_dimensions['C'].width = 26
    sources.column_dimensions['D'].width = 75
    sources.append(['Fonte', 'Stato indice', 'Associazioni candidate', 'Problemi', 'Aggiornamento'])
    for row in store.db.execute('SELECT * FROM sources ORDER BY id'):
        sources.append(list(row))
    output.parent.mkdir(parents=True, exist_ok=True)
    core.atomic_save(workbook, output)
    atomic_text(marker, json.dumps({'input': store.get('input'), 'state': str(store.folder)}))
    metrics = {'total': total, 'states': dict(counts), 'people_with_sources': store.db.execute('SELECT COUNT(DISTINCT pid) FROM candidates').fetchone()[0],
               'urls': dict(store.db.execute('SELECT state,COUNT(*) FROM urls GROUP BY state').fetchall()), 'updated': core.utc_now()}
    atomic_text(output.with_suffix('.summary.json'), json.dumps(metrics, ensure_ascii=False, indent=2))
    return metrics


def run(args):
    core.import_dependencies(Path(args.log_dir) / 'massivo_bootstrap.log')
    lock = RunLock(args.state_dir)
    store = Store(args.state_dir)
    try:
        logging.info('Importazione e controllo di tutto l’archivio: %s', args.input)
        store.import_people(args.input, args.sheet)
        people = {r['pid']: core.Person(**json.loads(r['data'])) for r in store.db.execute('SELECT pid,data FROM people')}
        names = Names(people)
        logging.info('Archivio: %s persone; %s nomi ripetuti', len(people), sum(len(v) - 1 for v in names.full.values()))
        output = args.output or Path('output/risultati_massivi.xlsx')
        # Il controllo deve avvenire prima di discovery/download: mai sprecare ore
        # per scoprire soltanto all'esportazione che il nome è già occupato.
        validate_output_target(store, args.input, output)
        fetcher = Fetcher(store.folder, args.max_http_requests, args.host_delay)
        if args.retry_errors:
            store.db.execute("UPDATE urls SET state='pending',error='' WHERE state='error'")
            store.db.commit()
        import_manifest(store, args.mass_manifest, people)
        import_local_documents(store, names, (args.cv_dir, args.cv_review_dir))
        interrupted = False
        try:
            if not args.export_only:
                if not args.skip_discovery:
                    discover(store, fetcher, names, args.source_catalog, args.refresh_sources, args.workers)
                def checkpoint(processed):
                    metrics = export_report(store, args.input, args.sheet, output, args.cv_dir)
                    logging.info('CHECKPOINT MASSIVO | documenti=%s | report=%s | stati=%s',
                                 processed, Path(output).resolve(), metrics['states'])
                    print(f'Checkpoint salvato: {processed} documenti | {Path(output).resolve()}', flush=True)

                process_queue(store, fetcher, names, args.workers, args.max_documents,
                              checkpoint=checkpoint,
                              checkpoint_every=args.mass_checkpoint_every)
        except KeyboardInterrupt:
            interrupted = True
            logging.warning('Interruzione richiesta: esporto i risultati già salvati nella coda.')
        metrics = export_report(store, args.input, args.sheet, output, args.cv_dir)
        logging.info('MASSIVO: %s; HTTP=%s; Search API=0', metrics, fetcher.requests)
        print(json.dumps(metrics, ensure_ascii=False, indent=2), flush=True)
        print(f'Report completo: {Path(output).resolve()}\nRipetere lo stesso comando per continuare dalla coda salvata. Search API: 0.')
        return 130 if interrupted else 0
    finally:
        store.close()
        lock.close()
