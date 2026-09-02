from __future__ import annotations

import argparse
import ipaddress
import logging
import os
import re
import shutil
import socket
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from ddgs import DDGS
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


# ============================================================
# VERSIONE / OBIETTIVO
# ============================================================

VERSION = "V7.0 STRUCTURE WORKERS"

# Questa versione è focalizzata SOLO sul dato realmente utile per i farmacisti:
# - dove lavorano attualmente;
# - se risultano in una farmacia privata;
# - se l'ultima struttura pubblica trovata è storica/non confermata.
#
# CV e Gemini sono volutamente esclusi dal flusso standard.
# Questo riduce drasticamente tempi, chiamate e falsi positivi.

INPUT_COLUMNS = (
    "Pers_ID",
    "Cognome",
    "Nome",
    "Data di Nascita",
)

RESULT_COLUMNS = (
    "Stato ricerca",
    "Struttura",
    "Tipo struttura",
    "Citta",
    "Attualita",
    "Confidenza",
    "Note",
    "Fonti",
    "Termine ricerca usato",
    "Ultimo aggiornamento UTC",
    "Errore",
)

DEFAULT_WORKERS = 8
DEFAULT_SEARCH_CONCURRENCY = 4
DEFAULT_SEARCH_RESULTS = 6
DEFAULT_BACKEND = os.getenv("SEARCH_BACKEND", "duckduckgo")
HTTP_TIMEOUT = 6
MAX_PAGE_BYTES = 2 * 1024 * 1024
MAX_PAGE_CHARS = 12000
DEFAULT_SAVE_EVERY = 25

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/151.0 Safari/537.36"
)

LOG_DIR = Path("logs")

SEARCH_SEMAPHORE: threading.BoundedSemaphore | None = None
PRINT_LOCK = threading.Lock()

# ============================================================
# MODELLI DATI
# ============================================================

@dataclass(frozen=True)
class Person:
    row: int
    pers_id: str
    surname: str
    name: str
    birth_date: str


@dataclass(frozen=True)
class SearchResult:
    title: str
    url: str
    snippet: str
    query: str
    query_index: int


@dataclass
class Candidate:
    structure: str
    structure_type: str
    city: str
    freshness: str
    confidence: str
    score: int
    source_url: str
    query: str
    notes: str


# ============================================================
# CLI
# ============================================================

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=f"Ricerca struttura attuale farmacisti - {VERSION}"
    )
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--sheet", default="Dati")
    parser.add_argument("--limit", "--max-rows", dest="max_rows", type=int)
    parser.add_argument("--start-row", type=int, default=2)
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS)
    parser.add_argument(
        "--search-concurrency",
        type=int,
        default=DEFAULT_SEARCH_CONCURRENCY,
        help="Numero massimo di ricerche DDGS contemporanee. "
             "Tenerlo più basso di --workers riduce 403/429.",
    )
    parser.add_argument("--search-results", type=int, default=DEFAULT_SEARCH_RESULTS)
    parser.add_argument("--backend", default=DEFAULT_BACKEND)
    parser.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    parser.add_argument(
        "--retry-all",
        action="store_true",
        help="Rielabora anche righe con Struttura già valorizzata.",
    )
    parser.add_argument(
        "--no-page-fetch",
        action="store_true",
        help="Non apre le pagine dei risultati: più veloce ma meno preciso.",
    )
    parser.add_argument("--log-dir", type=Path, default=LOG_DIR)

    args = parser.parse_args()

    if args.start_row < 2:
        parser.error("--start-row deve essere >= 2")
    if args.max_rows is not None and args.max_rows <= 0:
        parser.error("--limit deve essere > 0")
    if args.workers <= 0:
        parser.error("--workers deve essere > 0")
    if args.search_concurrency <= 0:
        parser.error("--search-concurrency deve essere > 0")
    if args.search_results <= 0:
        parser.error("--search-results deve essere > 0")
    if args.save_every <= 0:
        parser.error("--save-every deve essere > 0")

    return args


# ============================================================
# LOG
# ============================================================

def configure_logging(log_dir: Path) -> Path:
    log_dir = log_dir.expanduser().resolve()
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / f"run_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

    logger = logging.getLogger()
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    fmt = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(threadName)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    fh = logging.FileHandler(log_file, encoding="utf-8")
    sh = logging.StreamHandler()
    fh.setFormatter(fmt)
    sh.setFormatter(fmt)
    logger.addHandler(fh)
    logger.addHandler(sh)

    return log_file


# ============================================================
# UTILITY
# ============================================================

def clean(value: Any) -> str:
    if value is None:
        return ""
    value = str(value).strip()
    if value.casefold() in {"none", "null", "nan", "n/a", "nd", "n.d."}:
        return ""
    return value


def normalize(text: str) -> str:
    text = clean(text).casefold()
    text = text.replace("’", "'")
    text = re.sub(r"[\u00a0\t\r\n]+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def format_birth_date(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    return clean(value)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def canonical_url(url: str) -> str:
    return clean(url).split("#")[0].rstrip("/").casefold()


def get_domain(url: str) -> str:
    try:
        return (urlparse(url).hostname or "").casefold()
    except Exception:
        return ""


def person_full_name(person: Person) -> str:
    return f"{person.name} {person.surname}".strip()


def identity_score(text: str, person: Person) -> int:
    n = normalize(text)
    name = normalize(person.name)
    surname = normalize(person.surname)
    full = normalize(person_full_name(person))
    reverse = normalize(f"{person.surname} {person.name}")

    score = 0
    if full and full in n:
        score += 120
    if reverse and reverse in n:
        score += 120
    if surname and surname in n:
        score += 35
    if name and name in n:
        score += 25
    return score


# ============================================================
# EXCEL
# ============================================================

def normalized_headers(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for cell in ws[1]:
        label = clean(cell.value)
        if label:
            out[label.casefold()] = cell.column
    return out


def require_input_columns(headers: dict[str, int]) -> None:
    missing = [c for c in INPUT_COLUMNS if c.casefold() not in headers]
    if missing:
        raise ValueError("Colonne obbligatorie mancanti: " + ", ".join(missing))


def ensure_result_columns(ws) -> dict[str, int]:
    headers = normalized_headers(ws)
    for label in RESULT_COLUMNS:
        key = label.casefold()
        if key in headers:
            continue
        col = ws.max_column + 1
        cell = ws.cell(1, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers[key] = col
    ws.freeze_panes = "A2"
    return headers


def person_from_row(ws, row: int, headers: dict[str, int]) -> Person:
    return Person(
        row=row,
        pers_id=clean(ws.cell(row, headers["pers_id"]).value),
        surname=clean(ws.cell(row, headers["cognome"]).value),
        name=clean(ws.cell(row, headers["nome"]).value),
        birth_date=format_birth_date(ws.cell(row, headers["data di nascita"]).value),
    )


def output_path_for(input_path: Path, requested: Path | None) -> Path:
    if requested:
        return requested
    return input_path.with_name(f"{input_path.stem}_strutture.xlsx")


def prepare_output(input_path: Path, output_path: Path) -> None:
    if not input_path.exists():
        raise FileNotFoundError(f"Input non trovato: {input_path}")
    if input_path.suffix.casefold() != ".xlsx":
        raise ValueError("L'input deve essere .xlsx")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        logging.info("Riprendo output esistente: %s", output_path)
        return

    if input_path.resolve() != output_path.resolve():
        shutil.copy2(input_path, output_path)
    logging.info("Creata copia di lavoro: %s", output_path)


def atomic_save(workbook, output_path: Path) -> None:
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
    )
    os.close(fd)
    tmp = Path(temp_name)
    try:
        workbook.save(tmp)
        os.replace(tmp, output_path)
    finally:
        tmp.unlink(missing_ok=True)


# ============================================================
# QUERY: SOLO STRUTTURA ATTUALE / FARMACIA PRIVATA
# ============================================================

def build_queries(person: Person) -> list[tuple[int, str]]:
    full = f'"{person_full_name(person)}"'
    return [
        (1, f"{full} farmacista"),
        (2, f"{full} farmacia"),
        (3, f"{full} farmacista ospedale"),
    ]


# ============================================================
# SEARCH
# ============================================================

def search_query(
    query: str,
    query_index: int,
    max_results: int,
    backend: str,
) -> list[SearchResult]:
    global SEARCH_SEMAPHORE
    if SEARCH_SEMAPHORE is None:
        raise RuntimeError("SEARCH_SEMAPHORE non inizializzato")

    logging.info("SEARCH | q=%s | backend=%s", query, backend)

    try:
        with SEARCH_SEMAPHORE:
            raw = DDGS(timeout=HTTP_TIMEOUT).text(
                query,
                region="it-it",
                safesearch="moderate",
                max_results=max_results,
                backend=backend,
            ) or []
    except Exception as exc:
        logging.warning(
            "SEARCH FALLITA | q=%s | %s: %s",
            query,
            type(exc).__name__,
            exc,
        )
        return []

    results: list[SearchResult] = []
    seen: set[str] = set()

    for item in raw:
        url = clean(item.get("href") or item.get("url"))
        if not url.startswith(("http://", "https://")):
            continue

        key = canonical_url(url)
        if key in seen:
            continue
        seen.add(key)

        results.append(
            SearchResult(
                title=clean(item.get("title")),
                url=url,
                snippet=clean(item.get("body") or item.get("snippet")),
                query=query,
                query_index=query_index,
            )
        )

    logging.info("SEARCH OK | q=%s | risultati=%s", query, len(results))
    return results


def relevant_result(result: SearchResult, person: Person) -> bool:
    text = f"{result.title} {result.snippet} {result.url}"
    ident = identity_score(text, person)
    if ident < 120:
        return False

    n = normalize(text)

    context_terms = (
        "farmac",
        "ospedal",
        "asl",
        "ausl",
        "asst",
        "aou",
        "irccs",
        "policlin",
        "ulss",
        "azienda sanitaria",
        "azienda ospedal",
        "linkedin",
    )
    return any(term in n for term in context_terms)


def search_person(
    person: Person,
    max_results: int,
    backend: str,
) -> list[SearchResult]:
    collected: list[SearchResult] = []
    seen: set[str] = set()

    for query_index, query in build_queries(person):
        batch = search_query(query, query_index, max_results, backend)

        for result in batch:
            if not relevant_result(result, person):
                continue
            key = canonical_url(result.url)
            if key in seen:
                continue
            seen.add(key)
            collected.append(result)

        # Early stop: se abbiamo già almeno 2 risultati forti,
        # non sprechiamo la terza query.
        strong = [
            r for r in collected
            if result_pre_score(r, person) >= 250
        ]
        if query_index >= 2 and len(strong) >= 2:
            logging.info(
                "EARLY STOP QUERY | Pers_ID=%s | forti=%s",
                person.pers_id,
                len(strong),
            )
            break

    collected.sort(
        key=lambda r: result_pre_score(r, person),
        reverse=True,
    )
    return collected


# ============================================================
# CLASSIFICAZIONE FONTI
# ============================================================

def is_linkedin(url: str) -> bool:
    return "linkedin.com" in get_domain(url)


def is_institutional(url: str) -> bool:
    d = get_domain(url)
    tokens = (
        "asl", "ausl", "asst", "ats", "aou", "ulss",
        "osped", "policlin", "irccs", "sanita", "salute",
        "regione", "gov.it", "asp.", "asur",
    )
    return any(t in d for t in tokens)


def is_private_pharmacy_domain(url: str) -> bool:
    d = get_domain(url)
    if is_institutional(url):
        return False
    tokens = (
        "farmacia", "farmacie", "farmaci",
    )
    return any(t in d for t in tokens)


def is_historical_document(text: str) -> bool:
    n = normalize(text)
    terms = (
        "graduatoria",
        "concorso",
        "candidati",
        "ammessi",
        "delibera",
        "deliberazione",
        "determinazione",
        "avviso pubblico",
        "bando",
        "commissione",
        "verbale",
        "prova orale",
        "prova scritta",
    )
    return any(t in n for t in terms)


def has_currentness_cues(text: str) -> bool:
    n = normalize(text)
    cues = (
        "attualmente",
        "currently",
        "lavora presso",
        "works at",
        "presso ",
        "in servizio",
        "staff",
        "team",
        "chi siamo",
        "professionisti",
        "responsabile",
        "direttore",
        "collabora",
        "linkedin",
    )
    return any(c in n for c in cues)


def result_pre_score(result: SearchResult, person: Person) -> int:
    text = f"{result.title} {result.snippet} {result.url}"
    n = normalize(text)

    score = identity_score(text, person)

    if is_institutional(result.url):
        score += 100
    if is_linkedin(result.url):
        score += 80
    if is_private_pharmacy_domain(result.url):
        score += 70
    if "farmacista ospedaliero" in n or "dirigente farmacista" in n:
        score += 60
    elif "farmacista" in n:
        score += 35
    if "farmacia " in n:
        score += 25
    if has_currentness_cues(text):
        score += 35
    if is_historical_document(text):
        score -= 70

    bad_domains = (
        "facebook.com", "instagram.com", "pinterest.", "tiktok.com",
        "paginebianche.it", "geneanet.", "ancestry.",
    )
    d = get_domain(result.url)
    if any(x in d for x in bad_domains):
        score -= 120

    return score


# ============================================================
# HTTP: APRE SOLO I MIGLIORI RISULTATI
# ============================================================

def is_public_http_url(url: str) -> bool:
    try:
        p = urlparse(url)
        if p.scheme not in {"http", "https"} or not p.hostname:
            return False
        addresses = socket.getaddrinfo(
            p.hostname,
            p.port or (443 if p.scheme == "https" else 80),
            type=socket.SOCK_STREAM,
        )
        for address in addresses:
            ip = ipaddress.ip_address(address[4][0])
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_reserved
                or ip.is_multicast
            ):
                return False
        return True
    except Exception:
        return False


def fetch_page_text(url: str) -> str:
    if not is_public_http_url(url):
        return ""

    try:
        r = requests.get(
            url,
            headers={"User-Agent": USER_AGENT, "Accept-Language": "it-IT,it;q=0.9"},
            timeout=HTTP_TIMEOUT,
            allow_redirects=True,
        )
        r.raise_for_status()

        if len(r.content) > MAX_PAGE_BYTES:
            return ""

        ctype = r.headers.get("Content-Type", "").casefold()
        # In questa versione non analizziamo PDF: per la struttura attuale
        # preferiamo pagine web correnti e snippet.
        if "pdf" in ctype or r.content.startswith(b"%PDF"):
            return ""

        soup = BeautifulSoup(r.content, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "noscript", "svg"]):
            tag.decompose()
        text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True)).strip()
        return text[:MAX_PAGE_CHARS]
    except Exception as exc:
        logging.debug("FETCH FALLITO | %s | %s", url, exc)
        return ""


# ============================================================
# ESTRAZIONE STRUTTURA
# ============================================================

GENERIC_BAD_STRUCTURE = (
    "farmacia ospedaliera",
    "servizio farmaceutico",
    "azienda sanitaria",
    "azienda ospedaliera",
    "farmacia",
    "ospedale",
    "policlinico",
    "linkedin",
)

STRUCTURE_PATTERNS = [
    # Enti sanitari con sigla + denominazione
    re.compile(
        r"\b(?:ASL|AUSL|ASST|ATS|AOU|AO|ULSS|ASP|ASUR)\s+"
        r"[A-ZÀ-ÖØ-Ý0-9][A-Za-zÀ-ÿ0-9'’.\- ]{1,70}",
        re.IGNORECASE,
    ),
    re.compile(
        r"\bIRCCS\s+[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ0-9'’.\- ]{2,80}",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b(?:Azienda\s+(?:Ospedaliero[- ]Universitaria|Ospedaliera|Sanitaria(?:\s+Locale)?))\s+"
        r"[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ0-9'’.\- ]{2,80}",
        re.IGNORECASE,
    ),
    # Ospedali / policlinici
    re.compile(
        r"\b(?:Ospedale|Policlinico|Presidio\s+Ospedaliero)\s+"
        r"(?:di\s+)?[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ0-9'’.\- ]{2,80}",
        re.IGNORECASE,
    ),
    # Farmacia privata / commerciale
    re.compile(
        r"\bFarmacia\s+(?!Ospedaliera\b|Territoriale\b|Clinica\b)"
        r"[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ0-9'’.\-]{2,60}",
        re.IGNORECASE,
    ),
]

CITY_PATTERNS = [
    re.compile(r"\b(?:a|di|presso)\s+([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ'’\- ]{2,35})\b"),
]


def trim_structure(value: str) -> str:
    value = re.split(
        r"\s(?:\||–|—|-)\s|,|\.\s|;|\b(?:dal|dalla|nel|nella|come|con|presso|dove)\b",
        value,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    value = re.sub(r"\s+", " ", value).strip(" -–—|,.;:")
    return value[:100]


def extract_structure(text: str) -> tuple[str, str]:
    for pattern in STRUCTURE_PATTERNS:
        match = pattern.search(text)
        if not match:
            continue

        value = trim_structure(match.group(0))
        nv = normalize(value)

        if nv in GENERIC_BAD_STRUCTURE:
            continue
        if len(value) < 5:
            continue

        if nv.startswith("farmacia ") and "ospedaliera" not in nv:
            return value, "FARMACIA_PRIVATA"

        return value, "STRUTTURA_SANITARIA"

    return "", ""


def extract_structure_from_linkedin_title(
    result: SearchResult,
    person: Person,
) -> tuple[str, str]:
    if not is_linkedin(result.url):
        return "", ""

    title = clean(result.title)
    # Titoli frequenti:
    # "Mario Rossi - Farmacista presso ASST X | LinkedIn"
    # "Mario Rossi - Pharmacist - Farmacia Y | LinkedIn"
    parts = [
        p.strip()
        for p in re.split(r"\s(?:-|–|—|\|)\s", title)
        if p.strip()
    ]

    full = normalize(person_full_name(person))
    for part in reversed(parts):
        np = normalize(part)
        if np == "linkedin" or full in np:
            continue
        if any(x in np for x in ("farmacista", "pharmacist", "dirigente")):
            # Se il pezzo è solo il ruolo, non è una struttura.
            continue
        if len(part) >= 4:
            if normalize(part).startswith("farmacia "):
                return trim_structure(part), "FARMACIA_PRIVATA"
            return trim_structure(part), "STRUTTURA_SANITARIA"

    # prova nel testo "presso X" / "at X"
    combined = f"{result.title} {result.snippet}"
    m = re.search(
        r"\b(?:presso|at|works at|lavora presso)\s+"
        r"([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ0-9&'’.\- ]{2,80})",
        combined,
        re.IGNORECASE,
    )
    if m:
        structure = trim_structure(m.group(1))
        if normalize(structure).startswith("farmacia "):
            return structure, "FARMACIA_PRIVATA"
        return structure, "STRUTTURA_SANITARIA"

    return "", ""


def extract_city(text: str, structure: str) -> str:
    # Evita di "inventare" città da match troppo generici.
    # Cerca solo pattern espliciti vicini a parole struttura/sede.
    patterns = [
        re.compile(r"\b(?:sede|sede di|con sede a|ubicata a|situata a)\s+([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ'’\- ]{2,35})", re.I),
        re.compile(r"\b(?:ospedale|farmacia|asl|ausl|asst|aou|ulss|irccs)[^.;]{0,80}\b(?:di|a)\s+([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÿ'’\- ]{2,35})", re.I),
    ]
    for pattern in patterns:
        m = pattern.search(text)
        if m:
            city = trim_structure(m.group(1))
            if 2 < len(city) <= 40:
                return city
    return ""


def build_candidate(
    result: SearchResult,
    person: Person,
    page_text: str,
) -> Candidate | None:
    base = f"{result.title} {result.snippet}"
    combined = f"{base} {page_text}".strip()

    if identity_score(base, person) < 120:
        return None

    structure = ""
    structure_type = ""

    # LinkedIn spesso espone struttura direttamente nel title/snippet.
    structure, structure_type = extract_structure_from_linkedin_title(result, person)

    if not structure:
        structure, structure_type = extract_structure(combined)

    if not structure:
        return None

    n = normalize(combined)
    historical = is_historical_document(combined)
    current_cues = has_currentness_cues(base)

    score = result_pre_score(result, person)

    if structure:
        score += 120

    if structure_type == "FARMACIA_PRIVATA":
        score += 40

    if current_cues:
        score += 50

    if historical:
        score -= 90

    if is_institutional(result.url):
        score += 50

    if is_linkedin(result.url):
        score += 35

    # Attualità
    if current_cues and not historical:
        freshness = "ATTUALE_PROBABILE"
        confidence = "alta" if score >= 380 else "media"
    elif historical:
        freshness = "STORICA_DA_VERIFICARE"
        confidence = "bassa"
    else:
        freshness = "DA_VERIFICARE"
        confidence = "media" if score >= 320 else "bassa"

    city = extract_city(combined, structure)

    notes = []
    if structure_type == "FARMACIA_PRIVATA":
        notes.append("Possibile attività in farmacia privata/commerciale.")
    if historical:
        notes.append("La fonte sembra storica/amministrativa: non prova da sola l'impiego attuale.")
    elif current_cues:
        notes.append("La fonte contiene indicatori compatibili con una posizione corrente.")
    else:
        notes.append("Struttura trovata, ma l'attualità non è esplicitamente confermata.")

    return Candidate(
        structure=structure,
        structure_type=structure_type,
        city=city,
        freshness=freshness,
        confidence=confidence,
        score=score,
        source_url=result.url,
        query=result.query,
        notes=" ".join(notes),
    )


# ============================================================
# WORKER
# ============================================================

def research_person(
    person: Person,
    max_results: int,
    backend: str,
    fetch_pages: bool,
) -> dict[str, Any]:
    started = time.perf_counter()

    logging.info(
        "START | row=%s | Pers_ID=%s | %s",
        person.row,
        person.pers_id,
        person_full_name(person),
    )

    results = search_person(
        person=person,
        max_results=max_results,
        backend=backend,
    )

    if not results:
        elapsed = time.perf_counter() - started
        return {
            "row": person.row,
            "status": "NESSUN_RISULTATO",
            "structure": "",
            "structure_type": "",
            "city": "",
            "freshness": "",
            "confidence": "nessuna",
            "notes": "Nessuna fonte pertinente trovata.",
            "sources": "",
            "query": build_queries(person)[-1][1],
            "error": "",
            "elapsed": elapsed,
        }

    candidates: list[Candidate] = []

    # Prima usa solo title/snippet: costo quasi zero.
    for result in results[:6]:
        candidate = build_candidate(result, person, "")
        if candidate:
            candidates.append(candidate)

    # Se non abbiamo una struttura buona, apriamo SOLO i migliori 2 risultati.
    if fetch_pages and not any(c.score >= 380 for c in candidates):
        for result in results[:2]:
            # LinkedIn spesso blocca scraping: lo snippet è più utile della fetch.
            if is_linkedin(result.url):
                continue
            page_text = fetch_page_text(result.url)
            if not page_text:
                continue
            candidate = build_candidate(result, person, page_text)
            if candidate:
                candidates.append(candidate)

    # Deduplica stessa struttura.
    best_by_structure: dict[str, Candidate] = {}
    for c in candidates:
        key = normalize(c.structure)
        previous = best_by_structure.get(key)
        if previous is None or c.score > previous.score:
            best_by_structure[key] = c

    candidates = sorted(
        best_by_structure.values(),
        key=lambda c: c.score,
        reverse=True,
    )

    elapsed = time.perf_counter() - started

    if not candidates:
        return {
            "row": person.row,
            "status": "DA_VERIFICARE",
            "structure": "",
            "structure_type": "",
            "city": "",
            "freshness": "",
            "confidence": "bassa",
            "notes": "Trovate fonti riferibili alla persona, ma nessuna struttura estratta con sufficiente affidabilità.",
            "sources": "\n".join(r.url for r in results[:5]),
            "query": results[0].query,
            "error": "",
            "elapsed": elapsed,
        }

    best = candidates[0]

    if best.freshness == "ATTUALE_PROBABILE" and best.confidence in {"alta", "media"}:
        status = "COMPLETATO"
    else:
        status = "DA_VERIFICARE"

    # Se due fonti forti dicono due strutture diverse, non fingiamo certezza.
    if len(candidates) >= 2:
        c2 = candidates[1]
        if (
            normalize(c2.structure) != normalize(best.structure)
            and c2.score >= best.score - 35
        ):
            status = "DA_VERIFICARE"
            best.confidence = "bassa"
            best.notes += (
                f" Conflitto: altra fonte indica '{c2.structure}'."
            )

    source_urls = [best.source_url]
    for c in candidates[1:3]:
        if c.source_url not in source_urls:
            source_urls.append(c.source_url)

    logging.info(
        "DONE | row=%s | Pers_ID=%s | %.1fs | struttura=%s | tipo=%s | attualita=%s",
        person.row,
        person.pers_id,
        elapsed,
        best.structure,
        best.structure_type,
        best.freshness,
    )

    return {
        "row": person.row,
        "status": status,
        "structure": best.structure,
        "structure_type": best.structure_type,
        "city": best.city,
        "freshness": best.freshness,
        "confidence": best.confidence,
        "notes": best.notes,
        "sources": "\n".join(source_urls),
        "query": best.query,
        "error": "",
        "elapsed": elapsed,
    }


# ============================================================
# SCRITTURA
# ============================================================

def set_cell(ws, row: int, headers: dict[str, int], label: str, value: Any) -> None:
    ws.cell(row, headers[label.casefold()], value)


def write_worker_result(ws, headers: dict[str, int], result: dict[str, Any]) -> None:
    row = int(result["row"])

    values = {
        "Stato ricerca": result.get("status", ""),
        "Struttura": result.get("structure", ""),
        "Tipo struttura": result.get("structure_type", ""),
        "Citta": result.get("city", ""),
        "Attualita": result.get("freshness", ""),
        "Confidenza": result.get("confidence", ""),
        "Note": result.get("notes", ""),
        "Fonti": result.get("sources", ""),
        "Termine ricerca usato": result.get("query", ""),
        "Ultimo aggiornamento UTC": utc_now(),
        "Errore": result.get("error", ""),
    }
    for label, value in values.items():
        set_cell(ws, row, headers, label, value)


# ============================================================
# MAIN
# ============================================================

def main() -> int:
    global SEARCH_SEMAPHORE

    load_dotenv()
    args = parse_args()

    SEARCH_SEMAPHORE = threading.BoundedSemaphore(args.search_concurrency)

    log_file = configure_logging(args.log_dir)

    input_path = args.input.expanduser().resolve()
    output_path = output_path_for(input_path, args.output).expanduser().resolve()

    prepare_output(input_path, output_path)

    wb = load_workbook(output_path)
    if args.sheet not in wb.sheetnames:
        raise ValueError(
            f"Foglio '{args.sheet}' non trovato. Disponibili: {wb.sheetnames}"
        )

    ws = wb[args.sheet]
    headers = normalized_headers(ws)
    require_input_columns(headers)
    headers = ensure_result_columns(ws)
    atomic_save(wb, output_path)

    logging.info("=" * 60)
    logging.info("%s", VERSION)
    logging.info("Input: %s", input_path)
    logging.info("Output: %s", output_path)
    logging.info("Workers: %s", args.workers)
    logging.info("Search concurrency: %s", args.search_concurrency)
    logging.info("Backend: %s", args.backend)
    logging.info("Query/persona: massimo 3, con early-stop")
    logging.info("CV: DISATTIVATO")
    logging.info("Gemini: DISATTIVATO")
    logging.info("Salvataggio ogni %s risultati", args.save_every)
    logging.info("Page fetch: %s", "NO" if args.no_page_fetch else "SI, max 2 risultati/persona")
    logging.info("Log: %s", log_file)
    logging.info("=" * 60)

    people: list[Person] = []

    for row in range(max(2, args.start_row), ws.max_row + 1):
        if args.max_rows is not None and len(people) >= args.max_rows:
            break

        person = person_from_row(ws, row, headers)
        if not person.name and not person.surname:
            continue
        if not person.pers_id:
            logging.warning("SKIP | riga=%s | Pers_ID mancante", row)
            continue

        existing_structure = clean(
            ws.cell(row, headers["struttura"]).value
        ) if "struttura" in headers else ""

        existing_status = clean(
            ws.cell(row, headers["stato ricerca"]).value
        ).upper() if "stato ricerca" in headers else ""

        # Per correggere i vecchi risultati V6.x:
        # COMPLETATO senza Struttura viene RIELABORATO.
        if (
            not args.retry_all
            and existing_structure
            and existing_status == "COMPLETATO"
        ):
            continue

        people.append(person)

    if not people:
        logging.info("Nessuna persona da elaborare.")
        return 0

    total = len(people)
    logging.info("Persone da elaborare: %s", total)

    started_all = time.perf_counter()
    completed = 0
    found = 0
    verify = 0
    none_found = 0
    errors = 0
    pending_save = 0

    # IMPORTANTE:
    # ThreadPoolExecutor è adatto perché il lavoro è quasi tutto I/O di rete.
    # Excel viene scritto SOLO dal thread principale.
    with ThreadPoolExecutor(
        max_workers=args.workers,
        thread_name_prefix="worker",
    ) as executor:
        future_map = {
            executor.submit(
                research_person,
                person,
                args.search_results,
                args.backend,
                not args.no_page_fetch,
            ): person
            for person in people
        }

        for future in as_completed(future_map):
            person = future_map[future]

            try:
                result = future.result()
            except Exception as exc:
                logging.exception(
                    "WORKER ERROR | row=%s | Pers_ID=%s",
                    person.row,
                    person.pers_id,
                )
                result = {
                    "row": person.row,
                    "status": "ERRORE",
                    "structure": "",
                    "structure_type": "",
                    "city": "",
                    "freshness": "",
                    "confidence": "nessuna",
                    "notes": "",
                    "sources": "",
                    "query": "",
                    "error": f"{type(exc).__name__}: {exc}",
                    "elapsed": 0.0,
                }

            write_worker_result(ws, headers, result)

            completed += 1
            pending_save += 1

            status = result.get("status")
            if status == "COMPLETATO":
                found += 1
            elif status == "DA_VERIFICARE":
                verify += 1
            elif status == "NESSUN_RISULTATO":
                none_found += 1
            elif status == "ERRORE":
                errors += 1

            if pending_save >= args.save_every:
                atomic_save(wb, output_path)
                pending_save = 0

            if completed % 10 == 0 or completed == total:
                elapsed = time.perf_counter() - started_all
                rate = completed / elapsed * 60 if elapsed > 0 else 0
                eta_min = (total - completed) / rate if rate > 0 else 0

                logging.info(
                    "PROGRESS | %s/%s | %.1f persone/min | ETA %.1f min | "
                    "completato=%s | verificare=%s | nessuno=%s | errori=%s",
                    completed,
                    total,
                    rate,
                    eta_min,
                    found,
                    verify,
                    none_found,
                    errors,
                )

    atomic_save(wb, output_path)

    elapsed = time.perf_counter() - started_all
    rate = completed / elapsed * 60 if elapsed > 0 else 0

    logging.info("=" * 60)
    logging.info("FINE %s", VERSION)
    logging.info("Elaborate: %s", completed)
    logging.info("COMPLETATO: %s", found)
    logging.info("DA_VERIFICARE: %s", verify)
    logging.info("NESSUN_RISULTATO: %s", none_found)
    logging.info("ERRORE: %s", errors)
    logging.info("Tempo: %.1f minuti", elapsed / 60)
    logging.info("Velocità media: %.1f persone/min", rate)
    logging.info("Output: %s", output_path)
    logging.info("=" * 60)

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\nInterrotto.")
        raise SystemExit(130)
