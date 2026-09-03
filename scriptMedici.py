from __future__ import annotations

import argparse
import base64
import hashlib
import ipaddress
import json
import logging
import os
import random
import re
import shutil
import socket
import tempfile
import time
from dataclasses import dataclass, asdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import quote_plus, urljoin, urlparse, parse_qs, unquote

VERSION = "V4.9 MEDICI - LEGACY DOC EXTRACTION + TARGETED CV SEARCH"

# Import caricati dopo il bootstrap, così i log esistono anche se manca una dipendenza.
requests = None
BeautifulSoup = None
load_dotenv = None
load_workbook = None
Alignment = Font = PatternFill = None
PdfReader = None

DEFAULT_SHEET = "Foglio1"
DEFAULT_SEARCH_RESULTS = 8
DEFAULT_SAVE_EVERY = 10
DEFAULT_DELAY_MIN = 0.15
DEFAULT_DELAY_MAX = 0.35

HTTP_TIMEOUT = 12
MAX_PDF_BYTES = 25 * 1024 * 1024
MAX_DOC_BYTES = 25 * 1024 * 1024
MAX_HTML_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_PDF_TEXT = 120_000
MAX_PAGE_TEXT = 70_000
MAX_LANDING_PAGES_PER_PERSON = 8
MAX_PDF_CANDIDATES_PER_PERSON = 16
MAX_PDF_LINKS_PER_LANDING = 25
MAX_CV_LINKS_PER_LANDING = 30

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/151.0.0.0 Safari/537.36"
)

OUTPUT_COLUMNS = (
    "Ricerca_Stato",
    "Specialita",
    "Specialita_Confidenza",
    "Specialita_Evidenza",
    "CV_Salvato",
    "CV_URL",
    "CV_Confidenza",
    "CV_Da_Verificare",
    "CV_Da_Verificare_URL",
    "Fonti_Ricerca",
    "Ricerca_Metodo",
    "Ricerca_Note",
    "Ultimo_Aggiornamento_UTC",
    "Ricerca_Errore",
)


# ============================================================
# MODELLI
# ============================================================

@dataclass(frozen=True)
class Person:
    row: int
    pers_id: str
    medico_id: str
    surname: str
    name: str
    birth_date: str
    fiscal_code: str
    city: str
    email: str

    @property
    def full_name(self) -> str:
        return f"{self.name} {self.surname}".strip()

    @property
    def output_code(self) -> str:
        mid = numericish(clean(self.medico_id))
        if mid and mid not in {"0", "0.0"}:
            return mid
        return numericish(clean(self.pers_id))


@dataclass(frozen=True)
class WebHit:
    title: str
    url: str
    snippet: str
    engine: str


# ============================================================
# UTIL / BOOTSTRAP
# ============================================================

def clean(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.casefold() in {"none", "null", "nan", "n/a", "n.d.", "nd"}:
        return ""
    return s


def numericish(value: str) -> str:
    s = clean(value)
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    return s


def normalize(text: str) -> str:
    s = clean(text).casefold()
    s = s.replace("’", "'").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", s).strip()


def safe_part(value: str) -> str:
    s = re.sub(r'[<>:"/\\|?*]', "_", clean(value))
    s = re.sub(r"\s+", "_", s).strip(" ._")
    return s or "ND"


def canonical_url(url: str) -> str:
    return clean(url).split("#", 1)[0].rstrip("/")


def decode_bing_redirect(url: str) -> str:
    """
    Decodifica i redirect Bing del tipo /ck/a?...&u=a1BASE64...
    Restituisce l'URL reale quando possibile.
    """
    raw = clean(url)
    try:
        p = urlparse(raw)
        host = (p.hostname or "").casefold()
        if "bing.com" not in host:
            return raw

        qs = parse_qs(p.query)
        values = qs.get("u", [])
        if not values:
            return raw

        value = clean(values[0])
        # Bing usa spesso prefisso a1 + Base64 URL-safe.
        payload = value[2:] if value.startswith("a1") else value
        padding = "=" * (-len(payload) % 4)

        try:
            decoded = base64.urlsafe_b64decode(payload + padding).decode("utf-8", errors="ignore")
            if decoded.startswith(("http://", "https://")):
                return decoded
        except Exception:
            pass

        # In alcuni casi il valore è già percent-encoded o contiene direttamente URL.
        if value.startswith(("http://", "https://")):
            return value
    except Exception:
        pass

    return raw


def normalize_result_url(url: str) -> str:
    u = decode_bing_redirect(url)
    return canonical_url(u)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def domain(url: str) -> str:
    try:
        return (urlparse(url).hostname or "").casefold()
    except Exception:
        return ""


def bootstrap_log_path(log_dir: Path) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    return log_dir / f"startup_medici_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"


def write_bootstrap_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | {message}\n")


def parse_early_paths(argv: list[str]) -> tuple[Path | None, Path | None, Path]:
    input_path = None
    output_path = None
    log_dir = Path("logs")

    for arg in argv[1:]:
        if not arg.startswith("-"):
            input_path = Path(arg)
            break

    if "--output" in argv:
        try:
            output_path = Path(argv[argv.index("--output") + 1])
        except Exception:
            pass

    if "--log-dir" in argv:
        try:
            log_dir = Path(argv[argv.index("--log-dir") + 1])
        except Exception:
            pass

    return input_path, output_path, log_dir


def default_output_path(input_path: Path) -> Path:
    return Path.cwd() / "output" / f"{input_path.stem}_specialita_cv_v4_9.xlsx"


def create_initial_output_copy(input_path: Path | None, output_path: Path | None) -> Path | None:
    if input_path is None:
        return None
    inp = input_path.expanduser().resolve()
    out = output_path.expanduser().resolve() if output_path else default_output_path(inp).resolve()

    if inp.exists() and inp.suffix.casefold() == ".xlsx":
        out.parent.mkdir(parents=True, exist_ok=True)
        if not out.exists() and inp != out:
            shutil.copy2(inp, out)
    return out


def import_dependencies(log_path: Path) -> tuple[bool, list[str]]:
    global requests, BeautifulSoup, load_dotenv
    global load_workbook, Alignment, Font, PatternFill, PdfReader

    missing = []

    try:
        import requests as _requests
        requests = _requests
    except Exception as exc:
        missing.append(f"requests ({exc})")

    try:
        from bs4 import BeautifulSoup as _BeautifulSoup
        BeautifulSoup = _BeautifulSoup
    except Exception as exc:
        missing.append(f"beautifulsoup4 ({exc})")

    try:
        from dotenv import load_dotenv as _load_dotenv
        load_dotenv = _load_dotenv
    except Exception as exc:
        missing.append(f"python-dotenv ({exc})")

    try:
        from openpyxl import load_workbook as _load_workbook
        from openpyxl.styles import Alignment as _Alignment, Font as _Font, PatternFill as _PatternFill
        load_workbook = _load_workbook
        Alignment, Font, PatternFill = _Alignment, _Font, _PatternFill
    except Exception as exc:
        missing.append(f"openpyxl ({exc})")

    try:
        from pypdf import PdfReader as _PdfReader
        PdfReader = _PdfReader
    except Exception as exc:
        missing.append(f"pypdf ({exc})")


    if missing:
        write_bootstrap_log(log_path, "ERRORE | Dipendenze mancanti: " + "; ".join(missing))
        return False, missing

    write_bootstrap_log(log_path, "OK | Dipendenze caricate.")
    return True, []


# ============================================================
# CLI / LOG
# ============================================================

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=VERSION)
    p.add_argument("input", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--sheet", default=DEFAULT_SHEET)
    p.add_argument("--limit", "--max-rows", dest="limit", type=int)
    p.add_argument("--start-row", type=int, default=2)

    p.add_argument("--provider", choices=("serper", "brave", "auto"), default="auto",
                   help="auto = Serper primario, Brave fallback se configurato.")
    p.add_argument("--search-results", type=int, default=DEFAULT_SEARCH_RESULTS)
    p.add_argument("--max-searches-per-person", type=int, default=2,
                   help="Budget massimo di query Search API per medico. Default 2.")
    p.add_argument("--deep", action="store_true",
                   help="Consente una terza query mirata solo se serve.")
    p.add_argument(
        "--cv-searches-per-person", type=int, default=1,
        help="Query CV dedicate aggiuntive per medico. Default 1.",
    )
    p.add_argument(
        "--cv-search-mode",
        choices=("targeted", "broad"),
        default="targeted",
        help=(
            "targeted = query CV solo con identità già forte; "
            "broad = comportamento V4.8 su ogni medico con almeno una fonte."
        ),
    )
    p.add_argument("--save-every", type=int, default=DEFAULT_SAVE_EVERY)
    p.add_argument("--delay-min", type=float, default=DEFAULT_DELAY_MIN)
    p.add_argument("--delay-max", type=float, default=DEFAULT_DELAY_MAX)

    p.add_argument("--retry-all", action="store_true")
    p.add_argument("--retry-errors", action="store_true")
    p.add_argument("--ignore-cache", action="store_true")

    p.add_argument("--cv-dir", type=Path, default=Path("cv_medici"))
    p.add_argument("--cv-review-dir", type=Path, default=Path("cv_medici_da_verificare"))
    p.add_argument("--cache-dir", type=Path, default=Path("cache_medici"))
    p.add_argument("--log-dir", type=Path, default=Path("logs"))

    args = p.parse_args()

    if args.limit is not None and args.limit <= 0:
        p.error("--limit deve essere > 0")
    if args.start_row < 2:
        p.error("--start-row deve essere >= 2")
    if args.search_results <= 0:
        p.error("--search-results deve essere > 0")
    if args.max_searches_per_person <= 0:
        p.error("--max-searches-per-person deve essere > 0")
    if args.cv_searches_per_person < 0:
        p.error("--cv-searches-per-person deve essere >= 0")
    if args.save_every <= 0:
        p.error("--save-every deve essere > 0")
    if args.delay_min < 0 or args.delay_max < args.delay_min:
        p.error("Delay non valido")
    return args


def configure_logging(log_dir: Path) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    path = log_dir / f"run_medici_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"

    root = logging.getLogger()
    root.handlers.clear()
    root.setLevel(logging.INFO)

    fmt = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    fh = logging.FileHandler(path, encoding="utf-8")
    sh = logging.StreamHandler()
    fh.setFormatter(fmt)
    sh.setFormatter(fmt)
    root.addHandler(fh)
    root.addHandler(sh)
    return path


# ============================================================
# EXCEL
# ============================================================

def headers_map(ws) -> dict[str, int]:
    out = {}
    for c in ws[1]:
        label = clean(c.value)
        if label:
            out[label.casefold()] = c.column
    return out


def require_input_columns(headers: dict[str, int]) -> None:
    req = ("Pers_Id", "Pers_Cognome", "Pers_Nome")
    missing = [x for x in req if x.casefold() not in headers]
    if missing:
        raise ValueError("Colonne obbligatorie mancanti: " + ", ".join(missing))


def ensure_output_columns(ws) -> dict[str, int]:
    headers = headers_map(ws)
    for label in OUTPUT_COLUMNS:
        if label.casefold() in headers:
            continue
        col = ws.max_column + 1
        cell = ws.cell(1, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers[label.casefold()] = col
    ws.freeze_panes = "A2"
    return headers


def getv(ws, row: int, headers: dict[str, int], label: str) -> Any:
    col = headers.get(label.casefold())
    return ws.cell(row, col).value if col else ""


def person_from_row(ws, row: int, headers: dict[str, int]) -> Person:
    dob = getv(ws, row, headers, "Pers_DataNascita")
    if isinstance(dob, (datetime, date)):
        dob = dob.strftime("%d/%m/%Y")

    return Person(
        row=row,
        pers_id=numericish(clean(getv(ws, row, headers, "Pers_Id"))),
        medico_id=numericish(clean(getv(ws, row, headers, "Medico_Id"))),
        surname=clean(getv(ws, row, headers, "Pers_Cognome")),
        name=clean(getv(ws, row, headers, "Pers_Nome")),
        birth_date=clean(dob),
        fiscal_code=clean(getv(ws, row, headers, "Pers_CodFis")).upper(),
        city=clean(getv(ws, row, headers, "Indirizzi_Citta")),
        email=clean(getv(ws, row, headers, "emailPredefinita")) or
              clean(getv(ws, row, headers, "Email")),
    )


def prepare_output(input_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not output_path.exists():
        shutil.copy2(input_path, output_path)


def atomic_save(wb, path: Path) -> None:
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.stem}_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def set_result(ws, headers: dict[str, int], row: int, result: dict[str, Any]) -> None:
    mapping = {
        "Ricerca_Stato": result.get("status", ""),
        "Specialita": result.get("specialty", ""),
        "Specialita_Confidenza": result.get("specialty_confidence", ""),
        "Specialita_Evidenza": result.get("specialty_evidence", ""),
        "CV_Salvato": result.get("cv_path", ""),
        "CV_URL": result.get("cv_url", ""),
        "CV_Confidenza": result.get("cv_confidence", ""),
        "CV_Da_Verificare": result.get("cv_review_paths", ""),
        "CV_Da_Verificare_URL": result.get("cv_review_urls", ""),
        "Fonti_Ricerca": result.get("sources", ""),
        "Ricerca_Metodo": result.get("method", ""),
        "Ricerca_Note": result.get("notes", ""),
        "Ultimo_Aggiornamento_UTC": result.get("updated", utc_now()),
        "Ricerca_Errore": result.get("error", ""),
    }
    for label, value in mapping.items():
        ws.cell(row, headers[label.casefold()], value)


# ============================================================
# CACHE
# ============================================================

def cache_path(person: Person, cache_dir: Path) -> Path:
    identity = f"{VERSION}|{person.pers_id}|{person.full_name}|{person.birth_date}|{person.fiscal_code}"
    digest = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:14]
    return cache_dir / f"{safe_part(person.pers_id)}_{digest}.json"


def read_cache(person: Person, cache_dir: Path) -> dict[str, Any] | None:
    p = cache_path(person, cache_dir)
    try:
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            if data.get("version") == VERSION:
                return data.get("result")
    except Exception:
        pass
    return None


def write_cache(person: Person, cache_dir: Path, result: dict[str, Any]) -> None:
    cache_dir.mkdir(parents=True, exist_ok=True)
    p = cache_path(person, cache_dir)
    payload = {"version": VERSION, "person": asdict(person), "result": result}
    fd, tmp_name = tempfile.mkstemp(prefix=".cache_", suffix=".json", dir=cache_dir)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(tmp, p)
    finally:
        tmp.unlink(missing_ok=True)


# ============================================================
# SEARCH API
# ============================================================

class SearchQuotaError(RuntimeError):
    pass


class SearchAuthError(RuntimeError):
    pass


class SearchClient:
    def __init__(self, provider: str, max_results: int,
                 delay_min: float, delay_max: float):
        self.provider = provider
        self.max_results = max_results
        self.delay_min = delay_min
        self.delay_max = delay_max

        self.serper_key = clean(os.getenv("SERPER_API_KEY"))
        self.brave_key = clean(os.getenv("BRAVE_SEARCH_API_KEY"))

        if provider == "serper" and not self.serper_key:
            raise RuntimeError("SERPER_API_KEY mancante nel file .env.")
        if provider == "brave" and not self.brave_key:
            raise RuntimeError("BRAVE_SEARCH_API_KEY mancante nel file .env.")
        if provider == "auto" and not (self.serper_key or self.brave_key):
            raise RuntimeError(
                "Nessuna Search API configurata. Inserisci SERPER_API_KEY nel file .env "
                "(consigliata) oppure BRAVE_SEARCH_API_KEY."
            )

        self.requests_total = 0
        self.requests_by_provider = {"serper": 0, "brave": 0}
        self.failures_by_provider = {"serper": 0, "brave": 0}

    def configured_providers(self) -> list[str]:
        if self.provider == "serper":
            return ["serper"]
        if self.provider == "brave":
            return ["brave"]

        out = []
        if self.serper_key:
            out.append("serper")
        if self.brave_key:
            out.append("brave")
        return out

    def _sleep(self):
        if self.delay_max > 0:
            time.sleep(random.uniform(self.delay_min, self.delay_max))

    def search(self, query: str) -> list[WebHit]:
        last_error = None
        providers = self.configured_providers()

        for idx, provider in enumerate(providers):
            try:
                hits = self._search_provider(provider, query)
                if hits:
                    return hits
            except SearchQuotaError:
                raise
            except SearchAuthError as exc:
                self.failures_by_provider[provider] += 1
                logging.error("SEARCH API AUTH FAIL | %s | %s", provider, exc)
                if self.provider != "auto" or idx == len(providers) - 1:
                    raise
                last_error = exc
                continue
            except Exception as exc:
                last_error = exc
                self.failures_by_provider[provider] += 1
                logging.warning(
                    "SEARCH API FAIL | %s | %s | %s: %s",
                    provider, query, type(exc).__name__, exc
                )

        if last_error:
            logging.warning("SEARCH API EMPTY/FAIL | %s", query)
        return []

    def _search_provider(self, provider: str, query: str) -> list[WebHit]:
        self._sleep()
        self.requests_total += 1
        self.requests_by_provider[provider] += 1

        if provider == "serper":
            return self._serper(query)
        if provider == "brave":
            return self._brave(query)
        raise ValueError(provider)

    def _serper(self, query: str) -> list[WebHit]:
        r = requests.post(
            "https://google.serper.dev/search",
            headers={
                "X-API-KEY": self.serper_key,
                "Content-Type": "application/json",
            },
            json={
                "q": query,
                "gl": "it",
                "hl": "it",
                "num": min(max(self.max_results, 1), 20),
            },
            timeout=HTTP_TIMEOUT,
        )

        if r.status_code == 429:
            raise SearchQuotaError("Serper: crediti/quota esauriti o rate limit.")
        if r.status_code in {401, 403}:
            raise SearchAuthError(
                f"Serper autenticazione rifiutata HTTP {r.status_code}. "
                "Controlla che SERPER_API_KEY sia corretta, attiva e associata a un account con accesso API."
            )
        r.raise_for_status()

        data = r.json()
        out = []
        for item in data.get("organic", [])[:self.max_results]:
            url = clean(item.get("link"))
            if not url.startswith(("http://", "https://")):
                continue
            out.append(WebHit(
                title=clean(item.get("title")),
                url=url,
                snippet=clean(item.get("snippet")),
                engine="serper",
            ))

        logging.info("SEARCH API OK | serper | %s | %s risultati", query, len(out))
        return out

    def _brave(self, query: str) -> list[WebHit]:
        r = requests.get(
            "https://api.search.brave.com/res/v1/web/search",
            headers={
                "Accept": "application/json",
                "X-Subscription-Token": self.brave_key,
            },
            params={
                "q": query,
                "count": min(max(self.max_results, 1), 20),
                "country": "it",
                "search_lang": "it",
            },
            timeout=HTTP_TIMEOUT,
        )

        if r.status_code == 429:
            raise SearchQuotaError("Brave Search: quota/rate limit.")
        if r.status_code in {401, 403}:
            raise SearchAuthError(
                f"Brave Search autenticazione rifiutata HTTP {r.status_code}. "
                "Controlla BRAVE_SEARCH_API_KEY."
            )
        r.raise_for_status()

        data = r.json()
        results = ((data.get("web") or {}).get("results") or [])
        out = []
        for item in results[:self.max_results]:
            url = clean(item.get("url"))
            if not url.startswith(("http://", "https://")):
                continue
            out.append(WebHit(
                title=clean(item.get("title")),
                url=url,
                snippet=clean(item.get("description")),
                engine="brave",
            ))

        logging.info("SEARCH API OK | brave | %s | %s risultati", query, len(out))
        return out

# ============================================================
# HTTP / PDF / PAGINE
# ============================================================

def is_public_http_url(url: str) -> bool:
    try:
        p = urlparse(url)
        if p.scheme not in {"http", "https"} or not p.hostname:
            return False

        try:
            infos = socket.getaddrinfo(
                p.hostname,
                p.port or (443 if p.scheme == "https" else 80),
                type=socket.SOCK_STREAM,
            )
            for info in infos:
                ip = ipaddress.ip_address(info[4][0])
                if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
                    return False
        except OSError:
            pass
        return True
    except Exception:
        return False


def get_bytes(url: str, max_bytes: int) -> tuple[Any, bytes] | None:
    if not is_public_http_url(url):
        return None
    try:
        r = requests.get(
            url,
            headers={
                "User-Agent": USER_AGENT,
                "Accept-Language": "it-IT,it;q=0.9,en;q=0.5",
            },
            timeout=HTTP_TIMEOUT,
            allow_redirects=True,
            stream=True,
        )
        r.raise_for_status()
        buf = bytearray()
        for chunk in r.iter_content(65536):
            if chunk:
                buf.extend(chunk)
                if len(buf) > max_bytes:
                    r.close()
                    return None
        return r, bytes(buf)
    except Exception:
        return None


def pdf_text(raw: bytes) -> str:
    import io
    try:
        reader = PdfReader(io.BytesIO(raw))
        chunks = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            t = page.extract_text() or ""
            if t:
                chunks.append(t)
        return "\n".join(chunks)[:MAX_PDF_TEXT]
    except Exception:
        return ""


def html_text(raw: bytes) -> str:
    try:
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript", "svg"]):
            tag.decompose()
        return clean(soup.get_text(" ", strip=True))[:MAX_PAGE_TEXT]
    except Exception:
        return ""


def page_text(url: str) -> str:
    got = get_bytes(url, MAX_HTML_BYTES)
    if not got:
        return ""
    r, raw = got
    try:
        ctype = normalize(r.headers.get("Content-Type", ""))
        if raw.startswith(b"%PDF") or "application/pdf" in ctype:
            return pdf_text(raw)
        return html_text(raw)
    finally:
        r.close()


# ============================================================
# IDENTITÀ / CV
# ============================================================

CV_POSITIVE = (
    "curriculum vitae", "curriculum professionale", "curriculum formativo",
    "europass", "esperienza professionale", "esperienze professionali",
    "istruzione e formazione", "titoli di studio", "attività professionale",
)

CV_NEGATIVE = (
    "graduatoria", "elenco candidati", "candidati ammessi", "verbale",
    "bando", "avviso pubblico", "prova orale", "prova scritta",
    "commissione esaminatrice", "delibera", "deliberazione",
    "privacy", "cookie", "informativa", "consenso informato",
    "modulo", "manuale", "guida", "brochure", "faq",
    "per saperne di piu", "per saperne di più",
    "conduite a tenir", "ricerca clinica", "ricerca scientifica",
    "patient", "paziente", "istruzioni", "protocollo",
    "linee guida", "raccomandazioni", "prescrizione",
)

MEDICAL_TERMS = (
    "medico", "medicina", "chirurgo", "specialista", "specializzazione",
    "ospedale", "azienda sanitaria", "asl", "ausl", "asst", "aou",
    "irccs", "policlinico", "ordine dei medici",
)


TRUSTED_MEDICAL_DOMAIN_HINTS = (
    "asl", "ausl", "asst", "ats", "aou", "ao-", "osped", "policlin",
    "irccs", "sanita", "salute", "regione", "univ", "universita",
    "ordinemedici", "fnomceo", "gov.it",
)

JUNK_DOMAIN_HINTS = (
    "amazon.", "reddit.", "zhihu.", "baidu.", "qiwa.", "spartex.",
    "facebook.", "instagram.", "tiktok.", "pinterest.", "youtube.",
    "linkedin.", "wikipedia.", "ebay.", "aliexpress.",
)

# Directory/profili in cui una pagina può citare molti medici diversi.
# Su questi domini non basta che il nome compaia nello snippet: il profilo
# deve appartenere chiaramente alla persona cercata.
PROFILE_DIRECTORY_HINTS = (
    "miodottore.", "doctoralia.", "doctolib.", "dottori.it",
    "topdoctors.", "paginebianche.", "paginegialle.",
)

WEAK_SPECIALTY_DOMAIN_HINTS = (
    "qsalute.", "oraridiapertura24.", "paginebianche.", "paginegialle.",
)

NON_PHYSICIAN_ROLE_TERMS = (
    "infermiere", "infermiera", "biologo", "biologa",
    "biologo nutrizionista", "biologa nutrizionista", "nutrizionista",
    "dietista", "psicologo", "psicologa", "fisioterapista",
    "ostetrica", "ostetrico", "farmacista", "tecnico sanitario",
    "tecnico di laboratorio",
)

def weak_specialty_domain(url: str) -> bool:
    d = domain(url)
    return any(x in d for x in WEAK_SPECIALTY_DOMAIN_HINTS)

def source_domain_key(url: str) -> str:
    d = domain(url)
    return d[4:] if d.startswith("www.") else d

def normalized_for_proximity(text: str) -> str:
    s = normalize(text)
    s = re.sub(r"[_/\\?&=+%.,;:|()\[\]{}<>-]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def identity_positions(text: str, person: Person) -> list[tuple[int, int]]:
    n = normalized_for_proximity(text)
    variants = {
        normalized_for_proximity(person.full_name),
        normalized_for_proximity(f"{person.surname} {person.name}"),
    }
    out = []
    for v in variants:
        if not v:
            continue
        pos = 0
        while True:
            i = n.find(v, pos)
            if i < 0:
                break
            out.append((i, i+len(v)))
            pos = i + max(1, len(v))
    return sorted(set(out))

def identity_windows(text: str, person: Person, radius: int = 260) -> list[str]:
    n = normalized_for_proximity(text)
    return [n[max(0,s-radius):min(len(n),e+radius)] for s,e in identity_positions(text, person)]

def target_role_conflict(text: str, person: Person, radius: int = 150) -> bool:
    return any(
        any(term in w for term in NON_PHYSICIAN_ROLE_TERMS)
        for w in identity_windows(text, person, radius)
    )

def disambiguation_score(text: str, person: Person) -> int:
    n = normalize(text)
    score = 0
    if person.fiscal_code and normalize(person.fiscal_code) in n:
        score += 700
    if person.birth_date:
        wanted = {normalize(v) for v in date_variants(person.birth_date)}
        if any(v and v in n for v in wanted):
            score += 350
    if person.city and normalize(person.city) in n:
        score += 90
    return score

def cv_header_region(text: str) -> str:
    """
    Parte del CV che ragionevolmente contiene l'intestatario.
    Si ferma prima delle sezioni dove possono comparire coautori/collaboratori.
    """
    n = text[:2500]
    markers = (
        "esperienze professionali", "esperienza professionale",
        "istruzione e formazione", "formazione",
        "pubblicazioni", "pubblicazione",
        "attivita scientifica", "attività scientifica",
        "esperienze lavorative", "esperienza lavorativa",
    )
    nn = normalize(n)
    cut = len(n)
    for marker in markers:
        idx = nn.find(normalize(marker))
        if idx >= 0:
            cut = min(cut, idx)
    return n[:cut]


def _identity_variants(person: Person) -> list[str]:
    variants = [
        normalized_for_proximity(person.full_name),
        normalized_for_proximity(f"{person.surname} {person.name}"),
    ]
    return [v for v in dict.fromkeys(variants) if v]


def _specialty_alias_occurrences(text: str) -> list[tuple[int, int, str]]:
    n = normalized_for_proximity(text)
    found: list[tuple[int, int, str]] = []

    for alias, canonical in sorted(
        SPECIALTY_ALIASES.items(),
        key=lambda x: len(x[0]),
        reverse=True,
    ):
        a = normalized_for_proximity(alias)
        if not a:
            continue

        start = 0
        while True:
            pos = n.find(a, start)
            if pos < 0:
                break
            found.append((pos, pos + len(a), canonical))
            start = pos + max(1, len(a))

    return found


def specialty_near_identity(
    text: str,
    person: Person,
    max_gap: int = 110,
) -> list[tuple[str, str]]:
    return specialty_relations_raw(text, person, max_gap=max_gap)


def explicit_specialty_near_identity(
    text: str,
    person: Person,
    max_radius: int = 150,
) -> list[tuple[str, str]]:
    """
    Pattern espliciti ("specialista in", "specializzazione in", ecc.) ammessi
    solo nello stesso segmento molto vicino al nome del target.
    """
    out: list[tuple[str, str]] = []
    seen = set()

    for window in identity_windows(text, person, radius=max_radius):
        for pat in EXPLICIT_PATTERNS:
            for m in pat.finditer(window):
                spec = normalize_specialty(clean(m.group(1)))
                if not spec:
                    continue
                key = normalize(spec)
                if key in seen:
                    continue
                seen.add(key)
                out.append((spec, clean(m.group(0))[:350]))

    return out


def directory_profile_location(url: str, person: Person) -> str:
    """
    Prova a ricavare la località strutturata dallo slug di una directory.
    Esempio Doctolib:
      /medico-di-medicina-generale/gragnano/giuseppe-abagnale
    -> gragnano
    """
    if not is_profile_directory(url):
        return ""

    path = unquote(urlparse(url).path or "")
    raw_segments = [x for x in path.split("/") if x]
    segments = [normalized_for_proximity(x) for x in raw_segments]

    name_tokens = set(normalized_for_proximity(person.name).split())
    surname_tokens = set(normalized_for_proximity(person.surname).split())

    identity_idx = -1
    for i, seg in enumerate(segments):
        tokens = set(seg.split())
        if name_tokens and surname_tokens and name_tokens.issubset(tokens) and surname_tokens.issubset(tokens):
            identity_idx = i
            break

    if identity_idx <= 0:
        return ""

    candidate = segments[identity_idx - 1]
    if not candidate:
        return ""

    # Il segmento precedente può essere una specialità, non una città.
    specialty_terms = {
        normalized_for_proximity(alias)
        for alias in SPECIALTY_ALIASES
    }
    generic = (
        "medico", "dottore", "specialista", "chirurgo", "professor",
        "profilo", "doctor", "dr",
    )

    if candidate in specialty_terms:
        return ""
    if any(g in candidate for g in generic):
        return ""

    return candidate


def directory_geo_conflict(
    hit: WebHit,
    person: Person,
) -> tuple[bool, str]:
    """
    Per nomi potenzialmente omonimi, rifiuta un profilo directory quando
    lo slug espone una città diversa da quella del record.
    """
    if not is_profile_directory(hit.url) or not person.city:
        return False, ""

    target_city = normalized_for_proximity(person.city)
    if not target_city:
        return False, ""

    blob = normalized_for_proximity(
        f"{unquote(hit.url)} {hit.title} {hit.snippet}"
    )

    if target_city in blob:
        return False, ""

    profile_city = directory_profile_location(hit.url, person)
    if profile_city and profile_city != target_city:
        return True, f"citta profilo incompatibile: {profile_city} != {target_city}"

    return False, ""


def strong_identity_context(text: str, person: Person) -> bool:
    """
    Per pagine non strutturate: l'identità è più forte se nome+cognome
    compaiono in apertura oppure insieme a città/data/CF.
    """
    if exact_identity_in(text[:1800], person):
        return True
    return disambiguation_score(text, person) >= 90



def _other_identity_boundary(segment: str, person: Person) -> bool:
    """
    Rileva un secondo nome proprio nel segmento che collega target e specialità.
    Opera sul testo originale quando possibile.
    """
    target = normalize(person.full_name)
    # Titolo + Nome Cognome
    titled = re.findall(
        r"\b(?:dott(?:\.|ore|oressa|ssa)?|dr\.?|prof(?:\.|essore|essoressa)?|"
        r"sig(?:\.|ra|nor|nora)?|medico|infermier[ea]|biolog[oa]|psicolog[oa])\s+"
        r"([A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ'`-]+(?:\s+[A-ZÀ-ÖØ-Ý][A-Za-zÀ-ÖØ-öø-ÿ'`-]+){1,3})",
        segment,
        flags=re.I,
    )
    for x in titled:
        if normalize(x) and normalize(x) not in target:
            return True

    # Nome Cognome in maiuscole iniziali.
    pairs = re.findall(
        r"\b[A-ZÀ-ÖØ-Ý][a-zà-öø-ÿ'`-]{2,}\s+[A-ZÀ-ÖØ-Ý][a-zà-öø-ÿ'`-]{2,}\b",
        segment,
    )
    for x in pairs:
        nx = normalize(x)
        if nx and nx not in target:
            return True
    return False


def _target_occurrences_raw(text: str, person: Person) -> list[tuple[int,int]]:
    """
    Occorrenze case-insensitive del nome target nel testo originale.
    """
    variants = [
        clean(person.full_name),
        clean(f"{person.surname} {person.name}"),
    ]
    out = []
    low = text.casefold()
    for variant in variants:
        v = variant.casefold()
        if not v:
            continue
        start = 0
        while True:
            p = low.find(v, start)
            if p < 0:
                break
            out.append((p, p+len(v)))
            start = p + max(1, len(v))
    return sorted(set(out))


def specialty_relations_raw(
    text: str,
    person: Person,
    max_gap: int = 100,
) -> list[tuple[str,str]]:
    """
    Associa alias di specialità al target sul testo originale.
    Se fra target e alias compare un'altra identità, la relazione è spezzata.
    Evita inoltre alias contenuti dentro specialità più lunghe:
    Neuropsichiatria Infantile non genera anche Psichiatria.
    """
    if not text:
        return []

    target_occ = _target_occurrences_raw(text, person)
    if not target_occ:
        return []

    candidates = []
    low = text.casefold()

    for alias, canonical in sorted(
        SPECIALTY_ALIASES.items(), key=lambda x: len(x[0]), reverse=True
    ):
        a = alias.casefold()
        start = 0
        while True:
            p = low.find(a, start)
            if p < 0:
                break
            candidates.append((p, p+len(a), alias, canonical))
            start = p + max(1, len(a))

    # Tieni prima gli alias più lunghi e sopprimi quelli interamente contenuti.
    candidates.sort(key=lambda x: (x[0], -(x[1]-x[0])))
    filtered = []
    for cand in candidates:
        if any(
            cand[0] >= kept[0] and cand[1] <= kept[1]
            for kept in filtered
        ):
            continue
        filtered.append(cand)

    out = []
    seen = set()
    for i0,i1 in target_occ:
        for s0,s1,alias,canonical in filtered:
            if s0 >= i1:
                gap = s0-i1
                segment = text[i1:s0]
            elif i0 >= s1:
                gap = i0-s1
                segment = text[s1:i0]
            else:
                gap = 0
                segment = ""

            if gap > max_gap:
                continue
            if segment and _other_identity_boundary(segment, person):
                continue

            key = normalize(canonical)
            if key in seen:
                continue
            seen.add(key)
            lo=max(0,min(i0,s0)-45)
            hi=min(len(text),max(i1,s1)+70)
            out.append((canonical, clean(text[lo:hi])))
    return out


def generic_source_disambiguated(text: str, person: Person) -> bool:
    if not person.city:
        return strong_identity_context(text, person)
    return disambiguation_score(text, person) >= 90


def structured_profile_bonus(hit: WebHit, person: Person) -> int:
    if not is_profile_directory(hit.url):
        return 0
    bad, _ = directory_geo_conflict(hit, person)
    if bad:
        return -1000
    blob = f"{hit.title} {hit.snippet} {unquote(hit.url)}"
    score = 180
    if person.city and normalize(person.city) in normalize(blob):
        score += 300
    return score


def cv_owner_identity(text: str, person: Person) -> tuple[bool, int, str]:
    if not text:
        return False, 0, "CV senza testo."

    header = cv_header_region(text)
    header_has_name = exact_identity_in(header, person)
    strong = disambiguation_score(text, person)

    if person.fiscal_code and normalize(person.fiscal_code) in normalize(text):
        return True, 1000 + strong, "Codice fiscale del target nel CV."

    if person.birth_date:
        m = re.search(
            r"(?:data\s+di\s+nascita|nato\s+il|nata\s+il).{0,50}"
            r"(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
            text, re.I | re.S,
        )
        if m:
            found = m.group(1).replace(".", "/").replace("-", "/")
            wanted = {x.replace(".", "/").replace("-", "/") for x in date_variants(person.birth_date)}
            if found in wanted and header_has_name:
                return True, 900 + strong, "Nome e data di nascita coerenti nel CV."

    if header_has_name:
        score = 540 + strong
        return True, score, "Nome del target nell'intestazione del CV."

    return False, strong, "Nome del target assente dall'intestazione del CV."


def identity_score(text: str, person: Person) -> int:
    n = normalize(text)
    full = normalize(person.full_name)
    reverse = normalize(f"{person.surname} {person.name}")
    score = 0
    if full and full in n:
        score += 180
    if reverse and reverse in n:
        score += 180
    if normalize(person.surname) in n:
        score += 50
    if normalize(person.name) in n:
        score += 35
    if person.fiscal_code and normalize(person.fiscal_code) in n:
        score += 600
    if person.city and normalize(person.city) in n:
        score += 35
    return score


def date_variants(dob: str) -> set[str]:
    if not dob:
        return set()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            d = datetime.strptime(dob[:10], fmt)
            return {
                d.strftime("%d/%m/%Y"), d.strftime("%d-%m-%Y"),
                d.strftime("%d.%m.%Y"), d.strftime("%Y-%m-%d"),
            }
        except Exception:
            continue
    return {dob}


def verify_cv(text: str, person: Person) -> tuple[bool, str, str]:
    n = normalize(text)
    owner_ok, owner_score, owner_reason = cv_owner_identity(text, person)
    if not owner_ok:
        return False, "bassa", owner_reason

    if person.fiscal_code:
        fiscal_codes = re.findall(r"\b[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]\b", text.upper())
        if fiscal_codes and person.fiscal_code not in fiscal_codes:
            return False, "bassa", "Codice fiscale incompatibile."

    if person.birth_date:
        m = re.search(
            r"(?:data\s+di\s+nascita|nato\s+il|nata\s+il).{0,50}"
            r"(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
            text, re.I | re.S,
        )
        if m:
            found = m.group(1).replace(".", "/").replace("-", "/")
            wanted = {x.replace(".", "/").replace("-", "/") for x in date_variants(person.birth_date)}
            if found not in wanted:
                return False, "bassa", "Data di nascita incompatibile."

    pos = sum(1 for x in CV_POSITIVE if x in n)
    neg = sum(1 for x in CV_NEGATIVE if x in n)
    med = sum(1 for x in MEDICAL_TERMS if x in n)

    if pos < 2 and "curriculum vitae" not in n and "europass" not in n:
        return False, "bassa", "Il documento non appare un CV."
    if med == 0:
        return False, "bassa", "Manca contesto medico."
    if neg >= 3 and "curriculum vitae" not in n:
        return False, "bassa", "Documento concorsuale/amministrativo."

    conf = "alta" if owner_score >= 850 or (owner_score >= 540 and pos >= 4) else "media"
    return True, conf, f"CV verificato sul contenuto. {owner_reason}"
def cv_destination(person: Person, cv_dir: Path) -> Path:
    return cv_dir / f"{safe_part(person.output_code)}-{safe_part(person.surname)}-{safe_part(person.name)}.pdf"


def unique_review_destination(person: Person, review_dir: Path, raw: bytes) -> Path:
    review_dir.mkdir(parents=True, exist_ok=True)
    base = f"{safe_part(person.output_code)}-{safe_part(person.surname)}-{safe_part(person.name)}"
    digest = hashlib.sha1(raw).hexdigest()[:10]
    candidate = review_dir / f"{base}-{digest}.pdf"
    return candidate


def looks_like_cv(text: str) -> tuple[bool, str]:
    n = normalize(text)
    pos = sum(1 for x in CV_POSITIVE if x in n)
    neg = sum(1 for x in CV_NEGATIVE if x in n)
    med = sum(1 for x in MEDICAL_TERMS if x in n)

    if "curriculum vitae" in n or "europass" in n:
        return True, "Titolo/struttura CV rilevata."
    if pos >= 2 and med >= 1 and neg < 4:
        return True, "Struttura compatibile con un CV medico."
    if pos >= 3 and neg < 3:
        return True, "Documento con più sezioni tipiche da CV."
    return False, "Documento non sufficientemente compatibile con un CV."


def save_review_candidate(raw: bytes, text: str, url: str, person: Person,
                          review_dir: Path, reason: str) -> tuple[str, str] | None:
    plausible, why = looks_like_cv(text)
    if not plausible:
        logging.info("CV SCARTATO | %s | %s | %s", person.pers_id, url, reason or why)
        return None

    dest = unique_review_destination(person, review_dir, raw)
    if not dest.exists():
        dest.write_bytes(raw)
    logging.info("CV DA VERIFICARE SALVATO | %s | %s | %s", person.pers_id, dest, reason or why)
    return str(dest), url


def try_pdf_url(url: str, person: Person, cv_dir: Path,
                review_dir: Path) -> tuple[str, str, str, str, str] | None:
    got = get_bytes(url, MAX_PDF_BYTES)
    if not got:
        return None
    r, raw = got
    try:
        ctype = normalize(r.headers.get("Content-Type", ""))
        if not (raw.startswith(b"%PDF") or "application/pdf" in ctype):
            return None

        text = pdf_text(raw)
        if not text:
            logging.info("CV SCARTATO | %s | %s | PDF senza testo estraibile", person.pers_id, url)
            return None

        ok, conf, reason = verify_cv(text, person)
        if ok:
            cv_dir.mkdir(parents=True, exist_ok=True)
            dest = cv_destination(person, cv_dir)
            dest.write_bytes(raw)
            logging.info("CV VERIFICATO | %s | %s | confidenza=%s", person.pers_id, dest, conf)
            return "verified", str(dest), text, conf, ""

        review = save_review_candidate(raw, text, url, person, review_dir, reason)
        if review:
            review_path, _ = review
            return "review", review_path, text, "bassa", reason

        return None
    finally:
        r.close()


def try_cv_candidate(url: str, person: Person, cv_dir: Path,
                     review_dir: Path) -> tuple[str, str, str, str, str] | None:
    direct = try_pdf_url(url, person, cv_dir, review_dir)
    if direct:
        kind, path, text, conf, reason = direct
        return kind, path, url, text, conf

    got = get_bytes(url, MAX_HTML_BYTES)
    if not got:
        return None
    r, raw = got
    try:
        soup = BeautifulSoup(raw, "html.parser")
        links = []

        # Link espliciti
        for a in soup.select("a[href]"):
            href = clean(a.get("href"))
            if not href:
                continue
            candidate = urljoin(url, href)
            label = normalize(f"{a.get_text(' ', strip=True)} {candidate}")

            score = 0
            if ".pdf" in candidate.casefold():
                score += 60
            if "curriculum" in label or "europass" in label:
                score += 120
            if normalize(person.surname) in label:
                score += 50
            if normalize(person.name) in label:
                score += 30
            if any(x in label for x in CV_NEGATIVE):
                score -= 120

            if score >= 60:
                links.append((score, candidate))

        # URL PDF scritti nel markup/testo ma non necessariamente dentro <a>
        html = raw.decode("utf-8", errors="ignore")
        for m in re.findall(r'https?://[^"\'<> ]+?\.pdf(?:\?[^"\'<> ]*)?', html, flags=re.I):
            links.append((70, m))

        seen = set()
        for _, candidate in sorted(links, reverse=True)[:20]:
            key = normalize_result_url(candidate)
            if key in seen:
                continue
            seen.add(key)
            found = try_pdf_url(candidate, person, cv_dir, review_dir)
            if found:
                kind, path, text, conf, reason = found
                return kind, path, candidate, text, conf
    finally:
        r.close()
    return None


def hit_looks_pdf(hit: WebHit) -> bool:
    blob = normalize(f"{hit.title} {hit.snippet} {hit.url}")
    return ".pdf" in hit.url.casefold() or "curriculum" in blob or "europass" in blob

def dedupe_hits(hits: Iterable[WebHit]) -> list[WebHit]:
    """Deduplica i risultati mantenendo il primo URL canonico."""
    out: list[WebHit] = []
    seen: set[str] = set()
    for hit in hits:
        real_url = normalize_result_url(hit.url)
        key = canonical_url(real_url)
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(WebHit(hit.title, real_url, hit.snippet, hit.engine))
    return out




def cv_url_extension(url: str) -> str:
    path = unquote(urlparse(url or "").path or "").casefold()
    for ext in (".pdf", ".docx", ".doc"):
        if path.endswith(ext):
            return ext
    return ""


def looks_like_cv_url(url: str, label: str = "") -> bool:
    blob = normalize(f"{unquote(url or '')} {label or ''}")
    blob = re.sub(r"[_/\\?&=+%.-]+", " ", blob)
    return any(x in f" {blob} " for x in (
        "curriculum", "curriculum vitae", "europass",
        " cv ", "cv medico", "cv dott", "cv dr",
    ))


def plausible_cv_document_url(
    url: str, label: str = "", person: Person | None = None
) -> bool:
    decoded_url = unquote(url or "")
    decoded_label = unquote(label or "")
    raw = f"{decoded_url} {decoded_label}".casefold()
    blob = re.sub(r"[_/\\?&=+%.-]+", " ", raw)
    blob = re.sub(r"\s+", " ", blob).strip()

    if any(normalize(x) in blob for x in CV_NEGATIVE):
        return False

    ext = cv_url_extension(decoded_url)
    cv_signal = looks_like_cv_url(decoded_url, decoded_label)

    has_person = False
    if person is not None:
        n, s = normalize(person.name), normalize(person.surname)
        has_person = bool(n and s and n in blob and s in blob)

    d = domain(decoded_url)
    if d == "media.doctolib.com" or d.endswith(".media.doctolib.com"):
        return cv_signal
    if "/legal/" in decoded_url.casefold():
        return cv_signal
    if d == "esante.gouv.fr" or d.endswith(".esante.gouv.fr"):
        return cv_signal

    if ext in (".pdf", ".docx", ".doc"):
        return cv_signal or has_person

    return cv_signal and (has_person or person is None)


def docx_text(raw: bytes) -> str:
    import zipfile
    import xml.etree.ElementTree as ET
    import io

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            chunks = []
            for name in (
                "word/document.xml",
                "word/header1.xml",
                "word/header2.xml",
                "word/footer1.xml",
            ):
                if name not in zf.namelist():
                    continue
                root = ET.fromstring(zf.read(name))
                vals = []
                for node in root.iter():
                    if node.tag.endswith("}t") and node.text:
                        vals.append(node.text)
                    elif node.tag.endswith("}br"):
                        vals.append("\n")
                if vals:
                    chunks.append(" ".join(vals))
            return "\n".join(chunks)[:MAX_PDF_TEXT]
    except Exception:
        return ""


def _decode_process_output(raw: bytes) -> str:
    for enc in ("utf-8", "cp1252", "latin1"):
        try:
            text = raw.decode(enc, errors="ignore")
            if clean(text):
                return text
        except Exception:
            pass
    return ""


def _legacy_doc_external(raw: bytes) -> tuple[str, str]:
    import subprocess

    tmp_dir = Path(tempfile.mkdtemp(prefix="medici_doc_"))
    src = tmp_dir / "input.doc"
    src.write_bytes(raw)

    try:
        antiword = shutil.which("antiword")
        if antiword:
            try:
                p = subprocess.run([antiword, str(src)], capture_output=True, timeout=30)
                text = _decode_process_output(p.stdout)
                if p.returncode == 0 and len(clean(text)) >= 80:
                    return text[:MAX_PDF_TEXT], "antiword"
            except Exception:
                pass

        catdoc = shutil.which("catdoc")
        if catdoc:
            try:
                p = subprocess.run([catdoc, str(src)], capture_output=True, timeout=30)
                text = _decode_process_output(p.stdout)
                if p.returncode == 0 and len(clean(text)) >= 80:
                    return text[:MAX_PDF_TEXT], "catdoc"
            except Exception:
                pass

        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if soffice:
            try:
                outdir = tmp_dir / "lo"
                outdir.mkdir(exist_ok=True)
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "txt:Text",
                     "--outdir", str(outdir), str(src)],
                    capture_output=True, timeout=45
                )
                files = list(outdir.glob("*.txt"))
                if files:
                    text = files[0].read_text(encoding="utf-8", errors="ignore")
                    if len(clean(text)) >= 80:
                        return text[:MAX_PDF_TEXT], "libreoffice"
            except Exception:
                pass

        if os.name == "nt":
            powershell = shutil.which("powershell") or shutil.which("pwsh")
            if powershell:
                try:
                    dst = tmp_dir / "word_export.txt"
                    ps_src = str(src).replace("'", "''")
                    ps_dst = str(dst).replace("'", "''")
                    script = (
                        "$ErrorActionPreference='Stop';"
                        "$w=New-Object -ComObject Word.Application;"
                        "$w.Visible=$false;"
                        f"$d=$w.Documents.Open('{ps_src}', $false, $true);"
                        f"$d.SaveAs2('{ps_dst}', 2);"
                        "$d.Close($false);$w.Quit();"
                    )
                    p = subprocess.run(
                        [powershell, "-NoProfile", "-NonInteractive",
                         "-Command", script],
                        capture_output=True, timeout=45
                    )
                    if p.returncode == 0 and dst.exists():
                        text = dst.read_text(encoding="utf-8", errors="ignore")
                        if len(clean(text)) < 80:
                            text = dst.read_text(encoding="cp1252", errors="ignore")
                        if len(clean(text)) >= 80:
                            return text[:MAX_PDF_TEXT], "word-com"
                except Exception:
                    pass

        return "", ""
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _legacy_doc_ole_heuristic(raw: bytes) -> str:
    import io as _io
    chunks = []

    try:
        import olefile
        if olefile.isOleFile(_io.BytesIO(raw)):
            ole = olefile.OleFileIO(_io.BytesIO(raw))
            try:
                for stream_name in (
                    "WordDocument", "1Table", "0Table",
                    "\x05SummaryInformation", "\x05DocumentSummaryInformation",
                ):
                    try:
                        if ole.exists(stream_name):
                            chunks.append(ole.openstream(stream_name).read())
                    except Exception:
                        pass
            finally:
                ole.close()
    except Exception:
        pass

    chunks.append(raw)
    texts = []
    for data in chunks:
        for enc in ("utf-16le", "cp1252"):
            try:
                decoded = data.decode(enc, errors="ignore")
            except Exception:
                continue
            texts.extend(re.findall(
                r"[A-Za-zÀ-ÿ0-9][A-Za-zÀ-ÿ0-9\s.,;:/()'’+@_-]{5,}",
                decoded
            ))

    out, seen = [], set()
    for x in texts:
        x = clean(re.sub(r"\s+", " ", x))
        if len(x) < 6:
            continue
        k = normalize(x)
        if not k or k in seen:
            continue
        seen.add(k)
        out.append(x)
    return "\n".join(out)[:MAX_PDF_TEXT]


def legacy_doc_text_with_method(raw: bytes) -> tuple[str, str]:
    text, method = _legacy_doc_external(raw)
    if text:
        return text, method
    text = _legacy_doc_ole_heuristic(raw)
    if len(clean(text)) >= 80:
        return text, "ole/binary-heuristic"
    return "", ""


def legacy_doc_text(raw: bytes) -> str:
    text, _ = legacy_doc_text_with_method(raw)
    return text


def content_extension(url: str, content_type: str, raw: bytes) -> str:
    ext = cv_url_extension(url)
    ctype = normalize(content_type)

    if raw.startswith(b"%PDF") or "application/pdf" in ctype:
        return ".pdf"
    if raw.startswith(b"PK") and ("wordprocessingml" in ctype or ext == ".docx"):
        return ".docx"
    if "application/msword" in ctype or ext == ".doc":
        return ".doc"
    if "text/html" in ctype or raw.lstrip().startswith((b"<!DOCTYPE", b"<html", b"<HTML")):
        return ".html"
    return ext


def extract_cv_document_text(raw: bytes, ext: str) -> str:
    if ext == ".pdf":
        return pdf_text(raw)
    if ext == ".docx":
        return docx_text(raw)
    if ext == ".doc":
        return legacy_doc_text(raw)
    if ext == ".html":
        return html_text(raw)
    return ""


def cv_destination_ext(person: Person, cv_dir: Path, ext: str) -> Path:
    ext = ext if ext in (".pdf", ".docx", ".doc", ".html") else ".bin"
    return cv_dir / (
        f"{safe_part(person.output_code)}-"
        f"{safe_part(person.surname)}-"
        f"{safe_part(person.name)}{ext}"
    )


def unique_review_destination_ext(
    person: Person, review_dir: Path, raw: bytes, ext: str
) -> Path:
    review_dir.mkdir(parents=True, exist_ok=True)
    base = (
        f"{safe_part(person.output_code)}-"
        f"{safe_part(person.surname)}-"
        f"{safe_part(person.name)}"
    )
    digest = hashlib.sha1(raw).hexdigest()[:10]
    ext = ext if ext in (".pdf", ".docx", ".doc", ".html") else ".bin"
    return review_dir / f"{base}-{digest}{ext}"


def save_review_document(
    raw: bytes, text: str, url: str, person: Person,
    review_dir: Path, reason: str, ext: str
):
    allowed, gate_reason = cv_review_allowed(text, url, person, ext)
    if not allowed:
        logging.info(
            "CV SCARTATO | Pers_Id=%s | formato=%s | motivo=%s | url=%s",
            person.pers_id, ext, gate_reason or reason, url
        )
        return None

    dest = unique_review_destination_ext(person, review_dir, raw, ext)
    if not dest.exists():
        dest.write_bytes(raw)

    logging.info(
        "CV DA VERIFICARE SALVATO | Pers_Id=%s | formato=%s | file=%s | motivo=%s | url=%s",
        person.pers_id, ext, dest, reason or gate_reason, url
    )
    return str(dest), url


def text_has_target_identity(text: str, person: Person) -> bool:
    if not text:
        return False
    n = normalize(text)
    return (
        exact_identity_in(text, person)
        or (normalize(person.name) in n and normalize(person.surname) in n)
    )


def cv_review_allowed(text: str, url: str, person: Person, ext: str) -> tuple[bool, str]:
    if text:
        if not text_has_target_identity(text, person):
            return False, "Identità target assente dal documento."
        plausible, why = looks_like_cv(text)
        if not plausible:
            return False, why
        return True, "CV plausibile con identità target; ownership non conclusiva."

    if ext == ".doc":
        blob = normalize(unquote(url))
        if (
            looks_like_cv_url(url)
            and normalize(person.name) in blob
            and normalize(person.surname) in blob
        ):
            return True, "DOC esplicito nome+CV ma testo non estraibile."

    return False, "Candidato insufficiente per verifica manuale."


def strong_identity_for_cv_search(relevant_hits: list[WebHit], person: Person) -> tuple[bool, str]:
    for hit in relevant_hits:
        blob = f"{hit.title} {hit.snippet} {unquote(hit.url)}"
        title_or_url = exact_identity_in(hit.title, person) or url_identity_match(hit.url, person)
        city_ok = not person.city or normalize(person.city) in normalize(blob)

        if title_or_url and trusted_medical_domain(hit.url) and city_ok:
            return True, "fonte sanitaria forte"

        if title_or_url and is_profile_directory(hit.url):
            bad, _ = directory_geo_conflict(hit, person)
            if not bad:
                return True, "profilo professionale strutturato"

        if title_or_url and disambiguation_score(blob, person) >= 90:
            return True, "identità disambiguata"

    return False, "nessuna identità professionale abbastanza forte"


def html_cv_signal(raw: bytes, person: Person) -> tuple[bool, str]:
    try:
        BS = BeautifulSoup
        if BS is None:
            from bs4 import BeautifulSoup as BS

        soup = BS(raw, "html.parser")
        title = clean(soup.title.get_text(" ", strip=True) if soup.title else "")
        headings = " ".join(
            clean(x.get_text(" ", strip=True))
            for x in soup.select("h1,h2,h3")[:20]
        )

        if BeautifulSoup is None:
            text = clean(soup.get_text(" ", strip=True))[:MAX_PAGE_TEXT]
        else:
            text = html_text(raw)

        header_blob = f"{title} {headings} {text[:4000]}"
        n = normalize(header_blob)
        identity = text_has_target_identity(header_blob, person)

        cv_terms = sum(
            1 for t in (
                "curriculum vitae", "curriculum", "europass",
                "esperienze professionali", "esperienza professionale",
                "istruzione e formazione", "formazione",
                "incarichi", "pubblicazioni",
            )
            if normalize(t) in n
        )

        if identity and cv_terms >= 2:
            return True, "HTML con identità target e struttura CV."
        if identity and ("curriculum vitae" in n or "europass" in n):
            return True, "HTML con titolo CV ed identità target."
        return False, "Pagina HTML non sufficientemente CV-specifica."
    except Exception as exc:
        return False, f"HTML non analizzabile: {type(exc).__name__}."


def inspect_cv_document(
    url: str, person: Person, cv_dir: Path, review_dir: Path
):
    got = get_bytes(url, MAX_DOC_BYTES)
    if not got:
        logging.info("CV DOWNLOAD FALLITO | Pers_Id=%s | url=%s", person.pers_id, url)
        return None

    r, raw = got
    try:
        ext = content_extension(url, r.headers.get("Content-Type", ""), raw)
        if ext not in (".pdf", ".docx", ".doc", ".html"):
            logging.info(
                "CV SCARTATO | Pers_Id=%s | formato=%s | formato non supportato | url=%s",
                person.pers_id, ext or "?", url
            )
            return None

        extraction_method = ext.lstrip(".")
        if ext == ".doc":
            text, extraction_method = legacy_doc_text_with_method(raw)
        elif ext == ".html":
            signal_ok, signal_reason = html_cv_signal(raw, person)
            if not signal_ok:
                logging.info(
                    "CV HTML SCARTATO | Pers_Id=%s | motivo=%s | url=%s",
                    person.pers_id, signal_reason, url
                )
                return None
            text = html_text(raw)
            extraction_method = "html"
        else:
            text = extract_cv_document_text(raw, ext)

        logging.info(
            "CV CANDIDATO ANALIZZATO | Pers_Id=%s | formato=%s | estrazione=%s | chars=%s | url=%s",
            person.pers_id, ext, extraction_method or "nessuna",
            len(text or ""), url
        )

        if ext == ".doc" and not text:
            review = save_review_document(
                raw, "", url, person, review_dir,
                "DOC legacy: nessun parser locale ha estratto testo.", ext
            )
            if review:
                return "review", review[0], "", "bassa", "DOC legacy da verificare"
            return None

        if not text:
            return None

        ok, conf, reason = verify_cv(text, person)
        if ok:
            cv_dir.mkdir(parents=True, exist_ok=True)
            dest = cv_destination_ext(person, cv_dir, ext)
            dest.write_bytes(raw)
            logging.info(
                "CV VERIFICATO | Pers_Id=%s | formato=%s | estrazione=%s | confidenza=%s | file=%s | url=%s",
                person.pers_id, ext, extraction_method, conf, dest, url
            )
            return "verified", str(dest), text, conf, ""

        review = save_review_document(raw, text, url, person, review_dir, reason, ext)
        if review:
            return "review", review[0], text, "bassa", reason
        return None
    finally:
        r.close()


def cv_search_queries(person: Person) -> list[str]:
    full = f'"{person.full_name}"'
    qs = []
    if person.city:
        qs.append(
            f'{full} curriculum OR "curriculum vitae" medico "{person.city}"'
        )
    qs.append(
        f'{full} curriculum OR "curriculum vitae" OR europass medico'
    )
    qs.append(
        f'{full} CV medico ospedale ASL OR AUSL OR ASST OR AOU OR IRCCS'
    )

    out, seen = [], set()
    for q in qs:
        k = q.casefold()
        if k not in seen:
            seen.add(k)
            out.append(q)
    return out


def cv_hit_relevance(hit: WebHit, person: Person):
    url = normalize_result_url(hit.url)
    blob = f"{hit.title} {hit.snippet} {unquote(url)}"

    if is_junk_domain(url):
        return False, -1000, "dominio irrilevante"

    if not (
        exact_identity_in(hit.title, person)
        or url_identity_match(url, person)
        or exact_identity_in(hit.snippet, person)
    ):
        return False, 0, "nome completo assente"

    if not plausible_cv_document_url(url, blob, person):
        return False, 0, "nessun segnale CV"

    geo_bad, geo_reason = directory_geo_conflict(hit, person)
    if geo_bad:
        return False, 0, geo_reason

    score = 300
    if exact_identity_in(hit.title, person):
        score += 180
    if url_identity_match(url, person):
        score += 180
    if looks_like_cv_url(url, f"{hit.title} {hit.snippet}"):
        score += 220
    if trusted_medical_domain(url):
        score += 140
    if person.city and normalize(person.city) in normalize(blob):
        score += 100

    return True, score, "candidato CV coerente"


def landing_cv_links(url: str, person: Person) -> list[str]:
    got = get_bytes(url, MAX_HTML_BYTES)
    if not got:
        return []

    r, raw = got
    try:
        ext = content_extension(url, r.headers.get("Content-Type", ""), raw)
        if ext in (".pdf", ".docx", ".doc"):
            return [url] if plausible_cv_document_url(url, person=person) else []

        soup = BeautifulSoup(raw, "html.parser")
        scored = []

        for a in soup.select("a[href]"):
            href = clean(a.get("href"))
            if not href:
                continue

            candidate = urljoin(url, href)
            if not candidate.startswith(("http://", "https://")):
                continue

            label = clean(a.get_text(" ", strip=True))
            ext = cv_url_extension(candidate)
            is_download = ext in (".pdf", ".docx", ".doc")
            cv_signal = looks_like_cv_url(candidate, label)

            if not (is_download or cv_signal):
                continue
            if not plausible_cv_document_url(candidate, label, person):
                continue

            blob = normalize(f"{label} {unquote(candidate)}")
            score = (180 if cv_signal else 0) + (80 if is_download else 0)
            if normalize(person.surname) in blob:
                score += 80
            if normalize(person.name) in blob:
                score += 60
            scored.append((score, candidate))
            logging.info(
                "CV LINK LANDING | Pers_Id=%s | score=%s | formato=%s | label=%s | url=%s",
                person.pers_id, score, ext or "html", label[:160], candidate
            )

        out, seen = [], set()
        for _, candidate in sorted(scored, reverse=True):
            key = normalize_result_url(candidate)
            if key and key not in seen:
                seen.add(key)
                out.append(candidate)
            if len(out) >= MAX_CV_LINKS_PER_LANDING:
                break
        return out
    finally:
        r.close()


def plausible_cv_pdf_url(
    url: str,
    label: str = "",
    person: Person | None = None,
) -> bool:
    """
    Filtro preventivo ad alta precisione.
    Scarta PDF chiaramente generici e, per i PDF scoperti da una landing page,
    richiede almeno un segnale concreto di curriculum/persona.
    """
    decoded_url = unquote(url or "")
    decoded_label = unquote(label or "")
    raw = f"{decoded_url} {decoded_label}".casefold()

    blob = re.sub(r"[_/\\?&=+%.-]+", " ", raw)
    blob = re.sub(r"\s+", " ", blob).strip()

    negative_terms = {normalize(x) for x in CV_NEGATIVE}
    if any(term and term in blob for term in negative_terms):
        return False

    d = domain(decoded_url)
    path = normalize(unquote(urlparse(decoded_url).path))
    path = re.sub(r"[_/\\?&=+%.-]+", " ", path)
    path = re.sub(r"\s+", " ", path).strip()

    cv_signals = (
        "curriculum",
        "curriculum vitae",
        "europass",
        "cv medico",
        "cv dott",
        "cv dr",
    )
    has_cv_signal = any(sig in blob for sig in cv_signals)

    # Media/legal Doctolib e documentazione tecnica sanitaria estera non sono CV,
    # salvo che il documento sia esplicitamente un curriculum.
    if d == "media.doctolib.com" or d.endswith(".media.doctolib.com"):
        return has_cv_signal

    if "/legal/" in decoded_url.casefold() or " legal " in f" {path} ":
        return has_cv_signal

    if d == "esante.gouv.fr" or d.endswith(".esante.gouv.fr"):
        return has_cv_signal

    # Se conosciamo la persona, un PDF può passare anche senza "curriculum"
    # quando nome e cognome sono realmente presenti nel link/anchor.
    if person is not None:
        n_name = normalize(person.name)
        n_surname = normalize(person.surname)
        has_person = bool(
            n_name and n_surname and
            n_name in blob and n_surname in blob
        )
        return has_cv_signal or has_person

    return True


def landing_pdf_links(url: str, person: Person) -> list[str]:
    """
    Apre una landing page via HTTP e raccoglie PDF plausibili.
    Non richiede che nome e cognome siano già nello snippet del motore.
    """
    got = get_bytes(url, MAX_HTML_BYTES)
    if not got:
        return []
    r, raw = got
    try:
        ctype = normalize(r.headers.get("Content-Type", ""))
        if raw.startswith(b"%PDF") or "application/pdf" in ctype:
            return [url] if plausible_cv_pdf_url(url, person=person) else []

        soup = BeautifulSoup(raw, "html.parser")
        scored = []

        for a in soup.select("a[href]"):
            href = clean(a.get("href"))
            if not href:
                continue
            candidate = urljoin(url, href)
            if not candidate.startswith(("http://", "https://")):
                continue

            label = normalize(f"{a.get_text(' ', strip=True)} {unquote(candidate)}")

            if ".pdf" in candidate.casefold() and not plausible_cv_pdf_url(
                candidate, label, person
            ):
                logging.info("PDF PRE-SCARTATO | landing non-CV | %s", candidate)
                continue

            score = 0
            cv_signal = (
                "curriculum" in label
                or "europass" in label
                or " cv " in f" {label} "
            )
            name_hit = normalize(person.name) in label
            surname_hit = normalize(person.surname) in label

            if ".pdf" in candidate.casefold():
                score += 20
            if cv_signal:
                score += 150
            if surname_hit:
                score += 70
            if name_hit:
                score += 50

            # Un PDF trovato dentro una landing page viene seguito solo se
            # mostra "curriculum/CV" oppure nome+cognome della persona.
            if cv_signal or (name_hit and surname_hit):
                scored.append((score, candidate))

        html = raw.decode("utf-8", errors="ignore")
        pdf_url_pattern = r"https?://[^\"'<> ]+?\.pdf(?:\?[^\"'<> ]*)?"
        for m in re.findall(pdf_url_pattern, html, flags=re.I):
            decoded_m = unquote(m)
            if not plausible_cv_pdf_url(decoded_m, person=person):
                logging.info("PDF PRE-SCARTATO | markup non-CV | %s", m)
                continue

            nm = normalize(decoded_m)
            nm = re.sub(r"[_/\\?&=+%.-]+", " ", nm)

            name_hit = normalize(person.name) in nm
            surname_hit = normalize(person.surname) in nm
            cv_signal = "curriculum" in nm or "europass" in nm or " cv " in f" {nm} "

            if not (cv_signal or (name_hit and surname_hit)):
                continue

            score = 20
            if cv_signal:
                score += 150
            if surname_hit:
                score += 70
            if name_hit:
                score += 50

            scored.append((score, m))

        out = []
        seen = set()
        for _, candidate in sorted(scored, reverse=True):
            key = normalize_result_url(candidate)
            if key and key not in seen:
                seen.add(key)
                out.append(candidate)
            if len(out) >= MAX_PDF_LINKS_PER_LANDING:
                break
        return out
    finally:
        r.close()


def inspect_pdf_candidate(url: str, person: Person, cv_dir: Path, review_dir: Path):
    return inspect_cv_document(url, person, cv_dir, review_dir)




# ============================================================
# SPECIALITÀ
# ============================================================

# Alias e titoli comuni. Non è usata per "indovinare": serve solo per normalizzare
# frasi esplicite trovate nelle fonti.
SPECIALTY_ALIASES = {
    "anestesia e rianimazione": "Anestesia e Rianimazione",
    "anestesia rianimazione": "Anestesia e Rianimazione",
    "cardiologia": "Cardiologia",
    "chirurgia generale": "Chirurgia Generale",
    "chirurgia vascolare": "Chirurgia Vascolare",
    "dermatologia": "Dermatologia e Venereologia",
    "dermatologia e venereologia": "Dermatologia e Venereologia",
    "ematologia": "Ematologia",
    "endocrinologia": "Endocrinologia",
    "gastroenterologia": "Gastroenterologia",
    "geriatria": "Geriatria",
    "ginecologia e ostetricia": "Ginecologia e Ostetricia",
    "malattie infettive": "Malattie Infettive",
    "medicina del lavoro": "Medicina del Lavoro",
    "medicina dello sport": "Medicina dello Sport",
    "medicina fisica e riabilitativa": "Medicina Fisica e Riabilitativa",
    "medicina interna": "Medicina Interna",
    "medicina legale": "Medicina Legale",
    "nefrologia": "Nefrologia",
    "neurologia": "Neurologia",
    "neurochirurgia": "Neurochirurgia",
    "oculistica": "Oftalmologia",
    "oftalmologia": "Oftalmologia",
    "oncologia": "Oncologia Medica",
    "oncologia medica": "Oncologia Medica",
    "ortopedia": "Ortopedia e Traumatologia",
    "ortopedia e traumatologia": "Ortopedia e Traumatologia",
    "otorinolaringoiatria": "Otorinolaringoiatria",
    "pediatria": "Pediatria",
    "pneumologia": "Malattie dell'Apparato Respiratorio",
    "psichiatria": "Psichiatria",
    "radiodiagnostica": "Radiodiagnostica",
    "radiologia": "Radiodiagnostica",
    "reumatologia": "Reumatologia",
    "urologia": "Urologia",
    "medicina generale": "Medicina Generale",
    "medico di medicina generale": "Medicina Generale",
    "medico medicina generale": "Medicina Generale",
    "medico di base": "Medicina Generale",
    "medico curante": "Medicina Generale",
    "medicina di famiglia": "Medicina Generale",
    "medico di famiglia": "Medicina Generale",
    "pediatra di libera scelta": "Pediatria",
    "pediatra di famiglia": "Pediatria",
    "chirurgo generale": "Chirurgia Generale",
    "neuropsichiatra infantile": "Neuropsichiatria Infantile",
    "neuropsichiatria infantile": "Neuropsichiatria Infantile",
    "fisiatra": "Medicina Fisica e Riabilitativa",
}

EXPLICIT_PATTERNS = (
    re.compile(r"\bspecialista\s+in\s+([^.;:\n]{3,100})", re.I),
    re.compile(r"\bspecializzato(?:a)?\s+in\s+([^.;:\n]{3,100})", re.I),
    re.compile(r"\bspecializzazione\s+in\s+([^.;:\n]{3,100})", re.I),
    re.compile(r"\bdiploma\s+di\s+specializzazione\s+in\s+([^.;:\n]{3,100})", re.I),
    re.compile(r"\bscuola\s+di\s+specializzazione\s+in\s+([^.;:\n]{3,100})", re.I),
)


def normalize_specialty(raw: str) -> str:
    r = normalize(raw)
    r = re.split(r"\b(?:presso|conseguita|conseguito|università|universita|nel|nell'|anno)\b", r)[0]
    r = r.strip(" ,.-;:")
    if not r:
        return ""

    for alias, canonical in sorted(SPECIALTY_ALIASES.items(), key=lambda x: len(x[0]), reverse=True):
        if alias in r:
            return canonical

    # Se non è in tassonomia, conserviamo solo una forma breve e plausibile.
    raw2 = clean(raw)
    raw2 = re.split(r"[.;:\n|•]", raw2)[0].strip(" ,.-")
    if 3 <= len(raw2) <= 80:
        invalid = {"medico", "medico chirurgo", "medicina e chirurgia", "dirigente medico"}
        if normalize(raw2) not in invalid:
            return raw2
    return ""


def specialty_candidates(
    text: str,
    person: Person,
    source_url: str,
) -> list[tuple[int, str, str, str]]:
    """
    V4.6:
    - niente alias generico dentro una finestra larga;
    - disciplina collegata al nome da distanza stretta o pattern esplicito;
    - fonti deboli non bastano;
    - pagine generiche richiedono identità primaria/forte.
    """
    if not text or identity_score(text, person) < 180:
        return []

    if target_role_conflict(text, person, radius=105):
        return []

    authoritative = trusted_medical_domain(source_url)
    directory = is_profile_directory(source_url)
    weak = weak_specialty_domain(source_url)
    dscore = disambiguation_score(text, person)

    # Fonte generica + città disponibile: nome/cognome da soli non bastano.
    if not authoritative and not directory:
        if not generic_source_disambiguated(text, person):
            return []

    out: list[tuple[int, str, str, str]] = []
    seen = set()

    # 1) Pattern espliciti molto vicini al target.
    for spec, evidence in explicit_specialty_near_identity(text, person, max_radius=145):
        key = normalize(spec)
        if key in seen:
            continue
        seen.add(key)

        score = 510 + min(dscore, 500)
        if authoritative:
            score += 190
        elif directory:
            score += 80
        if weak:
            score -= 220

        out.append((score, spec, source_url, evidence))

    # 2) Alias clinico entro distanza stretta dal nome.
    for spec, evidence in specialty_near_identity(text, person, max_gap=95):
        key = normalize(spec)
        if key in seen:
            continue
        seen.add(key)

        score = 430 + min(dscore, 500)
        if authoritative:
            score += 170
        elif directory:
            score += 70
        if weak:
            score -= 220

        out.append((score, spec, source_url, evidence[:350]))

    return out


def specialty_candidates_from_hit(
    hit: WebHit,
    person: Person,
) -> list[tuple[int, str, str, str]]:
    """
    V4.6:
    - PDF: nessuna specialità dalla SERP prima della verifica ownership;
    - directory: nome + disciplina nel titolo/URL e geografia compatibile;
    - istituzionale: titolo o snippet con relazione stretta nome-disciplina;
    - generico: solo titolo/URL strutturato, mai snippet laterale.
    """
    real_url = normalize_result_url(hit.url)

    # Regola fondamentale: da un PDF non si estrae alcuna specialità
    # finché inspect_pdf_candidate/verify_cv non conferma che è del target.
    if ".pdf" in real_url.casefold():
        return []

    ok, _, _ = hit_relevance(hit, person)
    if not ok:
        return []

    geo_bad, _ = directory_geo_conflict(hit, person)
    if geo_bad:
        return []

    title = hit.title or ""
    snippet = hit.snippet or ""
    decoded_url = unquote(real_url)

    title_id = exact_identity_in(title, person)
    url_id = url_identity_match(real_url, person)
    snippet_id = exact_identity_in(snippet, person)

    if not (title_id or url_id or snippet_id):
        return []

    if target_role_conflict(f"{title} {snippet}", person, radius=105):
        return []

    authoritative = trusted_medical_domain(real_url)
    directory = is_profile_directory(real_url)
    weak = weak_specialty_domain(real_url)
    dscore = disambiguation_score(
        f"{title} {snippet} {decoded_url}",
        person,
    )

    out: list[tuple[int, str, str, str]] = []
    seen = set()

    # ---- Directory personali: titolo/slug strutturato ----
    if directory:
        if not (title_id or url_id):
            return []

        structural_texts = []
        if title_id:
            structural_texts.append(title)
        if url_id:
            structural_texts.append(decoded_url)

        for structural in structural_texts:
            for spec, evidence in specialty_near_identity(
                structural, person, max_gap=115
            ):
                key = normalize(spec)
                if key in seen:
                    continue
                seen.add(key)
                score = 560 + min(dscore, 400)
                score += structured_profile_bonus(hit, person)
                if weak:
                    score -= 220
                out.append((score, spec, real_url, evidence[:350]))

        # Alcuni slug hanno città fra specialità e nome:
        # /specialita/citta/nome-cognome. In questo caso cerchiamo l'alias
        # nello URL intero, ma solo se URL identifica il target e la città
        # è compatibile.
        url_n = normalized_for_proximity(decoded_url)
        if url_id:
            for alias, canonical in sorted(
                SPECIALTY_ALIASES.items(),
                key=lambda x: len(x[0]),
                reverse=True,
            ):
                a = normalized_for_proximity(alias)
                if not a or a not in url_n:
                    continue
                key = normalize(canonical)
                if key in seen:
                    continue
                seen.add(key)
                score = 540 + min(dscore, 400)
                score += structured_profile_bonus(hit, person)
                if weak:
                    score -= 220
                out.append((
                    score,
                    canonical,
                    real_url,
                    clean(f"{title} | {decoded_url}")[:350],
                ))

        return out

    # ---- Fonte sanitaria autorevole ----
    if authoritative:
        # Titolo: forte.
        if title_id:
            for spec, evidence in explicit_specialty_near_identity(
                title, person, max_radius=120
            ):
                key = normalize(spec)
                if key not in seen:
                    seen.add(key)
                    out.append((
                        690 + min(dscore, 400),
                        spec, real_url, evidence[:350]
                    ))

            for spec, evidence in specialty_near_identity(
                title, person, max_gap=90
            ):
                key = normalize(spec)
                if key not in seen:
                    seen.add(key)
                    out.append((
                        640 + min(dscore, 400),
                        spec, real_url, evidence[:350]
                    ))

        # Snippet: ammesso solo con nome e disciplina nello stesso segmento
        # molto stretto.
        if snippet_id:
            for spec, evidence in explicit_specialty_near_identity(
                snippet, person, max_radius=120
            ):
                key = normalize(spec)
                if key not in seen:
                    seen.add(key)
                    out.append((
                        620 + min(dscore, 400),
                        spec, real_url, evidence[:350]
                    ))

            for spec, evidence in specialty_near_identity(
                snippet, person, max_gap=75
            ):
                key = normalize(spec)
                if key not in seen:
                    seen.add(key)
                    out.append((
                        570 + min(dscore, 400),
                        spec, real_url, evidence[:350]
                    ))

        return out

    # ---- Fonte generica ----
    if weak:
        return []

    # Mai snippet generico. Solo titolo o URL chiaramente intestato.
    if title_id:
        for spec, evidence in specialty_near_identity(
            title, person, max_gap=80
        ):
            key = normalize(spec)
            if key not in seen:
                seen.add(key)
                out.append((500, spec, real_url, evidence[:350]))

    if url_id:
        for spec, evidence in specialty_near_identity(
            decoded_url, person, max_gap=85
        ):
            key = normalize(spec)
            if key not in seen:
                seen.add(key)
                out.append((480, spec, real_url, evidence[:350]))

    return out


def choose_specialty(
    candidates: list[tuple[int, str, str, str]]
) -> tuple[str, str, str, str]:
    if not candidates:
        return "", "nessuna", "", ""

    grouped = {}
    for cand in candidates:
        k=normalize(cand[1])
        if k:
            grouped.setdefault(k,[]).append(cand)

    ranked=[]
    for _,group in grouped.items():
        by_domain={}
        for cand in group:
            dk=source_domain_key(cand[2]) or canonical_url(cand[2])
            if dk not in by_domain or cand[0]>by_domain[dk][0]:
                by_domain[dk]=cand
        independent=sorted(by_domain.values(),key=lambda x:x[0],reverse=True)
        if not independent:
            continue
        best=independent[0]
        nd=len(independent)
        ns=sum(c[0]>=500 for c in independent)
        weak_only=all(weak_specialty_domain(c[2]) for c in independent)
        structured=any(is_profile_directory(c[2]) and c[0]>=700 for c in independent)
        total=best[0]+min(220,110*(nd-1))
        ranked.append((total,best,nd,ns,weak_only,structured))

    if not ranked:
        return "", "nessuna", "", ""
    ranked.sort(key=lambda x:x[0],reverse=True)
    score,best,nd,ns,weak_only,structured=ranked[0]

    if weak_only:
        return "", "nessuna", "", ""

    if len(ranked)>1:
        s2,b2,nd2,ns2,w2,st2=ranked[1]
        if (
            not w2 and best[0]>=500 and b2[0]>=500
            and (s2>=score-140 or (structured and st2))
        ):
            return (
                "", "nessuna", "",
                f"CONFLITTO: {best[1]} ({source_domain_key(best[2])}) "
                f"vs {b2[1]} ({source_domain_key(b2[2])})"
            )

    if (
        (trusted_medical_domain(best[2]) and best[0]>=650)
        or structured
        or (nd>=2 and ns>=1 and score>=620)
    ):
        conf="alta"
    elif (
        (is_profile_directory(best[2]) and best[0]>=500)
        or best[0]>=500
        or (nd>=2 and score>=520)
    ):
        conf="media"
    else:
        return "", "nessuna", "", ""

    return best[1],conf,best[2],best[3]


# ============================================================
# QUERY
# ============================================================

def email_domain(email: str) -> str:
    if "@" not in email:
        return ""
    d = email.rsplit("@", 1)[1].strip().casefold()
    if d in {"gmail.com", "hotmail.com", "outlook.com", "libero.it", "virgilio.it", "yahoo.it", "yahoo.com"}:
        return ""
    return d


def research_queries(person: Person, deep: bool) -> list[str]:
    """Query ordinate per precisione; l'algoritmo si ferma appena ha evidenze forti."""
    full = f'"{person.full_name}"'
    reverse = f'"{person.surname} {person.name}"'
    qs = []

    # Prima disambiguiamo con la città: evita omonimi e directory rumorose.
    if person.city:
        qs.append(f'{full} medico "{person.city}"')

    # Query generale per profilo/struttura/specialità.
    qs.append(f'{full} medico specialista ospedale')

    # Query mirata alla disciplina, prima di passare al CV.
    qs.append(f'{full} specializzazione OR specialista OR "medico di"')

    # Un dominio email è utile solo se non è un provider personale. Se sembra
    # sanitario/istituzionale lo anticipiamo; altrimenti resta una query deep.
    d = email_domain(person.email)
    if d and trusted_medical_domain(f'https://{d}'):
        qs.append(f'{full} site:{d} medico specialista curriculum')

    # CV specifico, dopo aver tentato l'identificazione professionale.
    qs.append(f'{full} "curriculum vitae" medico filetype:pdf')

    if deep:
        if d and not trusted_medical_domain(f'https://{d}'):
            qs.append(f'{full} site:{d} medico')
        qs.append(f'{full} ASL OR AUSL OR ASST OR AOU OR IRCCS medico')
        qs.append(f'{reverse} medico specializzazione')

    out = []
    seen = set()
    for q in qs:
        k = q.casefold()
        if k not in seen:
            seen.add(k)
            out.append(q)
    return out


# ============================================================
# PIPELINE
# ============================================================

def is_junk_domain(url: str) -> bool:
    d = domain(url)
    return any(x in d for x in JUNK_DOMAIN_HINTS)


def trusted_medical_domain(url: str) -> bool:
    d = domain(url)
    return any(x in d for x in TRUSTED_MEDICAL_DOMAIN_HINTS)


def exact_identity_in(text: str, person: Person) -> bool:
    n = normalize(text)
    full = normalize(person.full_name)
    reverse = normalize(f"{person.surname} {person.name}")
    return bool((full and full in n) or (reverse and reverse in n))


def url_identity_match(url: str, person: Person) -> bool:
    try:
        parsed = urlparse(normalize_result_url(url))
        # Gli slug dei profili usano spesso -, _, / fra nome e cognome.
        path = re.sub(r"[-_/+.]+", " ", unquote(parsed.path))
        return exact_identity_in(path, person)
    except Exception:
        return False


def is_profile_directory(url: str) -> bool:
    d = domain(url)
    return any(x in d for x in PROFILE_DIRECTORY_HINTS)


def profile_identity_conflict(hit: WebHit, person: Person) -> bool:
    """Rifiuta profili directory chiaramente intestati a un'altra persona."""
    if not is_profile_directory(hit.url):
        return False
    # Per directory/profili la persona deve essere nel titolo o nello slug URL.
    return not (exact_identity_in(hit.title, person) or url_identity_match(hit.url, person))


def hit_relevance(hit: WebHit, person: Person) -> tuple[bool, int, str]:
    """Filtro SERP ad alta precisione: nessun cognome-only e nessun profilo altrui."""
    real_url = normalize_result_url(hit.url)
    d = domain(real_url)
    if not d:
        return False, -999, "dominio assente"
    if is_junk_domain(real_url):
        return False, -500, f"dominio irrilevante: {d}"
    if profile_identity_conflict(hit, person):
        return False, -450, "profilo directory intestato a un'altra persona"

    title_full = exact_identity_in(hit.title, person)
    url_full = url_identity_match(real_url, person)
    snippet_full = exact_identity_in(hit.snippet, person)

    surname = normalize(person.surname)
    name = normalize(person.name)
    title_n = normalize(hit.title)
    title_both = bool(surname and name and surname in title_n and name in title_n)

    blob = normalize(f"{hit.title} {hit.snippet} {real_url}")
    medical_context = any(x in blob for x in MEDICAL_TERMS)
    trusted = trusted_medical_domain(real_url)
    professional_domain = email_domain(person.email)
    same_prof_domain = bool(
        professional_domain and (d == professional_domain or d.endswith("." + professional_domain))
    )

    score = 0
    if title_full:
        score += 360
    elif title_both:
        score += 260
    if url_full:
        score += 320
    if snippet_full:
        score += 150
    if medical_context:
        score += 90
    if trusted:
        score += 130
    if same_prof_domain:
        score += 160
    if "curriculum" in blob or "europass" in blob or ".pdf" in real_url.casefold():
        score += 80
    if person.city and normalize(person.city) in blob:
        score += 35

    # Regole di accettazione:
    # 1. nome completo nel titolo o URL; oppure
    # 2. fonte sanitaria autorevole + nome completo nello snippet; oppure
    # 3. dominio professionale noto + nome completo nello snippet/titolo.
    relevant = (
        title_full
        or url_full
        or (trusted and snippet_full)
        or (same_prof_domain and (snippet_full or title_both))
    )

    if not relevant:
        return False, score, "identita primaria non verificata"
    if not medical_context and not trusted and not same_prof_domain and not hit_looks_pdf(hit):
        return False, score, "manca contesto medico"

    return True, score, "ok"


def landing_identity_ok(
    hit: WebHit,
    text: str,
    person: Person,
) -> tuple[bool, str]:
    if not text:
        return False, "pagina senza testo"

    page_id = identity_score(text, person)
    if page_id < 180:
        return False, "nome completo assente dalla pagina"

    if target_role_conflict(text[:5000], person, radius=105):
        return False, "ruolo incompatibile vicino al nome: possibile omonimo"

    if is_profile_directory(hit.url):
        if not (
            exact_identity_in(hit.title, person)
            or url_identity_match(hit.url, person)
        ):
            return False, "profilo directory di altra persona"

        geo_bad, geo_reason = directory_geo_conflict(hit, person)
        if geo_bad:
            return False, geo_reason

        return True, "profilo directory coerente"

    if trusted_medical_domain(hit.url):
        # La fonte sanitaria può contenere elenchi di molte persone, quindi
        # non basta il nome disperso nel corpo: deve comparire in apertura
        # oppure esserci un disambiguatore forte.
        if strong_identity_context(text, person):
            return True, "fonte sanitaria con identita primaria/disambiguata"

        # Per risultati intestati esattamente al target accettiamo la landing.
        if exact_identity_in(hit.title, person) or url_identity_match(hit.url, person):
            return True, "fonte sanitaria intestata al target"

        return False, "fonte sanitaria con identita solo incidentale"

    if exact_identity_in(hit.title, person) or url_identity_match(hit.url, person):
        if strong_identity_context(text, person):
            return True, "identita primaria confermata"
        return False, "fonte generica senza disambiguazione sufficiente"

    return False, "fonte generica con identita solo incidentale"


def score_hit_for_person(hit: WebHit, person: Person) -> int:
    ok, score, _ = hit_relevance(hit, person)
    if not ok:
        score -= 300
    return score


def research_person(person: Person, args: argparse.Namespace,
                    search_client: SearchClient, cv_dir: Path, review_dir: Path,
                    cache_dir: Path) -> dict[str, Any]:
    if not args.ignore_cache:
        cached = read_cache(person, cache_dir)
        if cached:
            cached["row"] = person.row
            cached["method"] = clean(cached.get("method")) + " + CACHE"
            return cached

    started = time.perf_counter()
    sources = []
    notes = []
    specialty_pool = []

    cv_path = ""
    cv_url = ""
    cv_conf = "nessuna"
    cv_review_paths = []
    cv_review_urls = []

    # -------- SEARCH API ADATTIVA --------
    all_hits = []
    relevant_hits = []
    queries_used = 0

    budget = args.max_searches_per_person + (1 if args.deep else 0)

    for q in research_queries(person, args.deep):
        if queries_used >= budget:
            break

        hits = search_client.search(q)
        queries_used += 1
        all_hits.extend(hits)
        all_hits = dedupe_hits(all_hits)

        relevant_hits = []
        for hit in all_hits:
            ok, rel_score, reason = hit_relevance(hit, person)
            if ok:
                relevant_hits.append(hit)
            else:
                logging.debug(
                    "HIT SCARTATO | Pers_Id=%s | score=%s | %s | %s",
                    person.pers_id, rel_score, reason, hit.url
                )

        # Stop anticipato soltanto quando abbiamo già estratto una
        # specialità affidabile oppure trovato un PDF plausibile come CV.
        if relevant_hits:
            specialty_preview = []
            for preview_hit in relevant_hits:
                specialty_preview.extend(
                    specialty_candidates_from_hit(preview_hit, person)
                )

            strong_pdf = any(
                hit_looks_pdf(h)
                and plausible_cv_pdf_url(h.url, f"{h.title} {h.snippet}", person)
                for h in relevant_hits
            )

            if specialty_preview or strong_pdf:
                break

    cv_queries_used = 0
    cv_gate_ok = False
    cv_gate_reason = "nessuna fonte rilevante"

    if relevant_hits:
        if args.cv_search_mode == "broad":
            cv_gate_ok = True
            cv_gate_reason = "modalità broad"
        else:
            cv_gate_ok, cv_gate_reason = strong_identity_for_cv_search(relevant_hits, person)

    if relevant_hits and args.cv_searches_per_person > 0:
        existing_cv = any(
            plausible_cv_document_url(
                h.url, f"{h.title} {h.snippet}", person
            )
            for h in relevant_hits
        )

        logging.info(
            "CV SEARCH GATE | Pers_Id=%s | mode=%s | allowed=%s | reason=%s | existing_cv=%s",
            person.pers_id, args.cv_search_mode, cv_gate_ok, cv_gate_reason, existing_cv
        )

        if not existing_cv and cv_gate_ok:
            for q in cv_search_queries(person):
                if cv_queries_used >= args.cv_searches_per_person:
                    break

                logging.info("CV SEARCH QUERY | Pers_Id=%s | q=%s", person.pers_id, q)
                hits = search_client.search(q)
                cv_queries_used += 1
                queries_used += 1

                for hit in hits:
                    ok, rel_score, reason = cv_hit_relevance(hit, person)
                    if ok:
                        relevant_hits.append(hit)
                        logging.info(
                            "CV HIT ACCETTATO | Pers_Id=%s | score=%s | reason=%s | url=%s",
                            person.pers_id, rel_score, reason, hit.url
                        )
                    else:
                        logging.debug(
                            "CV HIT SCARTATO | Pers_Id=%s | score=%s | %s | %s",
                            person.pers_id, rel_score, reason, hit.url
                        )

    ranked_hits = sorted(
        dedupe_hits(relevant_hits),
        key=lambda h: (
            1 if plausible_cv_document_url(
                h.url, f"{h.title} {h.snippet}", person
            ) else 0,
            score_hit_for_person(h, person)
        ),
        reverse=True,
    )

    logging.info(
        "HIT SUMMARY | Pers_Id=%s | query=%s | cv_query=%s | grezzi=%s | rilevanti=%s | provider=%s",
        person.pers_id, queries_used, cv_queries_used, len(all_hits), len(ranked_hits),
        search_client.requests_by_provider
    )

    # 1) CV/documenti diretti dai risultati: PDF, DOCX, DOC, HTML.
    doc_seen = set()
    doc_checked = 0

    for hit in ranked_hits:
        if doc_checked >= MAX_PDF_CANDIDATES_PER_PERSON:
            break

        label = f"{hit.title} {hit.snippet}"
        ext = cv_url_extension(hit.url)
        html_cv = not ext and looks_like_cv_url(hit.url, label)

        if ext not in (".pdf", ".docx", ".doc") and not html_cv:
            continue
        if not plausible_cv_document_url(hit.url, label, person):
            continue

        key = normalize_result_url(hit.url)
        if not key or key in doc_seen:
            continue

        doc_seen.add(key)
        doc_checked += 1

        found = inspect_cv_document(hit.url, person, cv_dir, review_dir)
        if not found:
            continue

        kind, found_path, cv_text, found_conf, reason = found
        sources.insert(0, hit.url)

        if kind == "verified":
            cv_path, cv_url, cv_conf = found_path, hit.url, found_conf
            notes.append(
                f"CV verificato trovato direttamente ({Path(found_path).suffix})."
            )
            if cv_text:
                specialty_pool.extend(
                    specialty_candidates(cv_text, person, hit.url)
                )
            break

        if found_path not in cv_review_paths:
            cv_review_paths.append(found_path)
        if hit.url not in cv_review_urls:
            cv_review_urls.append(hit.url)
        notes.append(
            f"Documento CV candidato salvato ({Path(found_path).suffix})."
        )

    # 2) Landing page: analizziamo i migliori risultati anche senza identità
    # già perfettamente visibile nello snippet.
    landing_checked = 0
    for hit in ranked_hits:
        if cv_path:
            break
        if landing_checked >= MAX_LANDING_PAGES_PER_PERSON:
            break
        if ".pdf" in hit.url.casefold():
            continue

        rank_score = score_hit_for_person(hit, person)
        blob = normalize(f"{hit.title} {hit.snippet} {hit.url}")
        if rank_score < 80 and not any(
            x in blob for x in (
                "curriculum", "medico", "osped", "asl", "ausl",
                "asst", "aou", "irccs", "policlin"
            )
        ):
            continue

        landing_checked += 1

        text = page_text(hit.url)
        page_ok, page_reason = landing_identity_ok(hit, text, person)
        if not page_ok:
            logging.info(
                "LANDING SCARTATA | Pers_Id=%s | %s | %s",
                person.pers_id, page_reason, hit.url
            )
            continue

        sources.append(hit.url)
        specialty_pool.extend(specialty_candidates(text, person, hit.url))

        for candidate in landing_cv_links(hit.url, person):
            if doc_checked >= MAX_PDF_CANDIDATES_PER_PERSON:
                break

            if not plausible_cv_document_url(candidate, person=person):
                continue

            key = normalize_result_url(candidate)
            if not key or key in doc_seen:
                continue

            doc_seen.add(key)
            doc_checked += 1

            found = inspect_cv_document(candidate, person, cv_dir, review_dir)
            if not found:
                continue

            kind, found_path, cv_text, found_conf, reason = found
            sources.insert(0, candidate)

            if kind == "verified":
                cv_path, cv_url, cv_conf = found_path, candidate, found_conf
                notes.append(
                    f"CV verificato trovato tramite landing ({Path(found_path).suffix})."
                )
                if cv_text:
                    specialty_pool.extend(
                        specialty_candidates(cv_text, person, candidate)
                    )
                break

            if found_path not in cv_review_paths:
                cv_review_paths.append(found_path)
            if candidate not in cv_review_urls:
                cv_review_urls.append(candidate)
            notes.append(
                f"Documento CV da landing salvato ({Path(found_path).suffix})."
            )

        if cv_path:
            break

    # 3) Specialità da SERP verificata.
    # Unica via SERP: parser V4.5 con prossimità e controllo identità.
    for hit in ranked_hits:
        specialty_pool.extend(specialty_candidates_from_hit(hit, person))

    if all_hits and not ranked_hits:
        notes.append("I motori hanno restituito risultati, ma nessuno ha superato il filtro identità/fonte sanitaria.")

    specialty, specialty_conf, specialty_source, specialty_evidence = choose_specialty(specialty_pool)
    if specialty_source:
        sources.insert(0, specialty_source)

    if specialty and cv_path:
        status = "COMPLETATO"
    elif specialty:
        status = "SPECIALITA_TROVATA"
    elif cv_path:
        status = "CV_TROVATO_SPECIALITA_DA_VERIFICARE"
    elif cv_review_paths:
        status = "DA_VERIFICARE"
    elif sources:
        status = "DA_VERIFICARE"
    else:
        status = "NESSUN_RISULTATO"

    sources = list(dict.fromkeys(normalize_result_url(u) for u in sources if u.startswith(("http://", "https://"))))

    result = {
        "row": person.row,
        "status": status,
        "specialty": specialty,
        "specialty_confidence": specialty_conf,
        "specialty_evidence": specialty_evidence,
        "cv_path": cv_path,
        "cv_url": cv_url,
        "cv_confidence": cv_conf,
        "cv_review_paths": "\n".join(cv_review_paths[:20]),
        "cv_review_urls": "\n".join(cv_review_urls[:20]),
        "sources": "\n".join(sources[:12]),
        "method": f"Search API V4.9 ({args.provider})",
        "notes": " ".join(notes),
        "updated": utc_now(),
        "error": "",
        "elapsed": time.perf_counter() - started,
    }

    write_cache(person, cache_dir, result)

    logging.info(
        "DONE | Pers_Id=%s | %.1fs | %s | specialita=%s | cv=%s | cv_review=%s",
        person.pers_id, result["elapsed"], status,
        specialty or "N/D", "SI" if cv_path else "NO", len(cv_review_paths)
    )
    return result


# ============================================================
# MAIN
# ============================================================

def main() -> int:
    args = parse_args()
    log_file = configure_logging(args.log_dir)

    input_path = args.input.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else default_output_path(input_path).resolve()
    )
    cv_dir = args.cv_dir.expanduser().resolve()
    review_dir = args.cv_review_dir.expanduser().resolve()
    cache_dir = args.cache_dir.expanduser().resolve()

    if not input_path.exists():
        raise FileNotFoundError(input_path)

    prepare_output(input_path, output_path)

    wb = load_workbook(output_path)
    if args.sheet not in wb.sheetnames:
        raise ValueError(f"Foglio '{args.sheet}' non trovato. Disponibili: {wb.sheetnames}")

    ws = wb[args.sheet]
    headers = headers_map(ws)
    require_input_columns(headers)
    headers = ensure_output_columns(ws)

    # Avvisa chiaramente se si sta riutilizzando un file con risultati vecchi.
    method_col = headers.get("Ricerca_Metodo")
    previous_methods = set()
    if method_col:
        max_probe = min(ws.max_row, 200)
        for probe_row in range(2, max_probe + 1):
            value = clean(ws.cell(probe_row, method_col).value)
            if value:
                previous_methods.add(value)

    if previous_methods and not all("V4.9" in m for m in previous_methods):
        logging.warning(
            "OUTPUT CONTIENE RISULTATI DI VERSIONI PRECEDENTI | %s",
            sorted(previous_methods)[:8]
        )
        logging.warning(
            "Per test confrontabili usa un file output nuovo, ad esempio risultatiMediciV4_9.xlsx"
        )

    atomic_save(wb, output_path)

    logging.info("=" * 78)
    logging.info("%s", VERSION)
    logging.info("Input: %s", input_path)
    logging.info("Output: %s", output_path)
    logging.info("Provider: %s", args.provider)
    logging.info("Deep: %s", args.deep)
    logging.info("Search results: %s", args.search_results)
    logging.info("Max searches/persona: %s (+1 con --deep)", args.max_searches_per_person)
    logging.info("Delay: %.1f - %.1f sec", args.delay_min, args.delay_max)
    logging.info("CV sicuri dir: %s", cv_dir)
    logging.info("CV da verificare dir: %s", review_dir)
    logging.info("Cache dir: %s", cache_dir)
    logging.info("Log: %s", log_file)
    logging.info("=" * 78)

    print(f"Excel output: {output_path}")
    print(f"Log: {log_file.resolve()}")

    people = []
    for row in range(args.start_row, ws.max_row + 1):
        if args.limit is not None and len(people) >= args.limit:
            break

        p = person_from_row(ws, row, headers)
        if not p.pers_id or not p.surname or not p.name:
            continue

        status = clean(getv(ws, row, headers, "Ricerca_Stato")).upper()
        if not args.retry_all and status in {"COMPLETATO", "SPECIALITA_TROVATA"}:
            continue
        if status == "ERRORE" and not args.retry_errors:
            continue

        people.append(p)

    logging.info("Medici da elaborare: %s", len(people))

    if not people:
        atomic_save(wb, output_path)
        return 0

    counts = {}
    start_all = time.perf_counter()
    unsaved = 0

    search_client = SearchClient(
        provider=args.provider,
        max_results=args.search_results,
        delay_min=args.delay_min,
        delay_max=args.delay_max,
    )

    for i, person in enumerate(people, start=1):
        try:
            result = research_person(
                person, args, search_client, cv_dir, review_dir, cache_dir
            )
        except SearchAuthError as exc:
            logging.error("AUTH SEARCH API | %s", exc)
            result = {
                "row": person.row,
                "status": "BLOCCATO_AUTENTICAZIONE_API",
                "specialty": "",
                "specialty_confidence": "nessuna",
                "specialty_evidence": "",
                "cv_path": "",
                "cv_url": "",
                "cv_confidence": "nessuna",
                "cv_review_paths": "",
                "cv_review_urls": "",
                "sources": "",
                "method": f"Search API V4.9 ({args.provider})",
                "notes": "Autenticazione Search API rifiutata. Elaborazione interrotta.",
                "updated": utc_now(),
                "error": str(exc),
            }

        except SearchQuotaError as exc:
            logging.error("QUOTA SEARCH API | %s", exc)
            result = {
                "row": person.row,
                "status": "BLOCCATO_QUOTA_RICERCA",
                "specialty": "",
                "specialty_confidence": "nessuna",
                "specialty_evidence": "",
                "cv_path": "",
                "cv_url": "",
                "cv_confidence": "nessuna",
                "cv_review_paths": "",
                "cv_review_urls": "",
                "sources": "",
                "method": f"Search API V4.9 ({args.provider})",
                "notes": "Quota Search API esaurita o rate limit.",
                "updated": utc_now(),
                "error": str(exc),
            }
        except Exception as exc:
            logging.exception("ERRORE PERSONA | %s | %s", person.pers_id, person.full_name)
            result = {
                "row": person.row,
                "status": "ERRORE",
                "specialty": "",
                "specialty_confidence": "nessuna",
                "specialty_evidence": "",
                "cv_path": "",
                "cv_url": "",
                "cv_confidence": "nessuna",
                "cv_review_paths": "",
                "cv_review_urls": "",
                "sources": "",
                "method": f"Search API V4.9 ({args.provider})",
                "notes": "",
                "updated": utc_now(),
                "error": f"{type(exc).__name__}: {exc}",
            }

        set_result(ws, headers, person.row, result)
        status = clean(result.get("status")) or "?"
        counts[status] = counts.get(status, 0) + 1
        unsaved += 1

        if unsaved >= args.save_every:
            atomic_save(wb, output_path)
            unsaved = 0
            logging.info("CHECKPOINT | %s/%s | %s", i, len(people), output_path)

        if i % 5 == 0 or i == len(people):
            elapsed = time.perf_counter() - start_all
            rate = i / elapsed * 60 if elapsed else 0.0
            eta = (len(people) - i) / rate if rate else 0.0
            logging.info(
                "PROGRESS | %s/%s | %.2f medici/min | ETA %.1f min | %s | API=%s",
                i, len(people), rate, eta, counts, search_client.requests_by_provider
            )

        if status in {"BLOCCATO_QUOTA_RICERCA", "BLOCCATO_AUTENTICAZIONE_API"}:
            atomic_save(wb, output_path)
            if status == "BLOCCATO_QUOTA_RICERCA":
                logging.error("STOP | Quota Search API esaurita.")
            else:
                logging.error("STOP | Autenticazione Search API rifiutata.")
            break

    atomic_save(wb, output_path)

    elapsed = time.perf_counter() - start_all
    logging.info("=" * 78)
    logging.info("FINE | %.1f min | %s", elapsed / 60, counts)
    logging.info("API REQUESTS | totale=%s | %s",
                 search_client.requests_total, search_client.requests_by_provider)
    logging.info("OUTPUT | %s", output_path)
    logging.info("=" * 78)

    print("")
    print(f"Excel salvato in: {output_path}")
    print(f"Log esecuzione: {log_file.resolve()}")
    print(f"CV sicuri salvati in: {cv_dir}")
    print(f"CV da verificare salvati in: {review_dir}")
    print(f"Richieste Search API: {search_client.requests_total} | {search_client.requests_by_provider}")
    return 0


if __name__ == "__main__":
    import sys

    early_input, early_output, early_log_dir = parse_early_paths(sys.argv)
    startup_log = bootstrap_log_path(early_log_dir)

    try:
        initial_output = create_initial_output_copy(early_input, early_output)
        if initial_output:
            write_bootstrap_log(startup_log, f"INFO | Excel iniziale: {initial_output}")
            print(f"Excel di output: {initial_output}")
            print(f"Log startup: {startup_log.resolve()}")
    except Exception as exc:
        write_bootstrap_log(
            startup_log,
            f"ERRORE | Creazione output iniziale: {type(exc).__name__}: {exc}"
        )

    ok, missing = import_dependencies(startup_log)
    if not ok:
        print("\nERRORE: mancano dipendenze obbligatorie.")
        print("Installa con: python -m pip install -r requirements_medici_v4.txt")
        print(f"Log diagnostico: {startup_log.resolve()}")
        for item in missing:
            print(f" - {item}")
        raise SystemExit(2)

    try:
        load_dotenv()
        script_env = Path(__file__).resolve().parent / ".env"
        if script_env.exists():
            load_dotenv(script_env, override=False)
        raise SystemExit(main())
    except KeyboardInterrupt:
        write_bootstrap_log(startup_log, "INTERRUZIONE | KeyboardInterrupt")
        print("\nInterrotto.")
        raise SystemExit(130)
    except Exception as exc:
        write_bootstrap_log(startup_log, f"FATAL | {type(exc).__name__}: {exc}")
        print(f"\nERRORE FATALE: {type(exc).__name__}: {exc}")
        print(f"Log diagnostico: {startup_log.resolve()}")
        raise
