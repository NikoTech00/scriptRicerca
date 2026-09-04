import importlib.util
import logging
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

spec = importlib.util.spec_from_file_location('medici_under_test', Path(__file__).resolve().parents[1] / 'scriptMedici.py')
m = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = m
spec.loader.exec_module(m)


class FailureHandlingTests(unittest.TestCase):
    def test_specialty_fallback_trims_degree_metadata(self):
        self.assertEqual(m.normalize_specialty("Scienza dell'Alimentazione nel 2003 col massimo dei voti"), "Scienza dell'Alimentazione")

    def test_shared_source_requires_demographic_match(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '03/04/1980', '', '', '')
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'cv.pdf'
            path.write_bytes(b'pdf')
            base = 'Curriculum vitae Anna Rossi. Medico. Istruzione e formazione. Specializzata in Cardiologia.'
            for suffix, expected in [('', 'DA_VERIFICARE'), (' Data di nascita 3 aprile 1980.', 'COMPLETATO')]:
                with patch.object(m, 'extract_cv_document_text', return_value=base + suffix):
                    result = m.research_local_person(person, [path], {str(path)})
                self.assertEqual(result['status'], expected)


    def test_birth_date_words_reject_homonym(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '24/04/1961', '', '', '')
        text = 'Curriculum vitae Anna Rossi. Data di nascita 13 agosto 1966. Medico. Istruzione e formazione.'
        self.assertFalse(m.verify_cv(text, person)[0])
        self.assertIn('incompatibile', m.verify_cv(text, person)[2])

    def test_equivalent_birth_date_formats(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '03/04/1980', '', '', '')
        for value in ['3/4/1980', '03.04.1980', '1980-04-03', '3 aprile 1980', '03 APRILE 1980']:
            with self.subTest(value=value):
                text = f'Curriculum vitae Anna Rossi. Data di nascita {value}. Medico. Istruzione e formazione.'
                self.assertTrue(m.verify_cv(text, person)[0])
        self.assertIsNone(m.parse_birth_date('31 febbraio 1980'))

    def test_training_and_teaching_are_not_completed_specialties(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        for text in ['Lavora in un centro specializzato in Cardiologia.',
                     'Medico specialista in formazione in Cardiologia.',
                     'Occupazione desiderata: Medico Chirurgo specializzato in Cardiologia.',
                     'Scuola di Specializzazione in Cardiologia.',
                     'Specializzazione in Cardiologia (in corso).',
                     'Medico specializzando in Cardiologia.',
                     'Docente nella Scuola di Specializzazione in Cardiologia.']:
            with self.subTest(text=text):
                self.assertEqual(m.verified_cv_specialty_candidates(text, person, 'cv.pdf'), [])
        self.assertEqual(m.verified_cv_specialty_candidates('Specializzata in Cardiologia.', person, 'cv.pdf')[0][1], 'Cardiologia')

    def test_completed_degree_not_lost_due_to_old_training(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        text = 'Diploma di Specializzazione in Pediatria, conseguito nel 2020.\nEsperienza passata: Scuola di Specializzazione in Pediatria.'
        result = m.verified_cv_specialty_candidates(text, person, 'cv.pdf')
        self.assertEqual([c[1] for c in result], ['Pediatria'])

    def test_direct_sources_cache_and_limit(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            manifest = folder / 'sources.csv'
            manifest.write_text('Pers_Id,URL\n1,https://example.org/a.pdf\n1,https://example.org/b.pdf\n')
            client = m.DirectSources(manifest, folder/'cache', 1)
            response = Mock(url='https://example.org/a.pdf', headers={'Content-Type':'application/pdf'})
            with patch.object(m, 'get_bytes', return_value=(response, b'%PDF-test')) as get, patch.object(m.time, 'sleep'):
                paths = client.documents(person)
                self.assertEqual(len(paths), 1)
                self.assertEqual(get.call_count, 1)
                self.assertIn('Limite', client.notes[0])
                self.assertEqual(client.documents(person), paths)
                self.assertEqual(get.call_count, 1)
                paths[0].write_bytes(b'corrupted')
                self.assertEqual(client.documents(person), [])

    def test_direct_download_failure_is_not_a_negative_search(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            manifest = folder / 'sources.csv'
            manifest.write_text('Pers_Id,URL\n1,https://example.org/a.pdf\n')
            client = m.DirectSources(manifest, folder/'cache', 1)
            with patch.object(m, 'get_bytes', return_value=None), patch.object(m.time, 'sleep'):
                self.assertEqual(client.documents(person), [])
            self.assertIn('non scaricabile', client.notes[0])
            self.assertEqual(m.research_local_person(person, [])['status'], 'DA_VERIFICARE')

    def test_redirect_to_private_address_is_not_fetched(self):
        response = Mock(status_code=302, headers={'Location':'http://127.0.0.1/private'})
        with patch.object(m, 'requests', Mock(get=Mock(return_value=response))) as req, patch.object(m, 'is_public_http_url', side_effect=lambda url: '127.0.0.1' not in url):
            self.assertIsNone(m.get_bytes('https://example.org/cv', 1000))
            self.assertEqual(req.get.call_count, 1)
            response.close.assert_called_once()

    def test_generic_biography_is_not_html_cv(self):
        from bs4 import BeautifulSoup
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        raw = b'<html><title>Anna Rossi - Biografia</title><h1>Anna Rossi</h1><nav>Curriculum Formazione Pubblicazioni</nav><p>Medico specialista in Cardiologia</p></html>'
        with patch.object(m, 'BeautifulSoup', BeautifulSoup):
            self.assertFalse(m.html_cv_signal(raw, person)[0])

    def test_multiple_completed_degrees_in_same_cv(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        with tempfile.TemporaryDirectory() as directory:
            doc = Path(directory)/'1-Rossi-Anna.pdf'
            doc.write_bytes(b'fake')
            text = 'Curriculum vitae Anna Rossi. Medico. Istruzione e formazione. Diploma di Specializzazione in Chirurgia Generale. Diploma di Specializzazione in Chirurgia Plastica e Ricostruttiva.'
            with patch.object(m, 'extract_cv_document_text', return_value=text):
                result = m.research_local_person(person, [doc])
            self.assertEqual(result['specialty'], 'Chirurgia Generale; Chirurgia Plastica e Ricostruttiva')

    def client(self, provider='serper'):
        with patch.dict(m.os.environ, {'SERPER_API_KEY': 'test-secret', 'BRAVE_SEARCH_API_KEY': 'test-brave'}):
            return m.SearchClient(provider, 8, 0, 0)

    def test_http400_is_error_and_secret_is_redacted(self):
        response = Mock(status_code=400)
        response.json.return_value = {'message': 'Invalid test-secret'}
        with patch.object(m, 'requests', Mock(post=Mock(return_value=response))):
            with self.assertRaises(m.SearchRequestError) as caught:
                self.client().search('query')
        self.assertIn('400', str(caught.exception))
        self.assertNotIn('test-secret', str(caught.exception))

    def test_timeout_is_not_empty_result(self):
        client = self.client()
        with patch.object(client, '_search_provider', side_effect=TimeoutError('timeout')):
            with self.assertRaises(m.SearchRequestError):
                client.search('query')

    def test_valid_empty_response_stays_empty(self):
        client = self.client()
        with patch.object(client, '_search_provider', return_value=[]):
            self.assertEqual(client.search('query'), [])

    def test_auto_fallback_can_recover(self):
        client = self.client('auto')
        hit = m.WebHit('title', 'https://example.org', 'snippet', 'brave')
        with patch.object(client, '_search_provider', side_effect=[m.SearchRequestError('400'), [hit]]):
            self.assertEqual(client.search('query'), [hit])

    def test_locked_output_preserves_complete_recovery_file(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'result.xlsx'
            target.write_bytes(b'previous')
            workbook = Mock()
            workbook.save.side_effect = lambda path: path.write_bytes(b'complete-new-workbook')
            with patch.object(m.os, 'replace', side_effect=PermissionError('locked')), patch.object(m.time, 'sleep'):
                with self.assertRaisesRegex(RuntimeError, 'recupero'):
                    m.atomic_save(workbook, target)
            self.assertEqual(target.read_bytes(), b'previous')
            recovery = list(Path(directory).glob('.result_*.xlsx'))
            self.assertEqual(len(recovery), 1)
            self.assertEqual(recovery[0].read_bytes(), b'complete-new-workbook')

    def test_transient_lock_retries_and_replaces(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'result.xlsx'
            workbook = Mock()
            workbook.save.side_effect = lambda path: path.write_bytes(b'new')
            real_replace = m.os.replace
            calls = []
            def replace(src, dst):
                calls.append(src)
                if len(calls) == 1:
                    raise PermissionError('temporary lock')
                real_replace(src, dst)
            with patch.object(m.os, 'replace', side_effect=replace), patch.object(m.time, 'sleep'):
                m.atomic_save(workbook, target)
            self.assertEqual(target.read_bytes(), b'new')
            self.assertEqual(list(Path(directory).glob('.result_*.xlsx')), [])

    def test_cleanup_does_not_mask_save_error(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook = Mock()
            workbook.save.side_effect = ValueError('original save failure')
            with patch.object(Path, 'unlink', side_effect=PermissionError('locked')):
                with self.assertRaisesRegex(ValueError, 'original save failure'):
                    m.atomic_save(workbook, Path(directory) / 'result.xlsx')

    def test_credit_exhaustion_is_quota(self):
        response = Mock(status_code=400)
        response.json.return_value = {'message': 'Not enough credits'}
        with patch.object(m, 'requests', Mock(post=Mock(return_value=response))):
            with self.assertRaises(m.SearchQuotaError):
                self.client().search('query')

    def test_request_limit_includes_failures(self):
        client = self.client()
        client.max_requests = 1
        with patch.object(client, '_serper', side_effect=TimeoutError('timeout')) as api:
            with self.assertRaises(m.SearchRequestError):
                client.search('one')
            with self.assertRaises(m.SearchBudgetError):
                client.search('two')
            self.assertEqual(api.call_count, 1)

    def test_query_cache_survives_new_client_and_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            client = self.client()
            client.search_cache_dir = Path(directory)
            hit = m.WebHit('title', 'https://example.org', 'text', 'serper')
            with patch.object(client, '_serper', return_value=[hit]) as api:
                self.assertEqual(client.search('same'), [hit])
                self.assertEqual(api.call_count, 1)
            other = self.client()
            other.search_cache_dir = Path(directory)
            other.requests_total = other.max_requests
            with patch.object(other, '_serper', side_effect=AssertionError('network forbidden')):
                self.assertEqual(other.search('same'), [hit])
            self.assertEqual(other.cache_hits, 1)

    def test_failed_query_not_cached(self):
        with tempfile.TemporaryDirectory() as directory:
            client = self.client()
            client.search_cache_dir = Path(directory)
            with patch.object(client, '_serper', side_effect=m.SearchQuotaError('quota')):
                with self.assertRaises(m.SearchQuotaError):
                    client.search('same')
            self.assertEqual(list(Path(directory).glob('*.json')), [])

    def test_old_result_cache_is_not_trusted(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            m.cache_path(person, folder).write_text(m.json.dumps({'version': m.VERSION, 'result': {'status': 'NESSUN_RISULTATO'}}))
            self.assertIsNone(m.read_cache(person, folder))

    def test_specialty_normalization_removes_false_conflict(self):
        self.assertEqual(m.normalize_specialty('Ginecologia & Ostetricia Ospedale San Gerardo'),
                         m.normalize_specialty('Ginecologia e Ostetricia'))
        self.assertEqual(m.normalize_specialty('Malattie dell’Apparato Respiratorio'),
                         m.normalize_specialty('Pneumologia'))
        self.assertEqual(m.normalize_specialty('Neurochirurgia'), 'Neurochirurgia')

    def test_valid_empty_query_is_cached(self):
        with tempfile.TemporaryDirectory() as directory:
            client = self.client()
            client.search_cache_dir = Path(directory)
            with patch.object(client, '_serper', return_value=[]) as api:
                self.assertEqual(client.search('same'), [])
                self.assertEqual(client.search('same'), [])
                self.assertEqual(api.call_count, 1)

    def test_refresh_query_cache_consumes_budget(self):
        with tempfile.TemporaryDirectory() as directory:
            client = self.client()
            client.search_cache_dir = Path(directory)
            with patch.object(client, '_serper', return_value=[]):
                client.search('same')
            client.refresh_cache = True
            client.max_requests = 1
            with self.assertRaises(m.SearchBudgetError):
                client.search('same')

    def test_document_matching_does_not_accept_another_name(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        index = {'1': [Path('1-Rossi-Anna.pdf'), Path('1-Rossi-Anna_Maria.pdf'), Path('1-Rossi-Anna-deadbeef12.doc')]}
        self.assertEqual(m.local_documents(person, index), [index['1'][0], index['1'][2]])

    def test_offline_rejects_cv_of_another_person(self):
        person = m.Person(2, '1', '', 'Rossi', 'Anna', '', '', '', '')
        with tempfile.TemporaryDirectory() as directory:
            doc = Path(directory) / '1-Rossi-Anna.pdf'
            doc.write_bytes(b'fake')
            with patch.object(m, 'extract_cv_document_text', return_value='Curriculum vitae Mario Bianchi. Medico specialista in Cardiologia. Istruzione e formazione.'), patch.object(m, 'get_bytes', side_effect=AssertionError('network forbidden')):
                result = m.research_local_person(person, [doc])
            self.assertEqual(result['status'], 'DA_VERIFICARE')
            self.assertEqual(result['specialty'], '')

    def test_offline_report_preserves_input_and_uses_zero_network(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            source = folder / 'input.xlsx'
            output = folder / 'output.xlsx'
            book = Workbook()
            book.active.title = 'Foglio1'
            book.active.append(['Pers_Id', 'Pers_Cognome', 'Pers_Nome', 'Specialita', 'CV_URL'])
            book.active.append(['1', 'Rossi', 'Anna', 'OLD', 'https://old.invalid'])
            book.active.append(['2', 'Bianchi', 'Mario', '', ''])
            book.save(source)
            before = source.read_bytes()
            docs = folder / 'docs'
            docs.mkdir()
            (docs / '1-Rossi-Anna.pdf').write_bytes(b'fake')
            args = SimpleNamespace(input=source, output=output, cv_dir=docs, cv_review_dir=folder/'missing', sheet='Foglio1', limit=None, start_row=2)
            text = 'Curriculum vitae Anna Rossi. Medico. Istruzione e formazione. Specializzazione in Cardiologia.'
            with patch.multiple(m, load_workbook=load_workbook, Font=Font, PatternFill=PatternFill, Alignment=Alignment), patch.object(m, 'extract_cv_document_text', return_value=text), patch.object(m, 'get_bytes', side_effect=AssertionError('network forbidden')), patch.object(m, 'SearchClient', side_effect=AssertionError('API forbidden')):
                self.assertEqual(m.run_local_recovery(args), 0)
                with self.assertRaises(ValueError):
                    m.run_local_recovery(args)
            self.assertEqual(source.read_bytes(), before)
            result = load_workbook(output)
            sheet = result.active
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet['D2'].value, 'Cardiologia')
            self.assertIsNone(sheet['E2'].value)
            result.close()


if __name__ == '__main__':
    logging.disable(logging.CRITICAL)
    unittest.main()
