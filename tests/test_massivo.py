import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import medici_massivo as m


def person(pid='1', name='Anna', surname='Rossi', dob='03/04/1980'):
    return m.core.Person(2, pid, '', surname, name, dob, '', 'Roma', '')


class MassivoTests(unittest.TestCase):
    def test_readable_cv_path_uses_person_code_surname_and_name(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'cache' / 'abc.pdf'
            source.parent.mkdir()
            source.write_bytes(b'%PDF-test')
            person = SimpleNamespace(pers_id='123', surname='De Rossi', name='Anna Maria')
            target = m.readable_cv_path(source, person, root / 'cv', 'https://example.test/cv.pdf')
            self.assertEqual(target.name, '123_De_Rossi_Anna_Maria.pdf')
            self.assertTrue(target.is_absolute())
            self.assertEqual(target.read_bytes(), source.read_bytes())

    def test_indexed_activity_is_nominal_and_not_a_documented_specialty(self):
        p = person()
        spec = {'id': 'Doctolib', 'activity_pattern': r'^/([^/]+)/', 'activity_map': {}}
        result = m.indexed_activity_result(
            spec, p, 'https://www.doctolib.it/cardiologo/roma/anna-rossi', False
        )
        self.assertEqual(result['identity'], 'solo_nome_completo')
        self.assertEqual(result['specialties'], [])
        self.assertEqual(result['activities'], ['Cardiologia'])
        self.assertIn('Categoria pubblica Doctolib', result['activity_evidence'][0])

    def test_indexed_activity_rejects_ambiguous_name(self):
        spec = {'id': 'Doctolib', 'activity_pattern': r'^/([^/]+)/', 'activity_map': {}}
        result = m.indexed_activity_result(
            spec, person(), 'https://www.doctolib.it/pediatra/roma/anna-rossi', True
        )
        self.assertEqual(result['identity'], 'omonimia')
        self.assertEqual(result['activities'], [])

    def test_dentist_role_is_a_declared_discipline(self):
        spec = {'id': 'Doctolib', 'activity_pattern': r'^/([^/]+)/',
                'activity_map': {'odontoiatra': 'Odontoiatria e Stomatologia'}}
        result = m.indexed_activity_result(
            spec, person(), 'https://www.doctolib.it/odontoiatra/roma/anna-rossi', False
        )
        self.assertEqual(result['specialties'], [])
        self.assertEqual(result['activities'], ['Odontoiatria e Stomatologia'])

    def test_numeric_profile_is_queued_then_matched_from_title(self):
        names = m.Names({'1': person()})
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://example.org/medico/123/0</loc></url></urlset>'
        spec = {'id': 'numeric', 'type': 'sitemap', 'urls': ['https://example.org/sitemap.xml'],
                'hosts': ['example.org'], 'profile_pattern': '/medico/', 'content_match': True}
        from unittest.mock import Mock
        fetcher = Mock(); fetcher.get.return_value = ({'final_url': spec['urls'][0]}, raw)
        found, errors = m.discover_one(spec, fetcher, names)
        self.assertEqual(found, {('', 'https://example.org/medico/123/0', 'probe')})
        self.assertEqual(errors, [])
        with tempfile.TemporaryDirectory() as folder:
            store = m.Store(folder)
            try:
                fetcher.folder = Path(folder)
                url = 'https://example.org/medico/123/0'; store.add_probe(url, 'numeric'); store.db.commit()
                info = {'file': 'fake.html', 'final_url': url, 'sha': 'hash'}
                outcome = {'info': info, 'text': 'Specializzata in Cardiologia.', 'cv': False,
                           'title': 'Dott.ssa Anna Rossi', 'links': []}
                m.save_outcome(store, fetcher, names, url, outcome)
                self.assertEqual(store.db.execute('SELECT pid FROM candidates').fetchone()[0], '1')
                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM evidence').fetchone()[0], 1)
            finally:
                store.close()

    def test_seo_title_exposes_declared_specialty(self):
        title, text, _ = m.visible_profile(
            b'<title>Dott. Anna Rossi: specialista in Oftalmologia a Roma | Directory</title>'
              b'<main><h1>Dott. Anna Rossi</h1></main>')
        result = m.analyze_content(person(), text, False, title, False)
        self.assertEqual(result['activities'], ['Oftalmologia'])

    def test_schema_org_medical_specialty_outside_main_is_extracted(self):
        raw = (b'<title>D.ssa Anna Rossi - Directory</title><main><h1>D.ssa Anna Rossi</h1></main>'
               b'<meta itemprop="medicalSpecialty" content="cardiologia">')
        title, text, _ = m.visible_profile(raw)
        result = m.analyze_content(person(), text, False, title, False)
        self.assertEqual(result['specialties'], ['Cardiologia'])
        self.assertEqual(result['evidence'], ['Specialista in cardiologia'])

    def test_institutional_specialty_label_is_declared_activity(self):
        result = m.analyze_content(
            person(),
            'Specialità: Radiodiagnostica, Radiologia',
            False,
            'Dott.ssa Anna Rossi',
            False,
        )
        self.assertEqual(result['specialties'], [])
        self.assertEqual(result['activities'], ['Radiodiagnostica'])
        self.assertIn('Specialità:', result['activity_evidence'][0])

    def test_specialty_inside_profile_header_survives_navigation_cleanup(self):
        raw = (b'<html><body><main><header><h1>Dott.ssa Anna Rossi</h1>'
               b'<p><strong>Specialit\xc3\xa0:</strong> Urologia</p></header>'
               b'<article>Ruolo: dirigente medico</article></main></body></html>')
        title, text, _ = m.visible_profile(raw)
        self.assertNotIn('Specialità:', text)
        self.assertIn('Disciplina dichiarata: Urologia', text)
        result = m.analyze_content(person(), text, False, title, False)
        self.assertEqual(result['activities'], ['Urologia'])

    def test_jsonld_physician_medical_specialty_is_extracted(self):
        raw = (b'<title>Dott.ssa Anna Rossi | Directory</title><main><h1>Dott.ssa Anna Rossi</h1></main>'
               b'<script type="application/ld+json">'
               b'{"@context":"https://schema.org","@type":"Physician","name":"Anna Rossi",'
               b'"medicalSpecialty":"Cardiologia"}</script>')
        title, text, _ = m.visible_profile(raw)
        result = m.analyze_content(person(), text, False, title, False)
        self.assertEqual(result['specialties'], [])
        self.assertEqual(result['activities'], ['Cardiologia'])
        self.assertEqual(result['activity_evidence'], ['Disciplina dichiarata: Cardiologia'])

    def test_sitemap_names_use_slug_not_parent_directory(self):
        names = m.Names({'1': person('1', 'Federica', 'Medici')})
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://example.org/medici/federica-de-matteis</loc></url></urlset>'
        spec = {'type': 'sitemap', 'urls': ['https://example.org/sitemap.xml'], 'hosts': ['example.org'], 'profile_pattern': '/medici/'}
        from unittest.mock import Mock
        fetcher = Mock()
        fetcher.get.return_value = ({'final_url': spec['urls'][0]}, raw)
        found, errors = m.discover_one(spec, fetcher, names)
        self.assertEqual(found, set())
        self.assertEqual(errors, [])

    def test_topdoctors_full_url_matches_profile_pattern(self):
        names = m.Names({'1': person()})
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://www.topdoctors.it/dottor/anna-rossi/</loc></url></urlset>'
        spec = {'type': 'sitemap', 'urls': ['https://www.topdoctors.it/doctors.xml'],
                'hosts': ['www.topdoctors.it'], 'profile_pattern': '/dottor/[^/]+/?$',
                'name_pattern': '^/dottor/([^/]+)'}
        from unittest.mock import Mock
        fetcher = Mock(); fetcher.get.return_value = ({'final_url': spec['urls'][0]}, raw)
        found, errors = m.discover_one(spec, fetcher, names)
        self.assertEqual(found, {('1', 'https://www.topdoctors.it/dottor/anna-rossi/', 'profile')})
        self.assertEqual(errors, [])

    def test_institutional_profile_with_title_prefix_matches_name(self):
        names = m.Names({'1': person('1', 'Stefano', 'Bandiera')})
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://www.ior.it/curarsi-al-rizzoli/dr-stefano-bandiera</loc></url></urlset>'
        spec = {'type': 'sitemap', 'urls': ['https://www.ior.it/people.xml'],
                'hosts': ['www.ior.it'],
                'profile_pattern': r'^https://www\.ior\.it/curarsi-al-rizzoli/(?:dr|drssa)-[^/]+$',
                'name_pattern': r'^/curarsi-al-rizzoli/(?:dr|drssa)-(.+)$'}
        from unittest.mock import Mock
        fetcher = Mock(); fetcher.get.return_value = ({'final_url': spec['urls'][0]}, raw)
        found, errors = m.discover_one(spec, fetcher, names)
        self.assertEqual(found, {('1', 'https://www.ior.it/curarsi-al-rizzoli/dr-stefano-bandiera', 'profile')})
        self.assertEqual(errors, [])

    def test_sitemap_legacy_host_is_rewritten_to_canonical_profile(self):
        names = m.Names({'1': person()})
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://www.legacy.test/123/anna-rossi</loc></url></urlset>'
        spec = {'id': 'canonical', 'type': 'sitemap', 'urls': ['https://canonical.test/sitemap.xml'],
                'hosts': ['www.legacy.test'],
                'profile_pattern': r'^https://www\.legacy\.test/[0-9]+/[a-z-]+$',
                'name_pattern': r'^/[0-9]+/([^/]+)$',
                'host_rewrite': {'www.legacy.test': 'canonical.test'}}
        from unittest.mock import Mock
        fetcher = Mock(); fetcher.get.return_value = ({'final_url': spec['urls'][0]}, raw)
        found, errors = m.discover_one(spec, fetcher, names)
        self.assertEqual(found, {('1', 'https://canonical.test/123/anna-rossi', 'profile')})
        self.assertEqual(errors, [])

    def test_html_title_when_heading_is_missing(self):
        title, text, _ = m.visible_profile(b'<title>Anna Rossi - Ospedale</title><main>Specializzata in Cardiologia</main>')
        self.assertEqual(m.analyze_content(person(), text, False, title, False)['specialties'], ['Cardiologia'])

    def test_lock_excludes_second_process_and_releases(self):
        with tempfile.TemporaryDirectory() as folder:
            first = m.RunLock(folder)
            try:
                with self.assertRaises(ValueError):
                    m.RunLock(folder)
            finally:
                first.close()
            second = m.RunLock(folder)
            second.close()

    def test_standard_specialties_previously_missing_from_dictionary(self):
        for label in ['Anatomia Patologica', 'Genetica Medica', 'Cardiochirurgia', 'Allergologia ed Immunologia Clinica']:
            with self.subTest(label=label):
                e = m.analyze_content(person(), 'Specializzata in ' + label + ' con lode.', False, 'Anna Rossi', False)
                self.assertEqual(len(e['specialties']), 1)
        self.assertEqual(m.core.normalize_specialty("Malattie dell'Apparato Cardiovascolare (70/70 e lode)"), 'Cardiologia')

    def test_nonmedical_degree_is_not_a_medical_specialty(self):
        e = m.analyze_content(person(), 'Specializzata in Diritto Tributario.', False, 'Anna Rossi', False)
        self.assertEqual(e['specialties'], [])
        self.assertIn('tassonomia', e['reason'])

    def test_biographical_birth_date(self):
        self.assertEqual(str(m.core.cv_birth_date('Anna Rossi nasce a Milano il 3 aprile 1980.')), '1980-04-03')
        self.assertEqual(str(m.core.cv_birth_date('Nata a Milano il 3/4/1980')), '1980-04-03')

    def test_clinical_role_is_not_a_degree(self):
        e = m.analyze_content(person(), 'Specializzazione\nRadiologo', False, 'Anna Rossi', False)
        self.assertEqual(e['specialties'], [])
        self.assertEqual(e['activities'], ['Radiodiagnostica'])

    def test_related_doctors_are_excluded(self):
        text = 'Anna Rossi\nStessa specialità\nLuca Verdi specializzato in Cardiologia.'
        self.assertEqual(m.analyze_content(person(), text, False, 'Anna Rossi', False)['specialties'], [])

    def test_header_name_identifies_html_cv(self):
        e = m.analyze_content(person(), 'Curriculum\nSpecializzata in Cardiologia.', True, 'Anna Rossi', False)
        self.assertEqual(e['specialties'], ['Cardiologia'])

    def test_queue_resume_and_new_person_for_processed_url(self):
        with tempfile.TemporaryDirectory() as folder:
            store = m.Store(folder)
            fetcher = m.Fetcher(folder)
            names = m.Names({'1': person(), '2': person('2')})
            url = 'https://example.org/anna'
            store.add('1', url, 'test')
            info = {'file': 'fake.html', 'final_url': url, 'sha': 'hash'}
            outcome = {'info': info, 'text': 'Anna Rossi. Specializzata in Cardiologia.', 'cv': False, 'title': 'Anna Rossi', 'links': []}
            try:
                with patch.object(m, 'process_url', return_value=outcome) as download:
                    self.assertEqual(m.process_queue(store, fetcher, names), 1)
                    self.assertEqual(m.process_queue(store, fetcher, names), 0)
                    store.add('2', url, 'new source')
                    self.assertEqual(m.process_queue(store, fetcher, names), 1)
                    self.assertEqual(download.call_count, 2)
                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM evidence').fetchone()[0], 2)
            finally:
                store.close()

    def test_robots_merges_groups_and_respects_longest_rule(self):
        policy = m.RobotPolicy()
        policy.parse(['User-agent: *', 'Allow: /', 'Disallow: /private/',
                      'User-agent: *', 'Disallow: /search?*', 'Allow: /private/public$'])
        self.assertFalse(policy.can_fetch(m.UA, 'https://example.org/private/a'))
        self.assertFalse(policy.can_fetch(m.UA, 'https://example.org/search?q=abc'))
        self.assertTrue(policy.can_fetch(m.UA, 'https://example.org/private/public'))
        self.assertFalse(policy.can_fetch(m.UA, 'https://example.org/private/public/x'))

    def test_profile_heading_survives_semantic_header(self):
        title, text, _ = m.visible_profile(b'<header><h1>Anna Rossi</h1></header><main>Specializzata in Cardiologia</main>')
        self.assertEqual(title, 'Anna Rossi')
        self.assertEqual(m.analyze_content(person(), text, False, title, False)['specialties'], ['Cardiologia'])

    def test_names_punctuation_longest_and_duplicates(self):
        names = m.Names({'1': person(), '2': person('2', 'Maria Anna'), '3': person('3'),
                         '4': person('4', 'Giulia', "De_Càrlo")})
        self.assertEqual(names.match('Dott.ssa Maria Anna Rossi'), {'2'})
        self.assertEqual(names.match('rossi-anna'), {'1', '3'})
        self.assertEqual(names.match('/it/giulia-de-carlo/'), {'4'})
        self.assertTrue(names.ambiguous(person()))

    def test_nominal_identity_never_becomes_confirmed(self):
        text = 'Anna Rossi\nSpecializzata in Cardiologia.'
        e = m.analyze_content(person(), text, False, 'Anna Rossi', False)
        self.assertEqual(e['identity'], 'solo_nome_completo')
        e.update(kind='Profilo', path='', url='https://example.org')
        row = m.summarize([e])
        self.assertEqual(row[0], 'SPECIALITA_CON_IDENTITA_NOMINALE')
        self.assertEqual(row[2], '')
        self.assertEqual(row[3], 'Cardiologia')

    def test_birth_date_and_ambiguity(self):
        text = 'Curriculum vitae Anna Rossi. Medico. Specializzata in Cardiologia.'
        self.assertEqual(m.analyze_content(person(), text, True, '', True)['identity'], 'omonimia')
        e = m.analyze_content(person(), text + ' Data di nascita 3 aprile 1980.', True, '', True)
        self.assertEqual(e['identity'], 'anagrafica_concordante')
        e = m.analyze_content(person(), text + ' Data di nascita 3 aprile 1981.', True, '', False)
        self.assertEqual(e['identity'], 'incompatibile')
        self.assertEqual(e['specialties'], [])

    def test_sitemap_ignores_images(self):
        raw = b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9" xmlns:image="http://www.google.com/schemas/sitemap-image/1.1"><url><loc>https://example.org/anna-rossi</loc><image:image><image:loc>https://example.org/photo.jpg</image:loc></image:image></url></urlset>'
        self.assertEqual(m.sitemap_entries(raw), (False, ['https://example.org/anna-rossi']))
        with self.assertRaises(ValueError):
            m.sitemap_entries(b'<html>blocked</html>')

    def test_never_infer_from_occupation_desired(self):
        e = m.analyze_content(person(), 'Anna Rossi\nOccupazione desiderata: Medico specializzato in Cardiologia.', False, 'Anna Rossi', False)
        self.assertEqual(e['specialties'], [])

    def test_full_export_preserves_rows_and_resume(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            inp, out = root / 'input.xlsx', root / 'report.xlsx'
            wb = Workbook(); ws = wb.active; ws.title = 'Foglio1'
            ws.append(['Pers_Id', 'Pers_Cognome', 'Pers_Nome', 'Pers_DataNascita', 'Pers_CodFis', 'Indirizzi_Citta', 'Medico_Id'])
            ws.append([1, 'Rossi', 'Anna', '03/04/1980', '', 'Roma', 0])
            ws.append([2, 'Verdi', 'Luca', '02/02/1981', '', 'Milano', 0]);wb.save(inp)
            before = m.sha_file(inp)
            store = m.Store(root / 'state')
            try:
                store.import_people(inp, 'Foglio1')
                store.import_people(inp, 'Foglio1')
                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM people').fetchone()[0], 2)
                store.add('1', 'https://example.org/anna', 'test')
                store.add('1', 'https://example.org/anna', 'test')
                store.add('2', 'https://example.org/anna', 'test')
                self.assertEqual(store.db.execute('SELECT COUNT(*) FROM urls').fetchone()[0], 1)
                report = m.export_report(store, inp, 'Foglio1', out)
                self.assertEqual(report['total'], 2)
                self.assertEqual(report['states']['IN_CODA'], 2)
                m.export_report(store, inp, 'Foglio1', out)
                result = load_workbook(out, read_only=True)
                self.assertEqual(len(list(result['Foglio1'].values)), 3)
                result.close()
                self.assertEqual(before, m.sha_file(inp))
                foreign = root / 'foreign.xlsx'; foreign.write_bytes(b'original')
                with self.assertRaises(ValueError):
                    m.export_report(store, inp, 'Foglio1', foreign)
                self.assertEqual(foreign.read_bytes(), b'original')
            finally:
                store.close()

    def test_cache_integrity_and_no_redownload(self):
        with tempfile.TemporaryDirectory() as temp:
            fetcher = m.Fetcher(temp)
            with patch.object(fetcher, 'allowed', return_value=True), patch.object(fetcher, '_request', return_value=(200, {'Content-Type': 'text/html'}, b'<h1>Anna Rossi</h1>')) as call:
                info, _ = fetcher.get('https://example.org/anna')
                fetcher.get('https://example.org/anna')
                self.assertEqual(call.call_count, 1)
                (fetcher.folder / info['file']).write_bytes(b'corrupt')
                fetcher.get('https://example.org/anna')
                self.assertEqual(call.call_count, 2)

    def test_redirect_rechecks_robots(self):
        with tempfile.TemporaryDirectory() as temp:
            fetcher = m.Fetcher(temp)
            with patch.object(fetcher, 'allowed', side_effect=[True, False]) as policy, patch.object(fetcher, '_request', return_value=(302, {'Location': 'https://other.example/cv'}, b'')) as request:
                with self.assertRaisesRegex(m.FetchError, 'ROBOTS_NON_CONSENTE'):
                    fetcher.get('https://example.org/anna')
                self.assertEqual(policy.call_count, 2)
                self.assertEqual(request.call_count, 1)

    def test_robots_does_not_allow_on_server_error(self):
        with tempfile.TemporaryDirectory() as temp:
            fetcher = m.Fetcher(temp)
            with patch.object(fetcher, '_request', return_value=(503, {}, b'')):
                with self.assertRaises(m.FetchError):
                    fetcher.allowed('https://example.org/a')

    def test_missing_source_is_not_a_negative_search_result(self):
        self.assertEqual(m.summarize([])[0], 'NON_COPERTO_DALLE_FONTI')
        self.assertEqual(m.summarize([], pending=True)[0], 'IN_CODA')
        self.assertEqual(m.summarize([], errors=True)[0], 'FONTE_NON_ACCESSIBILE')


if __name__ == '__main__':
    unittest.main()
